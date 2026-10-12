from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest

from worker import workbook as wb
from worker.guard import PathRefused, check_path

FIXTURES = Path(__file__).parent / "fixtures"
SOURCES = sorted(
    p for p in FIXTURES.rglob("*.xlsx") if "answer_key" not in p.parts and not p.name.endswith(".answer.json")
)
IDS = [str(p.relative_to(FIXTURES)) for p in SOURCES]


@pytest.mark.parametrize("path", SOURCES, ids=IDS)
def test_list_and_look_are_bounded(path: Path) -> None:
    info = wb.list_sheets(path)
    assert info["kind"] == "xlsx"
    assert info["sheets"]
    json.dumps(info)
    for s in info["sheets"]:
        assert {
            "name",
            "dimensions",
            "max_row",
            "max_column",
            "hidden",
            "merged_count",
            "has_pivot",
        } <= s.keys()
        out = wb.look(path, s["name"], 1, 500, 1, 500)
        json.dumps(out)
        lines = out["text"].splitlines()
        assert len([ln for ln in lines if re.match(r"^\s*\d+ \|", ln)]) <= wb.LOOK_MAX_ROWS
        assert out["last_col"] - out["first_col"] + 1 <= wb.LOOK_MAX_COLS
        assert out["last_row"] - out["first_row"] + 1 <= wb.LOOK_MAX_ROWS
        rr = wb.read_range(path, s["name"], "A1:Z200", max_cells=50)
        assert len(rr["cells"]) <= 50
        json.dumps(rr)


def _col_a_prefixes(text: str) -> set[str]:
    found = set()
    for line in text.splitlines():
        m = re.match(r"^\s*\d+ \| (\.*)\S", line)
        if m:
            found.add(m.group(1))
    return found


def test_compact_pivot_shows_three_indent_levels() -> None:
    path = FIXTURES / "messy" / "investors_compact_pivot.xlsx"
    name = wb.list_sheets(path)["sheets"][0]["name"]
    out = wb.look(path, name, 1, 80, 1, 6)
    assert len(_col_a_prefixes(out["text"])) >= 3


def test_pivot_fixture_shows_merged_note() -> None:
    path = FIXTURES / "messy" / "investors_pivot.xlsx"
    name = wb.list_sheets(path)["sheets"][0]["name"]
    out = wb.look(path, name, 1, 30, 1, 16)
    assert "[merged " in out["text"]


def test_multi_sheet_lists_summary_and_notes() -> None:
    info = wb.list_sheets(FIXTURES / "messy" / "investors_multi_sheet.xlsx")
    names = {s["name"].lower() for s in info["sheets"]}
    assert len(names) > 1
    assert "summary" in names and "notes" in names


@pytest.mark.parametrize("path", SOURCES, ids=IDS)
def test_read_range_matches_openpyxl_header_row(path: Path) -> None:
    book = openpyxl.load_workbook(path)
    row = None
    for ws in book.worksheets:
        row = next(
            (r for r in ws.iter_rows(min_row=1, max_row=40) if sum(isinstance(c.value, str) for c in r) >= 2),
            None,
        )
        if row is not None:
            break
    assert row is not None
    expected = {
        c.coordinate: c.value for c in row if isinstance(c.value, str) and not c.value.startswith("=")
    }
    r = row[0].row
    got = wb.read_range(path, ws.title, f"A{r}:{openpyxl.utils.get_column_letter(ws.max_column)}{r}")
    got_map = {c["addr"]: c["value"] for c in got["cells"]}
    for addr, text in expected.items():
        assert got_map[addr] == text


def test_read_range_reports_formulas_and_cap(tmp_path: Path) -> None:
    p = tmp_path / "f.xlsx"
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = "=SUM(A1:A2)"
    book.save(p)
    out = wb.read_range(p, ws.title, "A1:A3")
    assert out["cells"][2]["formula"] == "=SUM(A1:A2)"
    capped = wb.read_range(p, ws.title, "A1:A3", max_cells=2)
    assert capped["truncated"] and len(capped["cells"]) == 2


def test_find_and_column_profile() -> None:
    path = FIXTURES / "set7" / "Commitment_Register.xlsx"
    name = wb.list_sheets(path)["sheets"][0]["name"]
    prof = wb.column_profile(path, name, 1, 1, 50)
    assert prof["count"] + prof["blanks"] == prof["last_row"] - prof["first_row"] + 1
    assert len(prof["samples"]) <= 5
    first = prof["samples"][0]
    hits = wb.find(path, name, str(first).upper())
    assert hits["matches"]


def test_csv_image_and_unknown(tmp_path: Path) -> None:
    c = tmp_path / "data.csv"
    c.write_text("id,name,date\nA-001,Jane Roe,2024-01-31\nA-002,John Doe,2024-02-28\n")
    info = wb.list_sheets(c)
    assert info["kind"] == "csv" and len(info["sheets"]) == 1
    sheet = info["sheets"][0]["name"]
    assert wb.column_profile(c, sheet, "A", 2, 3)["shape"] == "id-like"
    assert wb.column_profile(c, sheet, "C", 2, 3)["shape"] == "date-like"
    assert wb.read_range(c, sheet, "A1:C1")["cells"][1]["value"] == "name"

    img = tmp_path / "pic.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    assert wb.list_sheets(img) == {"kind": "image", "bytes": 28}

    junk = tmp_path / "x.bin"
    junk.write_bytes(b"\x00\x01\x02\x03" * 10)
    assert wb.list_sheets(junk) == {"kind": "unknown"}


def test_guard_refuses_answer_keys_and_outside_roots(tmp_path: Path) -> None:
    roots = [FIXTURES]
    for bad in [
        *(FIXTURES / "messy").glob("*.answer.json"),
        FIXTURES / "set7" / "answer_key" / "GL.csv",
        tmp_path / "other.xlsx",
    ]:
        with pytest.raises(PathRefused):
            check_path(bad, roots)
    ok = check_path(FIXTURES / "set7" / "GP_Ledger.xlsx", roots)
    assert ok.name == "GP_Ledger.xlsx"
