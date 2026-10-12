"""F3: staged tables, the paused gate, open questions and decisions survive a restart of the process."""

from __future__ import annotations

import datetime as dt
import time
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from id_decision import decide_agent_ids
from scripted_model import ScriptedModel, call, say
from worker.registry import Registry
from worker.service import RunManager, ServiceError
from worker.staging import Stage

COLS = {"Investor_ID": "A", "Investor_Name": "B", "Investor_Type": "C"}


def _lift(**extra: Any) -> dict[str, Any]:
    return call("lift", path="inv.xlsx", sheet="Sheet", template="Investors", columns=COLS,
                constants={"GP_ID": "GP-001"}, header_row=1, first_row=2, why="copy", **extra)


def _manager(state: Path, script: list[Any]) -> RunManager:
    return RunManager(state, lambda: ScriptedModel(script=list(script)))


def _inbox(state: Path) -> None:
    (state / "inbox").mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["ID", "Name", "Type", "Admitted"])
    ws.append(["I1", "Alpha Capital", "Pension", dt.datetime(2024, 3, 1, 9, 30)])
    ws.append(["I2", "Beta Partners", "Endowment", dt.date(2024, 4, 2)])
    ws.append(["I3", "Gamma Trust", "Family", None])
    book.save(state / "inbox" / "inv.xlsx")


def _wait(m: RunManager, run_id: str, *statuses: str, timeout: float = 60) -> dict[str, Any]:
    end = time.time() + timeout
    while time.time() < end:
        run = m.get(run_id)
        if run["status"] in statuses:
            return run
        time.sleep(0.05)
    raise AssertionError(f"run stayed {m.get(run_id)['status']}: {m.get(run_id)['error']}")


def test_F3_a_paused_gate_survives_a_restart_and_approves_exactly_once(tmp_path: Path) -> None:
    state = tmp_path / "state"
    _inbox(state)
    decide_agent_ids(state / "ids", "Investors")
    first = _manager(state, [say("lifted", _lift()),
                             say("write", call("write_upload", stage_id="s1", template="Investors", why="deliver"))])
    run_id = first.create_run("inv.xlsx", "sponsor-a", "GP-001")["id"]
    waiting = _wait(first, run_id, "needs_approval")
    gate = waiting["gates"][0]
    assert gate["name"] == "write_upload" and gate["stage"]["rows"] == 3
    first.close()

    second = _manager(state, [say("done")])  # a new process: nothing in memory
    again = second.get(run_id)
    assert again["status"] == "needs_approval" and again["gates"][0]["hash"] == gate["hash"]
    approved = second.decide(run_id, gate["hash"], "approve", "checked")
    assert _wait(second, run_id, "idle")["gates"] == []
    uploads = [o["name"] for o in approved["outputs"] + second.get(run_id)["outputs"] if o["kind"] == "upload"]
    assert set(uploads) == {"Investors_Upload.xlsx"}
    sheet = openpyxl.load_workbook(second.output_path(run_id, "Investors_Upload.xlsx")).active
    rows = [[c.value for c in r] for r in sheet.iter_rows()]
    assert len(rows) == 4 and rows[1][1:3] == ["I1", "Alpha Capital"]
    assert len(Registry(state, "sponsor-a").entries("Investors")) == 3
    with pytest.raises(ServiceError, match="nothing is waiting"):
        second.decide(run_id, gate["hash"], "approve", "again")  # not twice
    assert len(Registry(state, "sponsor-a").entries("Investors")) == 3
    second.close()


def test_F3_staged_tables_keep_their_cell_types_and_meta(tmp_path: Path) -> None:
    out = tmp_path / "out"
    st = Stage(out)
    sid = st.lift(_inbox_file(tmp_path), "Sheet", {"Investor_ID": "A", "Investor_Name": "B", "Investor_Type": "D"},
                  1, 2, template="Investors")
    st.annotate(sid, "note", {"when": dt.datetime(2025, 1, 2, 3, 4), "ids": {"a", "b"}})
    fresh = Stage(out)
    assert fresh.restore() == [sid]
    assert fresh.rows(sid) == st.rows(sid)
    assert isinstance(fresh.rows(sid)[0]["Investor_Type"], dt.datetime)
    assert fresh.meta(sid)["note"] == {"when": dt.datetime(2025, 1, 2, 3, 4), "ids": {"a", "b"}}
    assert fresh.summary(sid) == st.summary(sid)
    nxt = fresh.lift(_inbox_file(tmp_path), "Sheet", {"Investor_ID": "A"}, 1, 2, template="Investors")
    assert nxt == "s2"  # the counter continues, no id is reused


def _inbox_file(tmp_path: Path) -> Path:
    _inbox(tmp_path / "state")
    return tmp_path / "state" / "inbox" / "inv.xlsx"


def test_F3_an_open_question_and_the_staged_table_are_there_after_a_restart(tmp_path: Path) -> None:
    state = tmp_path / "state"
    _inbox(state)
    first = _manager(state, [say("lifted", _lift()), say("Alpha Capital has no Source key. Add one?")])
    run_id = first.create_run("inv.xlsx", "sponsor-a", "GP-001")["id"]
    asked = _wait(first, run_id, "idle")
    assert asked["messages"][-1]["text"].startswith("Alpha Capital has no Source key")
    first.close()

    second = _manager(state, [say("write", call("write_upload", stage_id="s1", template="Investors",
                                                  why="the human said go ahead"))])
    assert second.get(run_id)["messages"][-1]["text"].startswith("Alpha Capital has no Source key")
    sent = second.send(run_id, "No, go ahead")
    assert sent["status"] in ("working", "needs_approval")
    waiting = _wait(second, run_id, "needs_approval")
    assert waiting["gates"][0]["stage"]["rows"] == 3  # the stage id still resolves in the new process
    second.close()


def test_F3_a_run_that_was_busy_when_the_process_died_can_be_continued(tmp_path: Path) -> None:
    state = tmp_path / "state"
    _inbox(state)
    first = _manager(state, [say("lifted", _lift()), say("done")])
    run_id = first.create_run("inv.xlsx", "sponsor-a", "GP-001")["id"]
    _wait(first, run_id, "idle")
    first.close()
    path = state / "runs" / run_id / "run.json"
    path.write_text(path.read_text(encoding="utf-8").replace('"status": "idle"', '"status": "working"'), encoding="utf-8")
    second = _manager(state, [say("write", call("write_upload", stage_id="s1", template="Investors", why="go"))])
    view = second.get(run_id)
    assert view["status"] == "interrupted" and "restarted" in view["error"]
    second.send(run_id, "continue")
    assert _wait(second, run_id, "needs_approval")["gates"][0]["args"]["stage_id"] == "s1"
    second.close()


def test_F3_id_reservations_and_supplied_values_are_carried_over(tmp_path: Path) -> None:
    from worker.context import RunContext
    from worker.idgen import IdScheme
    from worker.jobstate import JobState

    out = tmp_path / "out"
    ctx = RunContext(sponsor="sponsor-a", gp_id="GP-001", inbox_roots=[], out_dir=out,
                     audit=None, stage=Stage(out), id_root=tmp_path / "ids")  # type: ignore[arg-type]
    scheme = IdScheme(ctx.id_root, "sponsor-a", "Investors", {"mode": "sequential", "prefix": "I-", "width": 3})
    rsv = scheme.reserve({"I-002"})
    assert [rsv.take(), rsv.take()] == ["I-001", "I-003"]
    ctx.reservations["s2"] = (scheme, rsv, "Investor_ID", "Source_Investor_Key")
    ctx.job_feedback.append({"kind": "job", "action": "override"})
    ctx.instruction = "build me the sheet"
    state = JobState(out)
    state.specs["s1"] = {"kind": "lift", "columns": {"Investor_ID": "A"}}
    state.supplied.set_row(5, "Country", "UK", who="human", reason="asked")
    state.save(ctx)

    ctx2 = RunContext(sponsor="sponsor-a", gp_id="GP-001", inbox_roots=[], out_dir=out,
                      audit=None, stage=Stage(out), id_root=tmp_path / "ids")  # type: ignore[arg-type]
    state2 = JobState(out)
    assert state2.restore(ctx2)
    _scheme, rsv2, id_field, src_field = ctx2.reservations["s2"]
    assert (id_field, src_field, rsv2.first, rsv2.next, rsv2.issued) == (
        "Investor_ID", "Source_Investor_Key", 1, 4, ["I-001", "I-003"])
    assert rsv2.take() == "I-004"  # skipped and issued numbers are never handed out again
    assert state2.specs["s1"]["columns"] == {"Investor_ID": "A"} and state2.supplied.rows[(5, "Country")] == "UK"
    assert ctx2.job_feedback == [{"kind": "job", "action": "override"}] and ctx2.instruction == "build me the sheet"
