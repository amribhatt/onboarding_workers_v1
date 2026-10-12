from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from langgraph.checkpoint.memory import InMemorySaver

from id_decision import decide_agent_ids
from scripted_model import ScriptedModel, approve, call, say
from worker.agent import (
    GateRefused,
    build_agent,
    check_approval,
    decide,
    pending_actions,
    pending_gate,
    start,
)
from worker.audit import AuditLog
from worker.context import RunContext
from worker.registry import Registry
from worker.review import review
from worker.staging import Stage

LONG = "X" * 21


def _inbox(tmp_path: Path, rows: list[list[str]], name: str = "f.xlsx") -> Path:
    d = tmp_path / "inbox"
    d.mkdir(exist_ok=True)
    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["ID", "Name", "Type"])
    for r in rows:
        ws.append(r)
    book.save(d / name)
    return d


def _ctx(tmp_path: Path, inbox: Path, registry: Registry | None = None) -> RunContext:
    out = tmp_path / "out"
    return RunContext(
        sponsor="sponsor-a",
        gp_id="GP-001",
        inbox_roots=[inbox],
        out_dir=out,
        audit=AuditLog(out / "audit.jsonl"),
        stage=Stage(out),
        registry=registry,
        id_root=tmp_path / "ids",
    )


def _lift(constants: dict[str, str] | None = None) -> dict[str, Any]:
    return call(
        "lift",
        path="f.xlsx",
        sheet="Sheet",
        template="Investors",
        columns={"Investor_ID": "A", "Investor_Name": "B", "Investor_Type": "C"},
        constants={"GP_ID": "GP-001", **(constants or {})},
        header_row=1,
        first_row=2,
        why="copy",
    )


def _paused(tmp_path: Path, rows: list[list[str]], extra: list[Any] | None = None, **kw: Any) -> tuple[Any, RunContext]:
    ctx = _ctx(tmp_path, _inbox(tmp_path, rows), kw.get("registry"))
    decide_agent_ids(ctx.id_root, "Investors")
    script = [say("lift", _lift()), *(extra or []), say("write", call("write_upload", stage_id=kw.get("stage", "s1"),
              template="Investors", why="deliver")), say("done")]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    result = start(agent, ctx, "build", "t1")
    assert [g["name"] for g in pending_gate(result)] == ["write_upload"]
    return agent, ctx


def test_unresolved_err_blocks_approval_and_nothing_is_written(tmp_path: Path) -> None:
    agent, ctx = _paused(tmp_path, [["I1", "Alpha", "LP"], [LONG, "Beta", "LP"]])
    r = review(ctx, "s1")
    assert r["err_count"] == 1 and r["open_err"] and not r["can_approve"]
    with pytest.raises(GateRefused, match="ERR must be resolved"):
        decide(agent, ctx, "t1", {"type": "approve"}, "try", {"reviewer": "reviewer-a"})
    assert not list((tmp_path / "out").glob("*_Upload*"))
    # the refusal happens before the decision is audited as an approval
    assert not any(e["action"] == "gate.approve" for e in ctx.audit.read())


def test_warn_needs_a_reason_and_is_audited(tmp_path: Path) -> None:
    rows = [["I1", "Alpha", "LP"], ["I2", "Beta", "LP"]]
    ctx = _ctx(tmp_path, _inbox(tmp_path, rows))
    decide_agent_ids(ctx.id_root, "Investors")
    script = [say("lift", _lift({"GP_ID": "GP-009"})), say("write", call("write_upload", stage_id="s1",
              template="Investors", why="deliver")), say("done")]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "build", "t1")
    r = review(ctx, "s1")
    warn = [f for f in r["findings"] if f["code"] == "gp_id_differs"]
    assert warn and warn[0]["severity"] == "WARN" and r["err_count"] == 0
    with pytest.raises(GateRefused, match="WARN need an acknowledgement"):
        decide(agent, ctx, "t1", {"type": "approve"}, meta={"reviewer": "reviewer-a"})
    with pytest.raises(GateRefused):  # an empty reason is not an acknowledgement
        decide(agent, ctx, "t1", {"type": "approve"}, meta={"reviewer": "reviewer-a"}, acknowledged={warn[0]["id"]: "   "})
    decide(agent, ctx, "t1", {"type": "approve"}, "ok", {"reviewer": "alex"},
           {warn[0]["id"]: "GP-009 is the real GP for this fund"})
    acks = [e for e in ctx.audit.read() if e["action"] == "finding.acknowledged"]
    assert len(acks) == 1 and acks[0]["rationale"] == "GP-009 is the real GP for this fund"
    assert acks[0]["inputs"]["reviewer"] == "alex"
    assert (tmp_path / "out" / "Investors_Upload.xlsx").exists()


def test_resolving_an_err_changes_output_and_audit(tmp_path: Path) -> None:
    resolve = say(
        "resolve",
        call(
            "resolve_conflicts",
            stage_id="s1",
            decisions={"Sheet!3": {"action": "correct", "reason": "typo, real ID is I2", "values": {"Investor_ID": "I2"}}},
            why="human corrected the ID",
        ),
    )
    agent, ctx = _paused(tmp_path, [["I1", "Alpha", "LP"], [LONG, "Beta", "LP"]], [resolve], stage="s2")
    r = review(ctx, "s2")
    assert r["err_count"] == 0 and r["can_approve"]
    approve(agent, ctx, "t1")
    book = openpyxl.load_workbook(tmp_path / "out" / "Investors_Upload.xlsx")
    ids = [row[1] for row in book.active.iter_rows(min_row=2, values_only=True)]
    assert ids == ["I1", "I2"]
    resolved = [e for e in ctx.audit.read() if e["action"] == "conflict.resolved"]
    assert resolved[0]["rationale"] == "typo, real ID is I2" and resolved[0]["inputs"]["row"] == "Sheet!3"


def test_exclude_a_flagged_row_with_a_reason(tmp_path: Path) -> None:
    resolve = say(
        "resolve",
        call(
            "resolve_conflicts",
            stage_id="s1",
            decisions={"Sheet!3": {"action": "exclude_incorrect", "reason": "test row left in the file"}},
            why="human excluded it",
        ),
    )
    _agent, ctx = _paused(tmp_path, [["I1", "Alpha", "LP"], [LONG, "Beta", "LP"]], [resolve], stage="s2")
    r = review(ctx, "s2")
    assert r["can_approve"] and r["rows_out"] == 1
    assert r["reconciliation"]["excluded_by_analyst"] == 1


def test_reconciliation_numbers_add_up(tmp_path: Path) -> None:
    reg = Registry(tmp_path, "sponsor-a")
    reg.add("Investors", [{"key": "I1", "name": "Alpha"}], source="app batch", batch="b1")
    rows = [["I1", "Alpha", "LP"], ["I2", "Beta", "LP"], ["I2", "Beta", "LP"], ["I3", "Gamma", "LP"], ["I4", "Delta", "LP"]]
    ctx = _ctx(tmp_path, _inbox(tmp_path, rows), reg)
    script = [
        say("lift", _lift()),
        say("distinct", call("distinct_rows", stage_id="s1", key_fields=["Investor_ID"], why="one row per investor")),
        say("match", call("match_existing", stage_id="s2", why="dedup")),
        say(
            "resolve",
            call(
                "resolve_conflicts",
                stage_id="s3",
                decisions={"Sheet!5": {"action": "exclude_duplicate", "reason": "analyst: duplicate of I3"}},
                why="human decided",
            ),
        ),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "build", "t1")
    rec = review(ctx, "s4")["reconciliation"]
    assert rec == {
        "input_rows": 5,
        "collapsed_by_dedup": 1,
        "excluded_existing": 1,
        "excluded_by_analyst": 1,
        "dropped_duplicates": 0,
        "output_rows": 2,
        "unexplained": 0,
        "balanced": True,
    }


def test_duplicate_ids_in_the_output_are_an_err_unless_dropped(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path, _inbox(tmp_path, [["I1", "Alpha", "LP"], ["I1", "Alpha", "LP"]]))
    script = [say("lift", _lift()), say("done")]
    start(build_agent(ctx, ScriptedModel(script=script), InMemorySaver()), ctx, "go", "t1")
    assert review(ctx, "s1")["open_err"]
    dropped = review(ctx, "s1", drop_duplicates=True, duplicate_key=["Investor_ID"])
    assert dropped["can_approve"] and dropped["rows_out"] == 1
    assert dropped["reconciliation"]["dropped_duplicates"] == 1


def test_sponsor_level_override_turns_a_warn_into_an_err(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    base = tmp_path / "templates"
    (base / "sponsors" / "sponsor-a").mkdir(parents=True)
    (base / "sponsors" / "sponsor-a" / "levels.json").write_text(json.dumps({"gp_id_differs": "ERR"}), encoding="utf-8")
    monkeypatch.setenv("WORKER_TEMPLATES_DIR", str(base))
    ctx = _ctx(tmp_path, _inbox(tmp_path, [["I1", "Alpha", "LP"]]))
    script = [say("lift", _lift({"GP_ID": "GP-009"})), say("done")]
    start(build_agent(ctx, ScriptedModel(script=script), InMemorySaver()), ctx, "go", "t1")
    r = review(ctx, "s1")
    assert [f["severity"] for f in r["findings"] if f["code"] == "gp_id_differs"] == ["ERR"]
    assert not r["can_approve"]


def test_check_approval_ignores_other_gates(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path, _inbox(tmp_path, []))
    check_approval(ctx, [{"name": "propose_rule", "args": {}}])  # no raise
    assert pending_actions is not None
