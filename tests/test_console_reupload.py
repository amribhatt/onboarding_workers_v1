"""Replace a job's source file at any stage: decisions follow their rows by content, the rest is re-opened and listed,
the swap is audited before it takes effect, and an approved job never changes. Fictitious data only."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from worker import console_issues as CI
from worker import console_reupload as RU
from worker.console import Console

WHO = "Priya Shah"


@pytest.fixture
def c(tmp_path: Path) -> Console:
    return Console(tmp_path / "inbox")


def job5(c: Console) -> dict[str, Any]:
    c.act("start|Sponsor A|Investors")
    c.act("ask_set|cov|all")
    c.act("ask_sample")
    c.act("ask_go")
    c.act("ex|yes")
    c.act("ex|file")
    c.act("go")
    j = c.job()
    assert j is not None
    return j


def edited(j: dict[str, Any], tmp: Path, name: str, *, insert_top: bool = False, set_cell: tuple[int, int, str] | None = None) -> bytes:
    """The job's current file with a row inserted at the top and/or one cell changed (sheet row, column, value)."""
    wb = openpyxl.load_workbook(j["up"]["path"])
    ws = wb[wb.sheetnames[0]]
    if insert_top:
        ws.insert_rows(2)
        for col, v in enumerate(["HS-NEW-1", "A Brand New Investor LLC"], start=1):
            ws.cell(2, col).value = v
    if set_cell:
        ws.cell(set_cell[0] + (1 if insert_top else 0), set_cell[1]).value = set_cell[2]
    out = tmp / name
    wb.save(out)
    return out.read_bytes()


def _audit(c: Console, j: dict[str, Any]) -> list[dict[str, Any]]:
    path = c.job_dir(j) / "audit.jsonl"
    return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]


def decide_two(c: Console, j: dict[str, Any]) -> None:
    g = next(g for g in CI.model(c, j)["groups"] if g["cause"] == "idlong" and g["state"] == "open")
    CI.apply(c, j, [{"id": g["id"], "option": "out", "rows": [9]}], WHO)
    g = next(g for g in CI.model(c, j)["groups"] if g["cause"] == "idlong" and g["state"] == "open")
    CI.apply(c, j, [{"id": g["id"], "option": "edit", "value": "INV-B", "rows": [10]}], WHO)
    assert j["dec"]["8"].startswith("exclude") and j["dec"]["9"].endswith("INV-B")


def test_decisions_follow_their_rows_when_rows_move_and_a_changed_row_is_reopened(c: Console, tmp_path: Path) -> None:
    j = job5(c)
    decide_two(c, j)
    # one new row on top shifts everything down by one; the old row 10 is changed in the new file
    data = edited(j, tmp_path, "Investors_v2.xlsx", insert_top=True, set_cell=(10, 1, "HS-FIXED-IN-FILE"))
    rep = RU.replace(c, j, "Investors_v2.xlsx", data, WHO, note="client sent a corrected export")
    assert (rep["kept"], rep["reopened"]) == (1, 1)
    assert rep["reopened_listed"][0]["was"] == 9 and "INV-B" in rep["reopened_listed"][0]["what"]
    assert j["dec"]["9"].startswith("exclude") and "8" not in j["dec"]  # old record 8 -> new record 9
    assert j["up"]["file"] == "Investors_v2.xlsx" and j["up"]["own"] is True
    entry = next(a for a in _audit(c, j) if a["action"] == "source.replaced")
    assert entry["inputs"]["reviewer"] == WHO and entry["inputs"]["kept"] == 1 and entry["inputs"]["reopened"] == 1
    assert entry["inputs"]["from"]["sha256"].startswith("sha256:") and entry["inputs"]["to"]["sha256"] != entry["inputs"]["from"]["sha256"]
    assert "client sent a corrected export" in entry["rationale"]
    # the issue history follows: the excluded row is now sheet row 10; the edit lost its row, so it is reverted
    live = [h for h in CI._hist(j) if not h["rev"]]
    assert len(live) == 1 and live[0]["first"] == 10
    assert any(h["rev"] and "replaced" in h["released"]["note"] for h in CI._hist(j))
    # the job was converted again on the new file: 15 rows became 16
    assert j["runs"][-1]["t"]["read"] == 16 and len(j["runs"]) == 2


def test_the_file_can_be_replaced_back_and_forth(c: Console, tmp_path: Path) -> None:
    j = job5(c)
    original = Path(j["up"]["path"]).read_bytes()
    decide_two(c, j)
    RU.replace(c, j, "Investors_v2.xlsx", edited(j, tmp_path, "Investors_v2.xlsx", insert_top=True), WHO)
    assert j["dec"]["9"].startswith("exclude") and j["dec"]["10"].endswith("INV-B")  # both rows moved down by one
    back = RU.replace(c, j, "Investors_back.xlsx", original, WHO)
    assert back["kept"] == 2 and back["reopened"] == 0  # and back up again
    assert j["dec"]["8"].startswith("exclude") and j["dec"]["9"].endswith("INV-B")
    assert [a["action"] for a in _audit(c, j)].count("source.replaced") == 2
    assert len(j["replaced"]) == 2


def test_a_row_that_appears_twice_is_matched_one_to_one(c: Console) -> None:
    old = [{"i": 1, "row": 2, "v": {"a": "x"}}, {"i": 2, "row": 3, "v": {"a": "x"}}, {"i": 3, "row": 4, "v": {"a": "y"}}]
    new = [{"i": 1, "row": 2, "v": {"a": "y"}}, {"i": 2, "row": 3, "v": {"a": "x"}}]
    assert RU._moves(old, new) == {1: 2, 3: 1}  # the first x takes the only x; the second x and nothing else is lost


def test_refusals_leave_the_job_as_it_was(c: Console, tmp_path: Path) -> None:
    j = job5(c)
    same = Path(j["up"]["path"]).read_bytes()
    with pytest.raises(RU.ReuploadError, match="already reads"):
        RU.replace(c, j, "same.xlsx", same, WHO)
    for name, data in (("notes.txt", b"hello"), ("empty.xlsx", b""), ("broken.xlsx", b"not a workbook")):
        with pytest.raises(Exception):  # noqa: B017  the console's own refusal, whichever words it uses
            RU.replace(c, j, name, data, WHO)
    with pytest.raises(Exception):  # noqa: B017  a name is required: no one is asked to type it, but a reviewer must be a person
        RU.replace(c, j, "x.xlsx", edited(j, tmp_path, "x.xlsx", insert_top=True), "local")
    assert j["up"]["file"] == "Sponsor_A_Common_Investors.xlsx" and not j.get("replaced")
    j["approved"] = True
    with pytest.raises(RU.ReuploadError, match="approved"):
        RU.replace(c, j, "late.xlsx", edited(j, tmp_path, "late.xlsx", insert_top=True), WHO)


def test_if_the_audit_cannot_be_written_nothing_changes_and_the_new_file_is_not_kept(c: Console, tmp_path: Path) -> None:
    j = job5(c)
    decide_two(c, j)
    before = (dict(j["dec"]), j["up"]["file"], sorted(p.name for p in c.inbox.iterdir()))
    audit = c.job_dir(j) / "audit.jsonl"
    audit.unlink()
    audit.mkdir()  # the log cannot be opened for appending
    with pytest.raises(RU.ReuploadError, match="Nothing was changed"):
        RU.replace(c, j, "Investors_v3.xlsx", edited(j, tmp_path, "Investors_v3.xlsx", insert_top=True), WHO)
    assert (dict(j["dec"]), j["up"]["file"], sorted(p.name for p in c.inbox.iterdir())) == before
