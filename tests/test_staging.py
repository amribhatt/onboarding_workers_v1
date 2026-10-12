from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils.cell import column_index_from_string

from worker import workbook as wb
from worker.staging import SRC, Stage
from worker.templates import TEMPLATES

FIXTURES = Path(__file__).parent / "fixtures"
MESSY = FIXTURES / "messy"
MARKERS = {"n/a", "-", "--", "tbd"}


def _key(layout: str) -> dict:
    return json.loads((MESSY / f"investors_{layout}.answer.json").read_text(encoding="utf-8"))


def _sheet_name(path: Path) -> str:
    return wb.list_sheets(path)["sheets"][0]["name"]


def _col_of(path: Path, text: str) -> str:
    hit = wb.find(path, _sheet_name(path), text)["matches"][0]["addr"]
    match = re.match(r"[A-Z]+", hit)
    assert match
    return match.group(0)


def test_lift_fund_roster_matches_openpyxl(tmp_path: Path) -> None:
    path = FIXTURES / "set7" / "Fund_Roster.xlsx"
    st = Stage(tmp_path)
    sid = st.lift(
        path,
        _sheet_name(path),
        {"GP_ID": "A", "Fund_ID": "B"},
        header_row=1,
        first_row=2,
        template="Commitments",
        constants={"Investor_ID": "X"},
    )
    ws = openpyxl.load_workbook(path).active
    expected = [(r[0].value, r[1].value) for r in ws.iter_rows(min_row=2) if r[0].value is not None]
    rows = st._rows[sid]
    assert [(r["GP_ID"], r["Fund_ID"]) for r in rows] == expected
    assert all(r["Investor_ID"] == "X" for r in rows)
    assert rows[0][SRC].endswith("!2")
    assert (tmp_path / "staged" / f"{sid}.csv").exists()
    s = st.summary(sid)
    assert s["rows"] == len(expected)
    assert len(s["samples"]) <= 3


@pytest.mark.parametrize(
    ("layout", "id_header", "name_header", "header_row"),
    [("vertical_messy", "Inv ID", "LP Name", 7), ("wide_noise", "Inv ID", "Legal Name", 3)],
)
def test_lift_messy_row_count_and_blanks(
    layout: str, id_header: str, name_header: str, header_row: int, tmp_path: Path
) -> None:
    path = MESSY / f"investors_{layout}.xlsx"
    key = _key(layout)
    id_col, name_col = _col_of(path, id_header), _col_of(path, name_header)
    sheet = _sheet_name(path)
    st = Stage(tmp_path)
    sid = st.lift(
        path,
        sheet,
        {"Investor_ID": id_col, "Investor_Name": name_col},
        header_row=header_row,
        first_row=header_row + 1,
        template="Investors",
    )
    s = st.summary(sid)
    assert s["rows"] == key["unique_investor_count"] + len(key["duplicates"])
    rows = st._rows[sid]
    assert {t["Investor_ID"] for t in key["truth"]} <= {r["Investor_ID"] for r in rows}
    # blank markers counted against an independent read of the lifted cells
    ws = openpyxl.load_workbook(path)[sheet]
    last = int(s["last_src"].split("!")[1])
    cells = [
        ws.cell(r, column_index_from_string(c)).value
        for c in (id_col, name_col)
        for r in range(header_row + 1, last + 1)
    ]
    expected = sum(1 for v in cells if isinstance(v, str) and v.strip().lower() in MARKERS)
    assert s["blank_markers"] >= expected  # rows skipped as header repeats may add none
    assert not any(str(r["Investor_ID"]).strip().lower() in MARKERS for r in rows)
    dups = st.duplicates(sid, ["Investor_ID"])
    assert dups and all(len(g["rows"]) > 1 for g in dups)


def test_vertical_stops_at_total_row(tmp_path: Path) -> None:
    path = MESSY / "investors_vertical_messy.xlsx"
    st = Stage(tmp_path)
    sid = st.lift(path, _sheet_name(path), {"Investor_ID": "A"}, 7, 8, template="Investors")
    s = st.summary(sid)
    assert s["stopped_at"] and s["reason"]
    assert s["last_src"].endswith("!50")


def test_wide_noise_skips_repeated_header(tmp_path: Path) -> None:
    path = MESSY / "investors_wide_noise.xlsx"
    st = Stage(tmp_path)
    cols = {"Investor_ID": "P", "Investor_Name": "AD"}
    sid = st.lift(path, _sheet_name(path), cols, 3, 4, template="Investors")
    assert st.summary(sid)["repeated_headers"] == 1


def test_flatten_compact_pivot(tmp_path: Path) -> None:
    path = MESSY / "investors_compact_pivot.xlsx"
    key = _key("compact_pivot")
    st = Stage(tmp_path)
    sid = st.flatten_compact(
        path,
        _sheet_name(path),
        {0: "Investor_ID", 1: "Investor_Name", 2: "Investor_Status"},
        first_row=3,
        template="Investors",
    )
    rows = st._rows[sid]
    # the pivot merges case variants, so compare IDs case-insensitively (values stay as-is)
    assert {t["Investor_ID"] for t in key["truth"]} <= {r["Investor_ID"].upper() for r in rows}
    assert {t["Investor_Name"] for t in key["truth"]} <= {r["Investor_Name"] for r in rows}
    assert not any(str(r["Investor_ID"]).lower() in ("row labels", "grand total") for r in rows)
    assert all(r["Investor_Name"] != "(blank)" for r in rows)


def test_lift_blocks_stacked(tmp_path: Path) -> None:
    path = MESSY / "investors_stacked_blocks.xlsx"
    key = _key("stacked_blocks")
    blocks = [(b["header_row"], b["first_data_row"], b["total_row"]) for b in key["layout_details"]["blocks"]]
    st = Stage(tmp_path)
    sid = st.lift_blocks(
        path, _sheet_name(path), {"Investor_ID": "A", "Investor_Name": "B"}, blocks, template="Investors"
    )
    rows = st._rows[sid]
    ws = openpyxl.load_workbook(path).active
    expected = sum(1 for _, f, t in blocks for r in range(f, t + 1) if ws.cell(r, 1).value not in (None, ""))
    assert len(rows) == expected
    ids = {r["Investor_ID"] for r in rows}
    aliases = {k.split("@")[0] for k in key["layout_details"]["id_aliases"]}
    assert aliases <= ids  # copied as-is, never translated


def test_reconcile_and_write_upload(tmp_path: Path) -> None:
    path = MESSY / "investors_vertical_messy.xlsx"
    st = Stage(tmp_path)
    sheet = _sheet_name(path)
    cols = {"Investor_ID": "A", "Investor_Name": "B"}
    a = st.lift(path, sheet, cols, 7, 8, template="Investors", constants={"GP_ID": "GP-001"})
    b = st.lift(path, sheet, {"Investor_ID": "A"}, 7, 8, 30, template="Investors")
    rec = st.reconcile(a, b, "Investor_ID", "Investor_ID")
    assert rec["both"] == rec["right_distinct"]
    assert rec["only_right"] == 0
    assert rec["only_left"] == rec["left_distinct"] - rec["both"]

    out = tmp_path / "out"
    dups = st.duplicates(a, ["Investor_ID"])
    p1 = st.write_upload(a, "Investors", out, drop_duplicates=True, duplicate_key=["Investor_ID"])
    p2 = st.write_upload(a, "Investors", out, drop_duplicates=False, duplicate_key=[])
    p3 = st.write_upload(a, "Investors", out, drop_duplicates=False, duplicate_key=[])
    assert [p.name for p in (p1, p2, p3)] == [
        "Investors_Upload.xlsx",
        "Investors_Upload_v2.xlsx",
        "Investors_Upload_v3.xlsx",
    ]
    rows1 = list(openpyxl.load_workbook(p1).active.iter_rows(values_only=True))
    rows2 = list(openpyxl.load_workbook(p2).active.iter_rows(values_only=True))
    assert list(rows1[0]) == list(TEMPLATES["Investors"].fields)
    assert len(rows2) - len(rows1) == sum(len(g["rows"]) - 1 for g in dups)
    assert rows1[1][rows1[0].index("GP_ID")] == "GP-001"
    assert rows1[1][rows1[0].index("Source_Investor_Key")] is None


def test_unknown_field_refused(tmp_path: Path) -> None:
    path = MESSY / "investors_vertical_messy.xlsx"
    with pytest.raises(ValueError):
        Stage(tmp_path).lift(path, _sheet_name(path), {"Nope": "A"}, 7, 8, template="Investors")
