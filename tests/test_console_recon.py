"""The eleven GL reconciliation checks: one engine, eleven specs, the upload side from real lines, the client side
from the analyst's own report. Fictitious data only (sponsor-a style names, Ledgerline, Fundbook)."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from worker import console_demo as D
from worker import console_learning as L
from worker import recon as R
from worker import sample as S
from worker import templates as T
from worker.accounts import Coa
from worker.audit import AuditLog
from worker.console import Console
from worker.web import create_app

COA = {k: dict(v) for k, v in Coa.load(T.standards_dir() / "Golden_COA.csv").accounts.items()}
GOLD = {a: g.split(" ", 1)[0] for a, g, _f in D.ACCTS if g}
GOLD["6150"] = "42020"  # the guide's unrealised-gain account, for the fixture only


def rows_from(lines: list[S.Line]) -> list[R.Row]:
    """The engine's rows for a list of sample lines, as the console builds them."""
    out = []
    for ln in lines:
        a = GOLD[ln.acct]
        out.append(R.Row(a, COA[a]["FS_Grouping"], COA[a]["FS_Line_Item"], round(ln.dr - ln.cr, 2),
                         R.parse_date(ln.date), ln.inv, ln.tt, ln.invst, R.parse_date(ln.pr), ln.shares,
                         R._yes(ln.agn)))
    return out


def L_(j: str, n: int, acct: str, tt: str, dr: float, cr: float, d: str, **kw: Any) -> S.Line:
    return S.Line(j, n, acct, "", tt, dr, cr, "x", date=d, **kw)


# A small known fund GL. Year 2025; one 2024 line. All journals balance.
FUND = [
    L_("J1", 1, "1000", "Capital call", 1_000_000, 0, "2024-06-30"),
    L_("J1", 2, "3000", "Capital call", 0, 1_000_000, "2024-06-30"),
    L_("J2", 1, "1500", "Purchase", 400_000, 0, "2025-02-15", invst="Beta Energy Series B", shares=1000.0),
    L_("J2", 2, "1000", "Purchase", 0, 400_000, "2025-02-15"),
    L_("J3", 1, "1000", "Disposal", 45_000, 0, "2025-03-31"),
    L_("J3", 2, "8050", "Disposal", 0, 45_000, "2025-03-31"),
    L_("J4", 1, "6100", "Expense", 12_500, 0, "2025-04-10"),
    L_("J4", 2, "1000", "Expense", 0, 12_500, "2025-04-10"),
    L_("J5", 1, "3100", "Distribution", 250_000, 0, "2025-05-01"),
    L_("J5", 2, "1000", "Distribution", 0, 250_000, "2025-05-01"),
    L_("J6", 1, "1600", "Valuation", 80_000, 0, "2025-06-30", invst="Beta Energy Series B"),
    L_("J6", 2, "6150", "Valuation", 0, 80_000, "2025-06-30"),
]
# A small known LP GL (investor rows; offsets are added as the blank-investor rows).
LP = [
    L_("K1", 1, "3000", "Capital call", 0, 600_000, "2024-06-30", inv="LP-A", pr="2024-06-30", agn="Yes"),
    L_("K1", 2, "3000", "Capital call", 0, 400_000, "2024-06-30", inv="LP-B", pr="2024-07-15", agn="Yes"),
    L_("K2", 1, "3100", "Distribution", 150_000, 0, "2025-05-01", inv="LP-A", pr="2025-05-01", agn="Yes"),
    L_("K2", 2, "3100", "Recallable distribution", 10_000, 0, "2025-05-01", inv="LP-A", pr="2025-05-01", agn="Yes"),
    L_("K2", 3, "3100", "Distribution", 100_000, 0, "2025-05-01", inv="LP-B", pr="2025-05-01", agn="Yes"),
]
NAMES = {"LP-A": "Investor Alpha", "LP-B": "Investor Beta"}
ALIASES = {"lp-a": "LP-A", "investor alpha": "LP-A", "lp-b": "LP-B", "investor beta": "LP-B"}


def _resolve(line: str, tt: str) -> str:
    t = line.strip().split(" ")[0]
    return t if t in COA else ""


def run(spec_id: str, lines: list[S.Line], report: list[dict[str, Any]] | None = None, **kw: Any) -> dict[str, Any]:
    rows = rows_from(lines)
    if lines is LP:
        rows += [R.Row(r.acct, r.fsg, r.fsl, -r.amt, r.date, "", r.tt, offset=True) for r in rows]
    ctx = R.make_ctx(rows, NAMES)
    return R.run_check(R.SPEC[spec_id], rows, ctx, COA, report, _resolve, ALIASES, **kw)


def rep(line: str = "", amount: float | None = None, **kw: Any) -> dict[str, Any]:
    return {"line": line, "investor": "", "date": "", "tt": "", "measure": "", "period": "", "amount": amount, **kw}


def byl(res: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {r["label"]: r for r in res["rows"]}


def figure(res: dict[str, Any], label: str) -> float:
    return byl(res)[label]["a"]


# ---- the specs themselves (table-driven) ---------------------------------------------------------------------
GUIDE_MATRIX = {"tb": ["DE", "EL", "LP"], "bs": ["DE", "EL", "LP"], "soi": ["DE", "EL", "LP"],
                "is": ["DE", "EL", "LP"], "cap": ["DE", "EL", "LP"], "cf": ["DE", "EL", "LP"],
                "pcap": ["DLP", "LP"], "irr": ["DLP", "LP"], "unf": ["DLP", "LP"], "wf": ["DLP", "LP"],
                "evl": ["DE"]}
GUIDE_PERIOD = {"tb": "ITD", "bs": "ITD", "soi": "ITD", "is": "YTD", "cap": "ITD", "cf": "YTD",
                "pcap": "ITD + YTD", "irr": "ITD", "unf": "ITD", "wf": "ITD", "evl": "ITD"}


def test_eleven_specs_in_the_guides_order() -> None:
    assert [s["id"] for s in R.SPECS] == list(GUIDE_MATRIX)
    assert len(R.SPECS) == 11


@pytest.mark.parametrize("spec", R.SPECS, ids=lambda s: s["id"])
def test_every_spec_is_consistent_with_the_guide(spec: dict[str, Any]) -> None:
    assert spec["scope"] == GUIDE_MATRIX[spec["id"]]
    assert spec["per"] == GUIDE_PERIOD[spec["id"]]
    assert spec["sel"] and spec["rule"] and len(spec["steps"]) >= 4
    known = set(spec["measures"]) | set(spec.get("derived", {}))
    assert set(spec["order"]) <= known
    for d in spec.get("derived", {}).values():
        assert {p for p, _ in d["of"]} <= known
    assert spec["tol"] == ("tiered" if spec["id"] in ("tb", "bs") else "flat")
    if not spec.get("pair"):
        assert spec["layout"] and spec["layout_text"]
        assert R.layout_bytes(spec["id"])


def test_thresholds_tiered_and_flat() -> None:
    assert R.classify(100.0, 99.0, "money", "tiered") == (1.0, "PASS")
    assert R.classify(100.0, 96.0, "money", "tiered") == (4.0, "WARN")
    assert R.classify(100.0, 94.99, "money", "tiered") == (5.01, "ERROR")
    assert R.classify(100.0, 98.0, "money", "flat")[1] == "ERROR"
    assert R.classify(100.0, 99.0, "money", "flat")[1] == "PASS"
    assert R.classify(1000, 1001, "int", "flat")[1] == "ERROR"
    assert R.classify(11.82, 11.82, "pct", "flat")[1] == "PASS"
    assert R.classify(11.82, 11.5, "pct", "flat")[1] == "ERROR"


# ---- the upload side of each check, on the known GL --------------------------------------------------------
def test_tb_upload_side() -> None:
    res = run("tb", FUND)
    assert res["state"] == "needs-report"
    got = {r["label"].split(" ")[0]: r["a"] for r in res["rows"] if r["mid"] == "bal"}
    assert got == {"11010": 382_500.0, "15010": 400_000.0, "15910": 80_000.0, "31010": -1_000_000.0,
                   "32010": 250_000.0, "42010": -45_000.0, "42020": -80_000.0, "51020": 12_500.0}
    assert all(r["b"] is None for r in res["rows"] if r["mid"] == "bal")  # no client figure without a report
    net = byl(res)["All accounts net to zero"]
    assert net["a"] == 0.0 and net["status"] == "PASS"


def test_bs_upload_side_ties_assets_to_liabilities_and_capital() -> None:
    res = run("bs", FUND)
    r = byl(res)
    assert r["Cash and equivalents"]["a"] == 382_500.0
    assert r["Investments at cost"]["a"] == 400_000.0
    assert r["Capital contributions"]["a"] == 1_000_000.0
    assert r["Capital distributions"]["a"] == -250_000.0
    assert r["Total assets"]["a"] == 862_500.0 and r["Total capital"]["a"] == 862_500.0
    assert r["Assets less liabilities and capital"]["status"] == "PASS"


def test_soi_upload_side_cost_unrealised_fmv_shares() -> None:
    r = byl(run("soi", FUND))
    assert r["Beta Energy Series B · cost"]["a"] == 400_000.0
    assert r["Beta Energy Series B · unrealised"]["a"] == 80_000.0
    assert r["Beta Energy Series B · fair market value"]["a"] == 480_000.0
    assert r["Beta Energy Series B · shares"]["a"] == 1000.0


def test_soi_without_shares_column_says_so() -> None:
    lines = [S.Line(x.j, x.ln, x.acct, x.an, x.tt, x.dr, x.cr, x.desc, date=x.date, invst=x.invst) for x in FUND]
    res = run("soi", lines)
    assert not any(r["label"].endswith("shares") for r in res["rows"])
    assert any("NO_OF_SHARES is not in this GL file" in n for n in res["notes"])


def test_is_upload_side_is_year_to_date_only() -> None:
    r = byl(run("is", FUND))
    assert r["Realised gains and losses"]["a"] == 45_000.0
    assert r["Unrealised gains and losses"]["a"] == 80_000.0
    assert r["Professional services"]["a"] == 12_500.0
    assert r["Total revenue"]["a"] == 125_000.0 and r["Total expenses"]["a"] == 12_500.0
    assert r["Net income"]["a"] == 112_500.0


def test_cap_upload_side_rolls_forward() -> None:
    r = byl(run("cap", FUND))
    assert r["Beginning capital"]["a"] == 1_000_000.0  # the 2024 contribution only
    assert "Contributions (YTD)" not in r  # none this year: a line with no movement is not invented
    assert r["Distributions (YTD)"]["a"] == 250_000.0
    assert r["Net income (YTD)"]["a"] == 112_500.0
    assert r["Ending capital"]["a"] == 862_500.0
    tie = r["Ending capital less total capital on the balance sheet"]
    assert tie["a"] == 0.0 and tie["status"] == "PASS"


def test_cf_upload_side_sections_and_cash_tie() -> None:
    r = byl(run("cf", FUND))
    assert r["Net cash from operating"]["a"] == -12_500.0
    assert r["Net cash from investing"]["a"] == -355_000.0
    assert r["Net cash from financing"]["a"] == -250_000.0
    assert r["Net change in cash"]["a"] == -617_500.0
    assert r["Opening cash"]["a"] == 1_000_000.0 and r["Closing cash"]["a"] == 382_500.0
    assert r["Net change in cash less the cash accounts' movement"]["status"] == "PASS"
    assert r["Opening cash + net change less closing cash"]["status"] == "PASS"


def test_pcap_upload_side_per_investor_itd_and_ytd() -> None:
    res = run("pcap", LP)
    r = byl(res)
    assert r["Investor Alpha · contributions"]["a"] == 600_000.0
    assert r["Investor Alpha · distributions"]["a"] == 160_000.0
    assert r["Investor Alpha · ending balance"]["a"] == 440_000.0
    assert r["Investor Beta · ending balance"]["a"] == 300_000.0
    assert r["Investor Alpha · opening balance (YTD)"]["a"] == 600_000.0
    assert r["Investor Alpha · ending balance (YTD)"]["a"] == 440_000.0
    assert any("fund run" in n for n in res["notes"])  # no fund-level total without the fund run


def test_pcap_sum_of_investors_against_the_fund_run() -> None:
    fund = rows_from([L_("F1", 1, "3000", "Capital call", 0, 1_000_000, "2024-06-30"),
                      L_("F2", 1, "3100", "Distribution", 260_000, 0, "2025-05-01")])
    res = run("pcap", LP, fund_rows=fund)
    t = byl(res)["Sum of investors less the fund-level capital total"]
    assert t["a"] == 0.0 and t["status"] == "PASS"  # 440,000 + 300,000 = 740,000 = 1,000,000 - 260,000


def test_irr_upload_side_uses_pr_date_and_terminal_value() -> None:
    r = byl(run("irr", LP))
    assert r["Investor Alpha · 2024-06-30"]["a"] == -600_000.0
    assert r["Investor Alpha · 2025-05-01"]["a"] == 160_000.0
    assert r["Investor Beta · 2024-07-15"]["a"] == -400_000.0  # PR_DATE, not the GL date 2024-06-30
    assert r["Investor Alpha · ending capital (terminal value)"]["a"] == 440_000.0
    flagged = [x for x in r.values() if x["status"] == "REVIEW"]
    assert flagged and flagged[0]["a"] == 1  # one row's PR_DATE differs from its GL date


def test_unf_upload_side_called_and_recallable() -> None:
    res = run("unf", LP)
    r = byl(res)
    assert r["Investor Alpha · contributions called"]["a"] == 600_000.0
    assert r["Investor Alpha · recallable distributions"]["a"] == 10_000.0
    assert r["Investor Beta · contributions called"]["a"] == 400_000.0
    assert not any(x["label"].endswith("unfunded commitment") for x in res["rows"])  # no commitment, no figure


def test_wf_upload_side_groups_by_investor_pr_date_and_trans_type() -> None:
    r = byl(run("wf", LP))
    assert r["Investor Alpha · 2024-06-30 · Capital call"]["a"] == -600_000.0
    assert r["Investor Alpha · 2025-05-01 · Recallable distribution"]["a"] == 10_000.0
    assert r["Investor Beta · 2024-07-15 · Capital call"]["a"] == -400_000.0


def test_evl_compares_fund_and_lp_runs_and_the_offsets() -> None:
    fund = rows_from([L_("F1", 1, "3000", "Capital call", 0, 1_000_000, "2024-06-30"),
                      L_("F2", 1, "3100", "Distribution", 260_000, 0, "2025-05-01")])
    lp = rows_from(LP)
    lp += [R.Row(r.acct, r.fsg, r.fsl, -r.amt, r.date, "", r.tt, offset=True) for r in lp]
    res = R.run_check(R.SPEC["evl"], lp, R.make_ctx(lp, NAMES), COA, None, _resolve, ALIASES, fund, lp)
    assert res["state"] == "run"
    r = byl(res)
    assert r["31010 Partner contributions"]["status"] == "PASS"
    assert r["32010 Partner distributions"]["a"] == 260_000.0 and r["32010 Partner distributions"]["b"] == 260_000.0
    assert r["32010 Partner distributions · LP + offset"]["status"] == "PASS"
    short = rows_from([L_("F1", 1, "3000", "Capital call", 0, 999_000, "2024-06-30")])
    bad = R.run_check(R.SPEC["evl"], lp, R.make_ctx(lp, NAMES), COA, None, _resolve, ALIASES, short, lp)
    assert byl(bad)["31010 Partner contributions"]["status"] == "ERROR"


# ---- the client side: diffs, tiers, review, gaps --------------------------------------------------------------
def test_tb_client_report_pass_warn_error_and_review() -> None:
    report = [rep("11010 Operating bank account", 382_500.40),  # 0.40: PASS
              rep("15010", 400_003.0),  # 3: WARN (tiered)
              rep("32010", 245_000.0),  # 5,000: ERROR
              rep("31010", -1_000_000.0), rep("15910", 80_000.0), rep("42010", -45_000.0),
              rep("42020", -80_000.0), rep("51020", 12_500.0),
              rep("9999 Unmapped client account", 10.0)]  # no Ledgerline account: REVIEW, never matched by guess
    res = run("tb", FUND, report)
    assert res["state"] == "run"
    r = {x["label"].split(" ")[0]: x for x in res["rows"]}
    assert r["11010"]["status"] == "PASS" and r["11010"]["diff"] == -0.4
    assert r["15010"]["status"] == "WARN" and r["15010"]["diff"] == -3.0
    assert r["32010"]["status"] == "ERROR" and r["32010"]["diff"] == 5_000.0
    assert r["31010"]["status"] == "PASS"
    rv = [x for x in res["rows"] if x["status"] == "REVIEW"]
    assert len(rv) == 1 and "no Ledgerline account" in rv[0]["note"] and rv[0]["a"] is None


def test_missing_on_either_side_is_a_difference_not_a_pass() -> None:
    report = [rep("11010", 382_500.0), rep("31010", -1_000_000.0), rep("11020", 50.0)]
    res = run("tb", FUND, report)
    r = {x["label"].split(" ")[0]: x for x in res["rows"] if x["mid"] == "bal"}
    assert r["15010"]["status"] == "ERROR" and r["15010"]["b"] == 0.0  # in the upload, not on the client's report
    assert r["11020"]["status"] == "ERROR" and r["11020"]["a"] == 0.0


def test_flat_threshold_on_the_income_statement() -> None:
    report = [rep("Realised gains and losses", 45_000.5), rep("Unrealised gains and losses", 80_002.0),
              rep("Professional services", 12_500.0), rep("Net income", 112_500.0)]
    r = byl(run("is", FUND, report))
    assert r["Realised gains and losses"]["status"] == "PASS"
    assert r["Unrealised gains and losses"]["status"] == "ERROR"
    assert r["Net income"]["status"] == "PASS"


def test_bs_renamed_label_is_review_never_auto_pass() -> None:
    report = [rep("Cash", 382_500.0), rep("Total assets", 862_500.0)]
    res = run("bs", FUND, report)
    r = byl(res)
    assert r["Cash and equivalents"]["status"] == "REVIEW" and "rename" in r["Cash and equivalents"]["note"]
    assert r["Total assets"]["status"] == "PASS"
    assert any(x["status"] == "REVIEW" and x["label"] == "Cash" for x in res["rows"])


def test_soi_shares_are_whole_numbers() -> None:
    report = [{**rep("Beta Energy Series B", 400_000.0), "measure": "Cost"},
              {**rep("Beta Energy Series B", 480_000.0), "measure": "Fair value"},
              {**rep("Beta Energy Series B", 1001.0), "measure": "Shares"}]
    r = byl(run("soi", FUND, report))
    assert r["Beta Energy Series B · cost"]["status"] == "PASS"
    assert r["Beta Energy Series B · fair market value"]["status"] == "PASS"
    assert r["Beta Energy Series B · shares"]["status"] == "ERROR"


def test_irr_date_mismatch_is_an_error_and_xirr_is_checked() -> None:
    def cf(inv: str, d: str, a: float, line: str = "") -> dict[str, Any]:
        return {**rep(line, a), "investor": inv, "date": d}

    report = [cf("Investor Alpha", "2024-06-30", -600_000.0), cf("Investor Alpha", "2025-05-01", 160_000.0),
              cf("Investor Beta", "2024-07-20", -400_000.0),  # the client's date differs by five days
              cf("Investor Beta", "2025-05-01", 100_000.0),
              cf("Investor Alpha", "2025-12-31", 0.0, "Ending capital")]
    res = run("irr", LP, report)
    r = byl(res)
    assert r["Investor Alpha · 2024-06-30"]["status"] == "PASS"
    assert r["Investor Beta · 2024-07-15"]["status"] == "ERROR"  # ours has no flow on the client's date
    assert r["Investor Beta · 2024-07-20"]["status"] == "ERROR"  # the client's has none on ours
    assert any(x["label"].endswith("IRR (%)") and x["b"] is None for x in res["rows"])  # nothing shown to compare


def test_xirr_matches_a_hand_figure() -> None:
    # -100 then +121 two years later: 10% a year
    assert R.xirr([(date(2023, 1, 1), -100.0), (date(2025, 1, 1), 121.0)]) == pytest.approx(10.0, abs=0.05)
    assert R.xirr([(date(2023, 1, 1), 100.0), (date(2025, 1, 1), 121.0)]) is None


def test_unf_unfunded_from_the_clients_commitment() -> None:
    def x(inv: str, line: str, a: float) -> dict[str, Any]:
        return {**rep(line, a), "investor": inv}

    report = [x("Investor Alpha", "Commitment", 1_000_000.0), x("Investor Alpha", "Unfunded", 410_000.0),
              x("Investor Beta", "Commitment", 500_000.0), x("Investor Beta", "Unfunded", 99_000.0)]
    res = run("unf", LP, report)
    r = byl(res)
    # the fixture carries AGNST_COMMIT, so contributions called are exact: 1,000,000 - 600,000 + 10,000
    assert r["Investor Alpha · unfunded commitment"]["a"] == 410_000.0
    assert r["Investor Alpha · unfunded commitment"]["status"] == "PASS"
    assert r["Investor Beta · unfunded commitment"]["a"] == 100_000.0
    assert r["Investor Beta · unfunded commitment"]["status"] == "ERROR"


def test_unf_without_agnst_commit_is_review_not_pass() -> None:
    lines = [S.Line(x.j, x.ln, x.acct, x.an, x.tt, x.dr, x.cr, x.desc, inv=x.inv, date=x.date, pr=x.pr) for x in LP]
    res = run("unf", lines)
    assert any("AGNST_COMMIT" in n for n in res["notes"])
    called = byl(res)["Investor Alpha · contributions called"]
    assert called["status"] == "REVIEW"


def test_wf_both_orientations_and_pr_date_review() -> None:
    long = [{**rep("", -600_000.0), "investor": "LP-A", "date": "2024-06-30", "tt": "Capital call"},
            {**rep("", -400_000.0), "investor": "LP-B", "date": "2024-07-15", "tt": "Capital call"},
            {**rep("", 150_000.0), "investor": "LP-A", "date": "2025-05-01", "tt": "Distribution"},
            {**rep("", 10_000.0), "investor": "LP-A", "date": "2025-05-01", "tt": "Recallable distribution"},
            {**rep("", 100_000.0), "investor": "LP-B", "date": "2025-05-01", "tt": "Distribution"}]
    res = run("wf", LP, long)
    assert not [x for x in res["rows"] if x["status"] == "ERROR"]
    assert any(x["status"] == "REVIEW" for x in res["rows"])  # LP-B: PR_DATE 2024-07-15 vs GL date 2024-06-30
    transposed = [{**rep("", -600_000.0), "investor": "LP-A", "date": "2024-06-30"},
                  {**rep("", -400_000.0), "investor": "LP-B", "date": "2024-07-15"},
                  {**rep("", 160_000.0), "investor": "LP-A", "date": "2025-05-01"},
                  {**rep("", 100_000.0), "investor": "LP-B", "date": "2025-05-01"}]
    res2 = run("wf", LP, transposed)
    assert not [x for x in res2["rows"] if x["status"] == "ERROR"]


def test_pcap_client_diff_per_investor_bucket() -> None:
    def x(inv: str, line: str, a: float) -> dict[str, Any]:
        return {**rep(line, a), "investor": inv}

    report = [x("Investor Alpha", "Contributions", 600_000.0), x("Investor Alpha", "Distributions", 150_000.0),
              x("Investor Alpha", "Ending balance", 450_000.0), x("Investor Beta", "Ending balance", 300_000.0)]
    r = byl(run("pcap", LP, report))
    assert r["Investor Alpha · contributions"]["status"] == "PASS"
    assert r["Investor Alpha · distributions"]["status"] == "ERROR" and r["Investor Alpha · distributions"][
        "diff"] == 10_000.0
    assert r["Investor Beta · ending balance"]["status"] == "PASS"


# ---- gaps, applicability, honest empty states --------------------------------------------------------------
def test_pr_date_gap_is_reported_not_filled_from_gl_date() -> None:
    lines = [S.Line(x.j, x.ln, x.acct, x.an, x.tt, x.dr, x.cr, x.desc, inv=x.inv, date=x.date) for x in LP]
    for spec in ("irr", "wf"):
        res = run(spec, lines)
        assert res["state"] == "gap" and res["rows"] == []
        assert "PR_DATE is not in this GL file" in res["notes"][0]


def test_year_to_date_gap_without_gl_dates() -> None:
    lines = [S.Line(x.j, x.ln, x.acct, x.an, x.tt, x.dr, x.cr, x.desc) for x in FUND]
    for spec in ("is", "cap", "cf"):
        res = run(spec, lines)
        assert res["state"] == "gap" and "GL Date" in res["notes"][0]


def test_lp_checks_need_investors() -> None:
    res = run("pcap", FUND)
    assert res["state"] == "gap" and "not an LP-level upload" in res["notes"][0]


def test_incomplete_upload_never_shows_a_false_error_on_its_own_tests() -> None:
    half = FUND[:-1]  # the last journal lost one side
    res = run("tb", half, incomplete=True)
    net = byl(res)["All accounts net to zero"]
    assert net["a"] != 0 and net["status"] == "REVIEW"
    res2 = run("tb", half)
    assert byl(res2)["All accounts net to zero"]["status"] == "ERROR"


def test_layout_is_prefilled_for_fixed_line_checks() -> None:
    for cid in ("cap", "cf"):
        wb = openpyxl.load_workbook(__import__("io").BytesIO(R.layout_bytes(cid)))
        labels = [r[0] for r in wb["Layout"].iter_rows(min_row=2, values_only=True)]
        assert "Ending capital" in labels or "Closing cash" in labels
        assert wb["Notes"]["A2"].value


def test_read_report_csv_and_transposed(tmp_path: Path) -> None:
    p = tmp_path / "r.csv"
    p.write_text("Line,Amount\n11010 Operating,1234.50\n(parens),(10)\n", encoding="utf-8")
    rows = R.read_report(p)
    assert [r["amount"] for r in rows] == [1234.5, -10.0]
    t = tmp_path / "t.csv"
    t.write_text("Investor,2024-06-30,2025-05-01\nLP-A,-600000,160000\n", encoding="utf-8")
    rows = R.read_report(t)
    assert [(r["investor"], r["date"], r["amount"]) for r in rows] == [("LP-A", "2024-06-30", -600000.0),
                                                                        ("LP-A", "2025-05-01", 160000.0)]
    bad = tmp_path / "b.csv"
    bad.write_text("Nothing,Here\nx,y\n", encoding="utf-8")
    with pytest.raises(R.ReportError):
        R.read_report(bad)


# ---- the console: view, upload, audit ---------------------------------------------------------------------
@pytest.fixture
def c(tmp_path: Path) -> Console:
    return Console(tmp_path / "inbox")


def gl_fund(c: Console, tmp: Path) -> dict[str, Any]:
    c.act("prov|paid")
    c.act("start|Sponsor B|GL Fund")
    c.act("ask_set|fund|B Fund I")
    c.act("ask_sample")
    c.act("ask_go")
    c.act("go")
    j = c.job()
    assert j is not None
    return j


def write_gl(path: Path, lines: list[tuple[Any, ...]], extra: list[str] | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(S.GLSRC) + (extra or []))
    for ln in lines:
        ws.append(list(ln))
    wb.save(path)


def custom_job(c: Console, tmp: Path) -> dict[str, Any]:
    j = gl_fund(c, tmp)
    cols = len(S.GLSRC)
    def row(d: str, jn: str, n: int, acct: str, tt: str, dr: float, cr: float, inv: str = "", invst: str = "",
            shares: float | None = None) -> list[Any]:
        r: list[Any] = [None] * (cols + 1)
        for name, v in (("GL Date", d), ("Journal", jn), ("Line", n), ("GL Account", acct), ("Transaction Type", tt),
                        ("Debit", dr), ("Credit", cr), ("Description", "x"), ("Investor", inv),
                        ("Investment", invst)):
            r[S.GLSRC.index(name)] = v
        r[cols] = shares
        return r

    beta = "Beta Energy Series B"
    lines = [row("2024-06-30", "J1", 1, "1000", "Capital call", 1_000_000, 0),
             row("2024-06-30", "J1", 2, "3000", "Capital call", 0, 1_000_000),
             row("2025-02-15", "J2", 1, "1500", "Purchase", 400_000, 0, invst=beta, shares=1000),
             row("2025-02-15", "J2", 2, "1000", "Purchase", 0, 400_000),
             row("2025-03-31", "J3", 1, "1000", "Disposal", 45_000, 0),
             row("2025-03-31", "J3", 2, "8050", "Disposal", 0, 45_000, invst=beta),
             row("2025-04-10", "J4", 1, "6100", "Expense", 12_500, 0),
             row("2025-04-10", "J4", 2, "1000", "Expense", 0, 12_500),
             row("2025-05-01", "J5", 1, "3100", "Distribution", 250_000, 0),
             row("2025-05-01", "J5", 2, "1000", "Distribution", 0, 250_000)]
    p = tmp / "custom_gl.xlsx"
    write_gl(p, lines, ["No of Shares"])
    j["up"]["path"] = str(p)
    c.save()
    return j


def check(v: dict[str, Any], cid: str) -> dict[str, Any]:
    return next(k for k in v["checks"] if k["id"] == cid)


def test_view_computes_upload_side_from_the_jobs_own_lines(c: Console, tmp_path: Path) -> None:
    j = custom_job(c, tmp_path)
    v = L.learning_view(c, tmp_path, "recon", j["id"])
    tb = check(v, "tb")
    assert tb["state"] == "needs-report"
    bal = {r["label"].split(" ")[0]: r for r in tb["rows"]}
    assert bal["11010"]["a"] == 382_500.0 and bal["11010"]["b"] is None  # no client figure without a report
    assert tb["worst"] == "PASS" and "Needs the client's trial balance" in tb["note"]
    soi = check(v, "soi")
    assert any(r["label"].endswith("shares") and r["a"] == 1000.0 for r in soi["rows"])
    assert check(v, "is")["rows"][-1]["a"] == 32_500.0
    assert check(v, "pcap")["state"] == "gap"
    assert check(v, "evl")["state"] == "waiting"
    assert check(v, "tb")["layout"]["columns"] == ["Line", "Trans Type", "Amount"]


def test_report_upload_compares_and_is_audited(c: Console, tmp_path: Path) -> None:
    j = custom_job(c, tmp_path)
    data = (b"Line,Trans Type,Amount\n11010 Operating bank account,,382500.40\n15010,,400003\n"
            b"32010,,245000\n31010,,-1000000\n15910,,0\n42010,,-125000\n51020,,12500\n")
    out = L.learning_report(c, j["id"], "tb", "client_tb.csv", data, "reviewer-a")
    assert out["ok"] and out["rows"] == 7
    tb = check(out["view"], "tb")
    assert tb["state"] == "run" and tb["worst"] == "ERROR"
    by = {r["label"].split(" ")[0]: r for r in tb["rows"]}
    assert by["11010"]["status"] == "PASS" and by["15010"]["status"] == "WARN" and by["32010"]["status"] == "ERROR"
    assert by["15910"]["status"] == "PASS"  # the client's zero and the upload's absence agree
    assert tb["report"]["file"] == "client_tb.csv" and tb["report"]["by"] == "reviewer-a"
    log = (c.job_dir(j) / "audit.jsonl").read_text(encoding="utf-8")
    assert "recon.report_uploaded" in log and "reviewer-a" in log
    assert AuditLog(c.job_dir(j) / "audit.jsonl").first_break()[1] is None


def test_report_upload_refuses_unreadable_files_and_keeps_nothing(c: Console, tmp_path: Path) -> None:
    j = custom_job(c, tmp_path)
    with pytest.raises(L.LearningError):
        L.learning_report(c, j["id"], "tb", "x.csv", b"Nothing,Here\nx,y\n")
    with pytest.raises(L.LearningError):
        L.learning_report(c, j["id"], "tb", "x.pdf", b"%PDF")
    with pytest.raises(L.LearningError):
        L.learning_report(c, j["id"], "evl", "x.csv", b"Line,Amount\na,1\n")
    assert check(L.learning_view(c, tmp_path, "recon", j["id"]), "tb")["state"] == "needs-report"


def test_audit_failure_stops_the_report(c: Console, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    j = custom_job(c, tmp_path)

    def boom(*a: Any, **k: Any) -> None:
        raise OSError("disk")

    monkeypatch.setattr(AuditLog, "record", boom)
    with pytest.raises(OSError):
        L.learning_report(c, j["id"], "tb", "x.csv", b"Line,Amount\n11010,1\n")
    assert R.load_report(c.job_dir(j), "tb") is None


def test_combined_break_note_only_when_the_data_shows_it(c: Console, tmp_path: Path) -> None:
    j = custom_job(c, tmp_path)
    v = L.learning_view(c, tmp_path, "recon", j["id"])
    assert v["cross"]["DE"] == []  # no reports, no break
    # a 10,000 distribution the client shows as recallable: the same break in the TB and in the capital schedule
    tb = "Line,Amount\n11010,382500\n15010,400000\n31010,-1000000\n32010,240000\n42010,-45000\n51020,12500\n"
    cap = ("Line,Amount\nBeginning capital,1000000\nDistributions (YTD),240000\nNet income (YTD),32500\n"
           "Ending capital,792500\n")
    L.learning_report(c, j["id"], "tb", "tb.csv", tb.encode())
    out = L.learning_report(c, j["id"], "cap", "cap.csv", cap.encode())
    cross = out["view"]["cross"]["DE"]
    assert cross and cross[0]["amount"] == 10_000.0
    assert {"Trial Balance", "Schedule of Capital"} <= set(cross[0]["checks"])


def test_fund_type_applicability_and_tiles(c: Console, tmp_path: Path) -> None:
    j = custom_job(c, tmp_path)
    v = L.learning_view(c, tmp_path, "recon", j["id"])
    ids = {ft: [k["id"] for k in v["checks"] if ft in k["scope"] and k["id"] in R.SPEC] for ft in
           ("DE", "DLP", "EL", "LP")}
    assert ids["DE"] == ["tb", "bs", "soi", "is", "cap", "cf", "evl"]
    assert ids["DLP"] == ["pcap", "irr", "unf", "wf"]
    assert ids["EL"] == ["tb", "bs", "soi", "is", "cap", "cf"]
    assert ids["LP"] == ["tb", "bs", "soi", "is", "cap", "cf", "pcap", "irr", "unf", "wf"]
    de = v["by_ft"]["DE"]
    on = [k for k in v["checks"] if "DE" in k["scope"]]
    assert de["applies"] == len(on)
    assert de["pass"] + de["warn"] + de["error"] + de["review"] + de["needs_report"] + de["gap"] + de["waiting"] \
        == len(on)


def test_endpoints_upload_and_layout(tmp_path: Path) -> None:
    client = TestClient(create_app(str(tmp_path / "st")))
    for a in ("prov|paid", "start|Sponsor B|GL Fund", "ask_set|fund|B Fund I", "ask_sample", "ask_go", "go"):
        client.post("/api/console/act", json={"action": a})
    jid = client.get("/api/console/state").json()["job"]["id"]
    lay = client.get("/api/console/learning/layout/tb")
    assert lay.status_code == 200 and lay.content[:2] == b"PK"
    assert client.get("/api/console/learning/layout/evl").status_code == 404
    r = client.put(f"/api/console/learning/report/{jid}/tb", params={"name": "client_tb.csv"},
                   content=b"Line,Amount\n11010,1\n")
    assert r.status_code == 200 and r.json()["rows"] == 1
    bad = client.put(f"/api/console/learning/report/{jid}/tb", params={"name": "client_tb.csv"}, content=b"zzz")
    assert bad.status_code == 409
    assert client.put("/api/console/learning/report/j404/tb", content=b"x").status_code == 404


def test_reconcile_is_a_show_working_tab_for_gl_and_numbers_only_the_guide_checks() -> None:
    """Found in the browser: Reconcile was a fold under the job and the guide's checks were numbered 3-13."""
    from pathlib import Path

    static = Path(__file__).resolve().parents[1] / "src" / "worker" / "static"
    working = (static / "working.js").read_text(encoding="utf-8")
    learning = (static / "learning.js").read_text(encoding="utf-8")
    console = (static / "console.js").read_text(encoding="utf-8")
    assert '["xwalk", "Crosswalks"], ["recon", "Reconcile"]' in working  # the tab sits after Crosswalks, GL only
    assert 'guide.indexOf(c) + 1 : "·"' in learning  # the eleven are 1-11, the job's own tie-outs are dotted
    assert '/^GL/.test(j.type) ? [] : ["recon"]' in console  # no second copy as a fold under a GL job
