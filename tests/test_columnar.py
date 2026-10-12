"""Columnar ingest (`worker.columnar`): a sheet reads back cell for cell as the readers it replaced gave it."""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from worker import columnar


@pytest.fixture(autouse=True)
def _cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("WORKER_CACHE_DIR", str(tmp_path / "cache"))


def _csv_reference(path: Path) -> list[list[str]]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [list(r) for r in csv.reader(fh)]


def _rows(path: Path, sheet: str | None = None) -> list[list[Any]]:
    return [cells for _n, cells in columnar.rows(path, sheet)]


CSV_CASES = {
    "plain": "a,b,c\n1,2,3\n007,x y,\n",
    "bom": "﻿a,b\n1,2\n",
    "quoted_newline": 'a,b\n"line1\nline2",2\n3,4\n',
    "ragged": "a,b,c\n1,2\n1,2,3,4\n",
    "blank_lines": "a,b\n\n1,2\n\n",
    "no_trailing_newline": "a,b\n1,2",
    "quotes": 'a,b\n"he said ""hi""",2\n',
    "header_only": "a,b,c\n",
    "empty": "",
    "crlf": "a,b\r\n1,2\r\n3,4\r\n",
    "bare_cr": "a,b\r1,2\r3,4\r",
    "spaces": " a , b \n 1 , 2 \n",
    "stray_quote": 'a,b\nab"c,2\n',
    "unicode": "name,city\ncafé,Zürich\n日本,x\n",
    "empty_cells": "a,b,c\n,,\n1,,3\n",
}


@pytest.mark.parametrize("name", sorted(CSV_CASES))
def test_csv_reads_as_the_csv_reader_does(tmp_path: Path, name: str) -> None:
    p = tmp_path / f"{name}.csv"
    p.write_bytes(CSV_CASES[name].encode("utf-8"))
    got = _rows(p)
    want = _csv_reference(p)
    assert len(got) == len(want)
    for g, w in zip(got, want, strict=True):  # a blank line is a blank row (cells all empty), not a short one
        assert g == w or (not w and g == [""] * len(g))


def test_random_csv_reads_as_the_csv_reader_does(tmp_path: Path) -> None:
    import random

    rnd = random.Random(7)
    alphabet = ["a", "B", " ", ",", '"', "\n", "\r\n", "é", "7", "", "x y", "'", ";", "\t"]
    p = tmp_path / "random.csv"
    with open(p, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
        for _ in range(3000):
            w.writerow(["".join(rnd.choice(alphabet) for _ in range(rnd.randint(0, 6))) for _ in range(5)])
    want = _csv_reference(p)
    got = _rows(p)
    assert len(got) == len(want)
    assert all(g == w or (not w and g == [""] * len(g)) for g, w in zip(got, want, strict=True))


def test_csv_row_numbers_count_records(tmp_path: Path) -> None:
    p = tmp_path / "n.csv"
    p.write_text('a,b\n"x\ny",1\n2,3\n', encoding="utf-8")
    assert [n for n, _ in columnar.rows(p)] == [1, 2, 3]


def test_xlsx_keeps_cell_types_and_blank_cells(tmp_path: Path) -> None:
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Id", "Amount", "When", "Flag", "Mixed"])
    ws.append(["A1", 10, dt.datetime(2025, 3, 31), True, "text"])
    ws.append(["A2", 2.5, dt.datetime(2025, 4, 1, 13, 30), False, 7])
    ws.append([None, None, None, None, dt.datetime(2025, 1, 1)])
    ws["C9"] = "far"
    wb.create_sheet("Other")["B2"] = 3
    wb.save(p)
    book = openpyxl.load_workbook(p, read_only=True, data_only=True)
    want = [list(r) for r in book["Data"].iter_rows(values_only=True)]
    book.close()
    got = _rows(p, "Data")
    assert got == want
    assert [[type(v) for v in r] for r in got] == [[type(v) for v in r] for r in want]
    assert columnar.sheet_names(p) == ["Data", "Other"]
    assert _rows(p, "Other")[1] == [None, 3]


def test_xlsx_offset_data_is_padded_like_openpyxl(tmp_path: Path) -> None:
    p = tmp_path / "off.xlsx"
    wb = openpyxl.Workbook()
    wb.active["C3"] = "x"
    wb.active["D4"] = 5
    wb.save(p)
    book = openpyxl.load_workbook(p, read_only=True, data_only=True)
    want = [list(r) for r in book.active.iter_rows(values_only=True)]
    book.close()
    assert _rows(p) == want


def test_unknown_sheet_is_a_key_error(tmp_path: Path) -> None:
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(p)
    with pytest.raises(KeyError):
        _rows(p, "Nope")


def test_many_parts_read_in_order_and_the_cache_is_reused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(columnar, "BATCH", 7)
    p = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["k", "v"])
    for i in range(100):
        wb.active.append([f"K{i:03d}", i])
    wb.save(p)
    parts = columnar.parquet_parts(p)
    assert len(parts) >= 2
    stamp = [x.stat().st_mtime_ns for x in parts]
    assert [r[0] for r in _rows(p)][1:4] == ["K000", "K001", "K002"]
    assert [r[1] for r in _rows(p)][1:4] == [0, 1, 2]
    assert [n for n, _ in columnar.rows(p)][-1] == 101
    assert [x.stat().st_mtime_ns for x in columnar.parquet_parts(p)] == stamp
    wb = openpyxl.Workbook()
    wb.active.append(["k", "v"])
    wb.active.append(["new", 1])
    wb.save(p)  # other content, other cache entry
    assert _rows(p) == [["k", "v"], ["new", 1]]


def test_whole_numbers_are_ints_and_big_floats_stay_floats(tmp_path: Path) -> None:
    p = tmp_path / "n.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1.0, 2.5, 1e20, -3])
    wb.save(p)
    row = _rows(p)[0]
    assert row == [1, 2.5, 1e20, -3]
    assert [type(v) for v in row] == [int, float, float, int]


def test_a_big_sheet_cut_into_pieces_reads_as_the_whole_sheet(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Pieces parsed in separate processes join into the rows, types and blank rows the one-pass reader gives,
    including empty rows, a column only some rows use and mixed kinds."""
    p = tmp_path / "wide.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Id", "Amt", "When", "Note"])
    for i in range(1, 400):
        if i % 97 == 0:
            ws.append([])  # an empty row
            continue
        ws.append([f"K{i:04d}", i * 1.0 if i % 3 else i + 0.25, dt.datetime(2025, 1, 1 + i % 28),
                   "late" if i == 390 else None])
    ws["F2"] = "far"
    wb.save(p)
    whole = _rows(p, "Data")  # far below the cutting size: one pass
    monkeypatch.setenv("WORKER_CACHE_DIR", str(tmp_path / "cache2"))
    monkeypatch.setattr(columnar, "PARALLEL_BYTES", 1_000)
    monkeypatch.setattr(columnar, "BATCH", 50)
    assert columnar._row_cuts(p, "Data") is not None
    cut = _rows(p, "Data")
    assert cut == whole and len(whole) == 400
    assert [[type(v) for v in r] for r in cut] == [[type(v) for v in r] for r in whole]
    assert [n for n, _ in columnar.rows(p, "Data")] == list(range(1, 401))
