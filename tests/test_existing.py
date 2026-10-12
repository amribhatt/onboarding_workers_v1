from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pytest
from langgraph.checkpoint.memory import InMemorySaver

from id_decision import decide_agent_ids
from scripted_model import ScriptedModel, approve, call, say
from worker.agent import build_agent, pending_gate, start
from worker.audit import AuditLog
from worker.context import RunContext
from worker.registry import Registry
from worker.staging import Stage

KEYS = {"Investor_ID": "A", "Investor_Name": "B"}


def _xlsx(path: Path, rows: list[list[str]], header: tuple[str, str] = ("ID", "Name")) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    ws = book.active
    ws.append(list(header))
    for r in rows:
        ws.append(r)
    book.save(path)
    return path


def _registry(tmp_path: Path) -> Registry:
    reg = Registry(tmp_path, "sponsor-a")
    reg.add(
        "Investors",
        [
            {"key": "I1", "name": "Alpha Capital"},
            {"key": "I2", "name": "Beta Partners"},
            {"key": "I3", "name": "Gamma Holdings"},
        ],
        source="app batch",
        batch="job-1",
        fund="Fund A",
    )
    reg.add("Investors", [{"key": "I9", "name": "Omega Trust"}], source="outside export", batch="ledgerline.xlsx")
    return reg


def _ctx(tmp_path: Path, inbox: Path, reg: Registry | None) -> RunContext:
    out = tmp_path / "out"
    decide_agent_ids(tmp_path / "ids", "Investors")  # approval waits for the sponsor's ID decision
    return RunContext(
        sponsor="sponsor-a",
        gp_id="GP-001",
        inbox_roots=[inbox],
        out_dir=out,
        audit=AuditLog(out / "audit.jsonl"),
        stage=Stage(out),
        registry=reg,
        id_root=tmp_path / "ids",
    )


def _lift_call(name: str) -> dict[str, Any]:
    return call(
        "lift", path=name, sheet="Sheet", template="Investors", columns=KEYS, header_row=1, first_row=2, why="copy"
    )


SOURCE_ROWS = [
    ["I1", "Alpha Capital"],  # row 2: same key, same name -> already onboarded
    ["I2", "Beta Partners LLC"],  # row 3: same key, other name -> ERR
    ["I7", "Gamma Holdings"],  # row 4: other key, same name -> ERR
    ["I4", "Delta Fund"],  # row 5: new
    ["I5", "Echo Fund"],  # row 6: new
]


def _stage_with_registry(tmp_path: Path) -> tuple[Stage, str, Registry]:
    reg = _registry(tmp_path)
    st = Stage(tmp_path / "out")
    sid = st.lift(_xlsx(tmp_path / "src.xlsx", SOURCE_ROWS), "Sheet", KEYS, 1, 2, template="Investors")
    return st, sid, reg


def _match(st: Stage, sid: str, reg: Registry) -> str:
    return st.match_existing(sid, reg.entries("Investors"), ["Investor_ID"], "Investor_Name")


def test_three_match_cases(tmp_path: Path) -> None:
    st, sid, reg = _stage_with_registry(tmp_path)
    out = _match(st, sid, reg)
    meta = st.summary(out)
    assert meta["existing"]["already_onboarded"] == 1
    assert [x["src"] for x in st.left_out(out)] == ["Sheet!2"]
    assert meta["existing"]["matched_by_source"] == {"app batch": 1}
    kinds = {c["src"]: c["kind"] for c in meta["conflicts"]}
    assert kinds == {"Sheet!3": "same_key_other_name", "Sheet!4": "same_name_other_key"}
    # in = out + excluded; the two ERR rows stay in the stage until resolved
    assert meta["existing"]["staged"] == 5 == meta["rows"] + meta["existing"]["already_onboarded"]


def test_write_upload_refused_while_conflicts_are_open(tmp_path: Path) -> None:
    st, sid, reg = _stage_with_registry(tmp_path)
    out = _match(st, sid, reg)
    with pytest.raises(ValueError, match="still conflict"):
        st.write_upload(out, "Investors", tmp_path / "up")


def test_resolution_actions_and_output_has_no_excluded_row(tmp_path: Path) -> None:
    st, sid, reg = _stage_with_registry(tmp_path)
    out = _match(st, sid, reg)
    with pytest.raises(ValueError, match="reason is required"):
        st.resolve_conflicts(out, {"Sheet!3": {"action": "exclude_duplicate", "reason": " "}}, reg.entries("Investors"))
    done = st.resolve_conflicts(
        out,
        {
            "Sheet!3": {"action": "exclude_incorrect", "reason": "wrong name in the source"},
            "Sheet!4": {"action": "correct", "reason": "ID typo; real ID is I3", "values": {"Investor_ID": "I3"}},
        },
        reg.entries("Investors"),
    )
    assert st.summary(done)["conflict_count"] == 0
    # the corrected row is now an exact match of an existing record and leaves the file
    path = st.write_upload(done, "Investors", tmp_path / "up")
    book = openpyxl.load_workbook(path)
    ids = [r[1] for r in book.active.iter_rows(min_row=2, values_only=True)]
    assert ids == ["I4", "I5"]
    actions = {r["src"]: r["action"] for r in st.summary(done)["resolutions"]}
    assert actions == {"Sheet!3": "exclude_incorrect", "Sheet!4": "correct"}


def test_correct_that_still_conflicts_stays_open(tmp_path: Path) -> None:
    st, sid, reg = _stage_with_registry(tmp_path)
    out = _match(st, sid, reg)
    done = st.resolve_conflicts(
        out,
        {"Sheet!3": {"action": "correct", "reason": "try", "values": {"Investor_Name": "Beta Partner"}}},
        reg.entries("Investors"),
    )
    # same key as I2 with a different name is still a conflict after the edit
    assert {c["src"] for c in st.summary(done)["conflicts"]} == {"Sheet!3", "Sheet!4"}


def test_no_id_in_source_uses_name_and_flags_probable_merges(tmp_path: Path) -> None:
    reg = _registry(tmp_path)
    rows = [["", "alpha capital"], ["", "Beta Partner"], ["", "Zeta Group"]]
    st = Stage(tmp_path / "out")
    sid = st.lift(_xlsx(tmp_path / "s.xlsx", rows), "Sheet", {"Investor_Name": "B"}, 1, 2, template="Investors")
    out = _match(st, sid, reg)
    info = st.summary(out)["existing"]
    assert info["key_kinds"] == {"source id": 0, "name": 3}
    assert [x["src"] for x in st.left_out(out)] == ["Sheet!2"]  # same normalised name, no ID: the name is the key
    assert [(c["src"], c["kind"]) for c in st.summary(out)["conflicts"]] == [("Sheet!3", "probable_merge")]


def test_registry_provenance_says_where_the_count_comes_from(tmp_path: Path) -> None:
    reg = _registry(tmp_path)
    assert reg.provenance("Investors") == {"app batch · Fund A": 3, "outside export": 1}


def test_registry_is_per_sponsor_and_never_rewrites(tmp_path: Path) -> None:
    a, b = Registry(tmp_path, "sponsor-a"), Registry(tmp_path, "sponsor-b")
    rec = [{"key": "I1", "name": "Alpha"}]
    assert a.add("Investors", rec, source="app batch", batch="j1") == 1
    assert a.add("Investors", rec, source="app batch", batch="j2") == 0
    assert b.entries("Investors") == []
    assert a.counts() == {"Investors": 1}
    with pytest.raises(ValueError):
        Registry(tmp_path, "../x")
    with pytest.raises(ValueError):
        a.add("Investors", rec, source="guess", batch="j3")


def _agent(ctx: RunContext, script: list[Any]) -> Any:
    return build_agent(ctx, ScriptedModel(script=script), InMemorySaver())


def test_tools_audit_each_excluded_row_and_registry_grows_after_approval(tmp_path: Path) -> None:
    reg = _registry(tmp_path)
    inbox = tmp_path / "inbox"
    _xlsx(inbox / "f.xlsx", [["I1", "Alpha Capital"], ["I4", "Delta Fund"], ["I5", "Echo Fund"]])
    ctx = _ctx(tmp_path, inbox, reg)
    script = [
        say("lift", _lift_call("f.xlsx")),
        say("match", call("match_existing", stage_id="s1", why="dedup against existing records")),
        say("write", call("write_upload", stage_id="s2", template="Investors", why="deliver")),
        say("done"),
    ]
    agent = _agent(ctx, script)
    result = start(agent, ctx, "build", "t1")
    assert [g["name"] for g in pending_gate(result)] == ["write_upload"]
    assert len(reg.entries("Investors")) == 4  # nothing added before approval
    excluded = [e for e in ctx.audit.read() if e["action"] == "row.excluded"]
    assert len(excluded) == 1
    assert excluded[0]["inputs"]["row"] == "Sheet!2" and excluded[0]["inputs"]["key"] == "I1"
    assert excluded[0]["outputs"]["matched"]["source"] == "app batch"
    approve(agent, ctx, "t1")
    assert len(reg.entries("Investors")) == 6
    assert any(e["action"] == "registry.added" for e in ctx.audit.read())
    assert ctx.audit.verify()[0]


def test_import_of_an_export_is_gated(tmp_path: Path) -> None:
    reg = Registry(tmp_path, "sponsor-a")
    inbox = tmp_path / "inbox"
    _xlsx(inbox / "ledgerline_export.xlsx", [["X1", "One Co"], ["X2", "Two Co"]])
    ctx = _ctx(tmp_path, inbox, reg)
    script = [
        say(
            "import",
            call(
                "import_existing_records",
                path="ledgerline_export.xlsx",
                sheet="Sheet",
                template="Investors",
                columns=KEYS,
                header_row=1,
                first_row=2,
                why="records loaded outside the app",
            ),
        ),
        say("done"),
    ]
    agent = _agent(ctx, script)
    result = start(agent, ctx, "import", "t1")
    assert [g["name"] for g in pending_gate(result)] == ["import_existing_records"]
    assert reg.entries("Investors") == []
    approve(agent, ctx, "t1", "yes")
    assert [e["key"] for e in reg.entries("Investors")] == ["X1", "X2"]
    assert {e["source"] for e in reg.entries("Investors")} == {"outside export"}


def test_new_ids_are_matched_through_the_source_key_in_a_later_batch(tmp_path: Path) -> None:
    reg = Registry(tmp_path, "sponsor-a")
    reg.add(
        "Investors",
        [{"key": "GINV-1", "name": "Alpha Capital", "source_key": "I1"}],
        source="app batch",
        batch="job-1",
        fund="Fund A",
    )
    st = Stage(tmp_path / "out")
    sid = st.lift(
        _xlsx(tmp_path / "s.xlsx", [["I1", "Alpha Capital"], ["I8", "New Co"]]), "Sheet", KEYS, 1, 2, template="Investors"
    )
    out = _match(st, sid, reg)
    assert [x["key"] for x in st.left_out(out)] == ["I1"]


def test_resolve_conflicts_tool_audits_each_decision_and_replays(tmp_path: Path) -> None:
    from worker.memory import SponsorMemory, replay

    reg = _registry(tmp_path)
    inbox = tmp_path / "inbox"
    _xlsx(inbox / "f.xlsx", SOURCE_ROWS)
    ctx = _ctx(tmp_path, inbox, reg)
    ctx.memory = SponsorMemory(tmp_path / "memory", "sponsor-a")
    decisions = {
        "Sheet!3": {"action": "exclude_incorrect", "reason": "wrong name in the source"},
        "Sheet!4": {"action": "correct", "reason": "typo in ID", "values": {"Investor_ID": "I3"}},
    }
    script = [
        say("lift", _lift_call("f.xlsx")),
        say("match", call("match_existing", stage_id="s1", why="dedup")),
        say("resolve", call("resolve_conflicts", stage_id="s2", decisions=decisions, why="human decided")),
        say("write", call("write_upload", stage_id="s3", template="Investors", why="deliver")),
        say("done"),
    ]
    agent = _agent(ctx, script)
    start(agent, ctx, "build", "t1")
    resolved = [e for e in ctx.audit.read() if e["action"] == "conflict.resolved"]
    assert {e["inputs"]["row"]: e["rationale"] for e in resolved} == {
        "Sheet!3": "wrong name in the source",
        "Sheet!4": "typo in ID",
    }
    approve(agent, ctx, "t1")
    assert len(ctx.memory.cases()) == 1
    assert [r["changed"] for r in replay(ctx)] == [False]
