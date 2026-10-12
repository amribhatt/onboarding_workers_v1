"""Live run finding (docs/live-model-run.md): the CLI read the model settings from `.env` into a dict, but the roles
built inside a tool read `os.environ`. Live, `propose_layout_program` then fell back to the default provider without
a key ("Missing credentials") and the Coworker asked the person for OpenAI credentials; `WORKER_EXPLAINER=1` in
`.env` was ignored the same way. The job's settings now travel on `RunContext.env`. No model is called here."""

from __future__ import annotations

import argparse
import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from scripted_model import ScriptedModel
from worker import agent as A
from worker import console_reshape as CR
from worker import roles
from worker.audit import AuditLog
from worker.cli import _build_session
from worker.context import RunContext
from worker.staging import Stage
from worker.templates import TEMPLATES
from worker.tools import build_tools

JOB_ENV = {"WORKER_MODEL_PROVIDER": "openrouter", "WORKER_MODEL": "vendor/model-a:free",
           "OPENROUTER_API_KEY": "test-not-a-key", "WORKER_EXPLAINER": "1"}
PIVOT = Path(__file__).parent / "fixtures" / "messy" / "investors_pivot.xlsx"


@pytest.fixture(autouse=True)
def _clean_environ(monkeypatch: pytest.MonkeyPatch) -> None:
    for name in JOB_ENV:
        monkeypatch.delenv(name, raising=False)  # only the job's settings say which provider to use


def _ctx(tmp_path: Path, env: dict[str, str] | None) -> RunContext:
    out = tmp_path / "out"
    out.mkdir(parents=True, exist_ok=True)
    return RunContext(sponsor="sponsor-a", gp_id="GP-001", inbox_roots=[PIVOT.parent], out_dir=out,
                      audit=AuditLog(out / "audit.jsonl"), stage=Stage(out, dict(TEMPLATES)), env=env)


def test_the_cli_session_carries_the_env_file_settings(tmp_path: Path) -> None:
    args = argparse.Namespace(sponsor="sponsor-a", state_dir=str(tmp_path / "state"), model=None)
    s = _build_session(args, JOB_ENV, ScriptedModel(script=[]), None, input, io.StringIO(),
                       PIVOT.parent, "GP-001", tmp_path / "out", A.new_thread_id())
    assert s.ctx.env is not None and s.ctx.env["WORKER_MODEL_PROVIDER"] == "openrouter"


def test_the_layout_drafter_in_the_tool_runs_on_the_job_settings(tmp_path: Path,
                                                                 monkeypatch: pytest.MonkeyPatch) -> None:
    seen: dict[str, Any] = {}

    def fake_draft(runner: Any, flex: Any, **kw: Any) -> roles.RoleOutcome:
        seen["env"] = runner.env
        return roles.RoleOutcome("needs_human")

    monkeypatch.setattr(roles, "draft_layout", fake_draft)
    ctx = _ctx(tmp_path, dict(JOB_ENV))
    tools = {t.name: t for t in build_tools(ctx, sandbox=object())}
    sheet = load_workbook(PIVOT, read_only=True).sheetnames[0]
    tools["propose_layout_program"].invoke({"path": PIVOT.name, "sheet": sheet, "template": "Investors",
                                            "why": "an unknown layout"})
    assert seen["env"] == JOB_ENV


def test_the_explainer_switch_is_read_from_the_job_settings(tmp_path: Path) -> None:
    runner = A._explainer_for(_ctx(tmp_path, dict(JOB_ENV)), None, JOB_ENV)
    assert isinstance(runner, roles.RoleRunner) and runner.env == JOB_ENV
    assert A._explainer_for(_ctx(tmp_path, {}), None, {}) is None


def test_build_agent_passes_the_job_settings_not_os_environ(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    got: list[Any] = []
    real = A._explainer_for
    monkeypatch.setattr(A, "_explainer_for", lambda ctx, ex, env: got.append(env) or real(ctx, ex, env))
    from langgraph.checkpoint.memory import InMemorySaver

    A.build_agent(_ctx(tmp_path, dict(JOB_ENV)), ScriptedModel(script=[]), InMemorySaver())
    assert got == [JOB_ENV]


def test_the_console_reshape_runner_reads_the_env_file_and_is_capped(tmp_path: Path,
                                                                     monkeypatch: pytest.MonkeyPatch) -> None:
    env_file = tmp_path / "job.env"
    env_file.write_text("WORKER_MODEL_PROVIDER=openrouter\nWORKER_MODEL_PLAN=vendor/model-b:free\n"
                        "WORKER_BUDGET_CALLS=7\n", encoding="utf-8")
    monkeypatch.setenv("WORKER_ENV_FILE", str(env_file))
    runner = CR.role_runner(None, AuditLog(tmp_path / "a.jsonl"))
    assert runner.env is not None and runner.env["WORKER_MODEL_PLAN"] == "vendor/model-b:free"
    assert runner.budget is not None and runner.budget.max_calls == 7
