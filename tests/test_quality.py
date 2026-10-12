from __future__ import annotations

import csv
import os
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from worker.quality import (
    MARK_COLUMN,
    Supplied,
    groups_from_spec,
    just_these_rows,
    load_rules,
    marked_copy,
    scan,
)
from worker.templates import TEMPLATES

RULES = load_rules()


def _xlsx(path: Path, header: list[str], rows: list[list[Any]], sheet: str = "Sheet1") -> Path:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    book.save(path)
    return path


def _cause(report: dict[str, Any], rule: str, field: str | None = None) -> dict[str, Any]:
    found = [c for c in report["causes"] if c["rule"] == rule and (field is None or c["field"] == field)]
    assert len(found) == 1, [c["id"] for c in report["causes"]]
    return found[0]


def _investors(tmp_path: Path) -> Path:
    rows = [[f"HS-99{i:02d}", f"Investor {i}", "LP", "Fund A"] for i in range(1, 15)]  # rows 2..15
    rows[13][1] = None  # row 15: blank name
    return _xlsx(tmp_path / "Sponsor_A_Common_Investors.xlsx", ["Row Labels", "Investor", "Type", "Entity"], rows,
                 "Investor")


def test_blank_investor_name_is_one_cause_with_its_recipe_and_row(tmp_path: Path) -> None:
    path = _investors(tmp_path)
    report = scan(path, RULES, columns={"Investor_ID": "Row Labels", "Investor_Name": "Investor"},
                  record_type="Investors", sheet="Investor",
                  groups=groups_from_spec(TEMPLATES["Investors"], ["Investor_ID", "Investor_Name"]))
    cause = _cause(report, "required_blank", "Investor_Name")
    assert cause["count"] == 1 and cause["row_ranges"] == [[15, 15]] and cause["severity"] == "ERR"
    assert "column B \u201cInvestor\u201d = (Blanks)" in cause["locator"]["filter"]
    assert cause["locator"]["where"] == (
        "Sponsor_A_Common_Investors.xlsx \u00b7 sheet Investor, row 15 \u00b7 Investor!A15:D15"
    )
    assert cause["sample_rows"][0]["values"]["Investor_ID"] == "HS-9914"


def test_a_value_supplied_by_the_person_clears_the_cause_and_is_never_inferred(tmp_path: Path) -> None:
    path = _investors(tmp_path)
    kw: dict[str, Any] = {
        "columns": {"Investor_ID": "Row Labels", "Investor_Name": "Investor"}, "record_type": "Investors",
        "sheet": "Investor", "groups": groups_from_spec(TEMPLATES["Investors"], ["Investor_ID", "Investor_Name"]),
    }
    supplied = Supplied()
    supplied.set_row(15, "Investor_Name", "Alder Co LLC", who="alex")
    again = scan(path, RULES, supplied=supplied.rows, **kw)
    assert not [c for c in again["causes"] if c["rule"] == "required_blank"]
    assert supplied.log[0]["inferred"] is False and supplied.log[0]["reason"] == "supplied by the person"
    with pytest.raises(ValueError, match="cannot be blank"):
        supplied.set_row(15, "Investor_Name", "  ", who="alex")


def test_commitment_date_blank_suggests_the_batch_value(tmp_path: Path) -> None:
    rows = [[f"LP-{i}", "2025-06-30", 100 + i] for i in range(1, 9)]  # rows 2..9
    rows[5][1] = None  # row 7
    path = _xlsx(tmp_path / "commitments.xlsx", ["LP ID", "Closing Date", "Total Commitment"], rows)
    report = scan(
        path, RULES, record_type="Commitments",
        columns={"Investor_ID": "LP ID", "Commitment_Date": "Closing Date", "Commitment_Amount": "Total Commitment"},
        groups=groups_from_spec(TEMPLATES["Commitments"], ["Investor_ID", "Commitment_Date", "Commitment_Amount"]),
    )
    cause = _cause(report, "required_blank", "Commitment_Date")
    assert cause["row_ranges"] == [[7, 7]]
    assert "column B \u201cClosing Date\u201d = (Blanks)" in cause["locator"]["filter"]
    assert cause["suggestion"] == "Use 2025-06-30, the closing date of the other rows in this batch"


def test_affiliate_blank_name_follows_the_same_rule_with_no_special_case(tmp_path: Path) -> None:
    rows = [[f"AFF-{i}", f"Affiliate {i}"] for i in range(1, 7)]
    rows[3][1] = ""  # row 5
    path = _xlsx(tmp_path / "aff.xlsx", ["Affiliate ID", "Affiliate"], rows)
    report = scan(path, RULES, record_type="Affiliates",
                  columns={"Affiliate_ID": "Affiliate ID", "Affiliate_Name": "Affiliate"},
                  groups=groups_from_spec(TEMPLATES["Affiliates"], ["Affiliate_ID", "Affiliate_Name"]))
    cause = _cause(report, "required_blank", "Affiliate_Name")
    assert cause["row_ranges"] == [[5, 5]] and cause["message"] == "Affiliate_Name is blank"


def test_required_field_with_no_column_is_reported_once_and_covers_every_row(tmp_path: Path) -> None:
    rows = [[f"LP-{i}", 100] for i in range(1, 9)]
    path = _xlsx(tmp_path / "c.xlsx", ["LP ID", "Amount"], rows)
    report = scan(path, RULES, record_type="Commitments",
                  columns={"Investor_ID": "LP ID", "Commitment_Amount": "Amount", "Currency": "Currency"},
                  groups={"required": ["Investor_ID", "Commitment_Amount", "Currency"]})
    cause = _cause(report, "required_blank", "Currency")
    assert cause["message"] == "This file has no \u201cCurrency\u201d column"
    assert cause["count"] == 8 and cause["severity"] == "ERR" and cause["no_column"]
    assert report["rows"] == 8


def test_a_generated_id_that_is_blank_is_not_an_error(tmp_path: Path) -> None:
    path = _xlsx(tmp_path / "i.xlsx", ["Deal ID", "Deal Name"], [[None, "Beta Energy"], ["D-2", "Gamma"]])
    report = scan(path, RULES, record_type="Investments", columns={"Investment_ID": "Deal ID",
                                                                  "Investment_Name": "Deal Name"},
                  groups={"required": ["Investment_ID", "Investment_Name"]}, generated=["Investment_ID"])
    assert report["causes"] == []


def test_validity_uniqueness_and_consistency_families_each_fire(tmp_path: Path) -> None:
    rows = [
        ["I1" + "x" * 25, "2025-13-45", "abc", "Alpha"],  # row 2: long ID, bad date, bad amount
        ["I2", "2025-06-30", "100", "Beta"],
        ["I2", "2025-06-30", "200", "Beta"],  # row 4: duplicate key
        ["I3", "2025-06-30", "300", "Ghost"],  # row 5: investor not known
    ]
    path = _xlsx(tmp_path / "mix.xlsx", ["ID", "Date", "Amount", "Investor"], rows)
    groups = {"ids": ["Commitment_ID"], "dates": ["Commitment_Date"], "amounts": ["Commitment_Amount"],
              "key": ["Commitment_ID"]}
    report = scan(
        path, RULES, record_type="Commitments", groups=groups,
        columns={"Commitment_ID": "ID", "Commitment_Date": "Date", "Commitment_Amount": "Amount",
                 "Investor_ID": "Investor"},
        references={"investors": {"alpha", "beta"}},
    )
    assert _cause(report, "id_too_long")["row_ranges"] == [[2, 2]]
    assert _cause(report, "bad_date")["row_ranges"] == [[2, 2]]
    assert _cause(report, "bad_amount")["row_ranges"] == [[2, 2]]
    assert _cause(report, "duplicate_key")["row_ranges"] == [[3, 4]]
    assert _cause(report, "unknown_investor")["row_ranges"] == [[5, 5]]
    assert report["families"] == {"validity": 3, "uniqueness": 1, "consistency": 1}


def test_unbalanced_journals_are_one_balance_cause(tmp_path: Path) -> None:
    rows = [["J1", 100, 0], ["J1", 0, 100], ["J2", 50, 0], ["J2", 0, 49], ["J3", 10, 0], ["J3", 0, 10]]
    path = _xlsx(tmp_path / "gl.xlsx", ["Journal", "Debit", "Credit"], rows)
    report = scan(path, RULES, record_type="GL", columns={"Journal": "Journal", "Debit": "Debit", "Credit": "Credit"})
    cause = _cause(report, "journal_unbalanced")
    assert cause["row_ranges"] == [[4, 5]] and cause["count"] == 2 and cause["copy_values"] == ["J2"]


# ---- scenarios G: finding thousands of rows --------------------------------------------------------------
def _gl_csv(path: Path, plan: dict[int, list[tuple[int, int]]], total: int, deal: dict[str, str] | None = None) -> Path:
    """A GL CSV of `total` rows where each account owns the row blocks in `plan` (rows are 1-based file rows)."""
    owner: dict[int, str] = {}
    for acct, blocks in plan.items():
        for a, b in blocks:
            for n in range(a, b + 1):
                owner[n] = str(acct)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Journal", "Line", "GL Account", "Investment"])
        for n in range(2, total + 1):
            acct = owner.get(n, "1000")
            w.writerow([f"J{n // 2}", n, acct, (deal or {}).get(acct, "Deal X")])
    return path


def test_one_value_many_rows_shows_the_filter_and_ranges_never_a_list_of_rows(tmp_path: Path) -> None:
    blocks = [(1204, 7203), (256000, 258999), (909000, 910213)]
    path = _gl_csv(tmp_path / "gl.csv", {6150: blocks}, 910_300)
    report = scan(path, RULES, record_type="GL", columns={"Account": "GL Account"},
                  references={"coa_sources": {"1000"}})
    cause = _cause(report, "account_not_in_chart", "Account")
    assert cause["count"] == 10_214
    assert cause["locator"]["filter"] == (
        "In Excel: Data \u2192 Filter, column C \u201cGL Account\u201d = 6150 \u2192 10,214 lines"
    )
    assert cause["row_ranges"] == [[a, b] for a, b in blocks]
    assert cause["ranges_text"] == "1204\u20137203, 256000\u2013258999, 909000\u2013910213"
    assert len(cause["sample_rows"]) == 20 and cause["more_samples_in_marked_copy"]
    assert cause["copy_values"] == ["6150"]


def test_blank_deal_over_four_accounts_has_a_two_condition_recipe_and_a_breakdown(tmp_path: Path) -> None:
    plan = {1500: [(10, 10_009)], 1600: [(20_000, 29_999)], 8050: [(30_000, 39_999)], 9010: [(50_000, 59_999)]}
    deal = {"1500": "", "1600": "", "8050": "", "9010": ""}
    path = _gl_csv(tmp_path / "gl.csv", plan, 60_000, deal)
    report = scan(path, RULES, record_type="GL", columns={"Account": "GL Account", "Deal": "Investment"},
                  groups={"required": []}, default_breakdown="Account")
    cause = _cause(report, "deal_required", "Deal")
    assert cause["count"] == 40_000
    assert cause["locator"]["filter"] == (
        "In Excel: Data \u2192 Filter, column D \u201cInvestment\u201d = (Blanks), and column C \u201cGL Account\u201d "
        "= 1500, 1600, 8050 and 1 more \u2192 40,000 lines"
    )
    bd = {b["value"]: b for b in cause["breakdown"]}
    assert sorted(bd) == ["1500", "1600", "8050", "9010"]
    assert all(b["count"] == 10_000 for b in bd.values())
    assert bd["8050"]["copy_filter"] == "Investment = (Blanks); GL Account = 8050"
    assert bd["8050"]["row_ranges"] == [[30_000, 39_999]]


def test_marked_copy_keeps_every_original_cell_and_adds_only_the_check_column(tmp_path: Path) -> None:
    rows = [[f"J{i}", i, "6150" if i in (3, 4) else "1000"] for i in range(1, 9)]
    src = _xlsx(tmp_path / "gl.xlsx", ["Journal", "Line", "GL Account"], rows)
    before = src.read_bytes()
    report = scan(src, RULES, record_type="GL", columns={"Account": "GL Account"}, references={"coa_sources": {"1000"}})
    out = marked_copy(report, tmp_path / "gl_marked.xlsx")
    assert src.read_bytes() == before  # the user's file is untouched
    original = [list(r) for r in openpyxl.load_workbook(src).active.iter_rows(values_only=True)]
    marked = openpyxl.load_workbook(out).active
    got = [list(r) for r in marked.iter_rows(values_only=True)]
    assert got[0] == [*original[0], MARK_COLUMN]
    assert [r[:-1] for r in got] == original  # every original cell, byte for byte in value
    assert [r[-1] for r in got[1:]] == [None, None, "Account has no match in the chart of accounts",
                                         "Account has no match in the chart of accounts", None, None, None, None]
    assert marked.auto_filter.ref == "A1:D9"
    assert marked.cell(row=4, column=1).fill.start_color.rgb.endswith("FFF2CC")


def test_marked_copy_for_one_value_and_just_these_rows(tmp_path: Path) -> None:
    plan = {1500: [(2, 5)], 8050: [(6, 9)]}
    path = _gl_csv(tmp_path / "gl.csv", plan, 12, {"1500": "", "8050": ""})
    report = scan(path, RULES, record_type="GL", columns={"Account": "GL Account", "Deal": "Investment"},
                  groups={"required": []}, default_breakdown="Account")
    only = marked_copy(report, tmp_path / "gl_marked_8050.csv", only=("Account", "8050"))
    lines = list(csv.reader(open(only, encoding="utf-8-sig")))  # noqa: SIM115
    marks = {int(r[1]): r[-1] for r in lines[1:]}
    assert all(marks[n] for n in (6, 7, 8, 9)) and not any(marks[n] for n in (2, 3, 4, 5, 10, 11, 12))
    rows = just_these_rows(report, tmp_path / "just.csv", only=("Account", "8050"))
    got = list(csv.reader(open(rows, encoding="utf-8-sig")))  # noqa: SIM115
    assert got[0][0] == "Original row" and [int(r[0]) for r in got[1:]] == [6, 7, 8, 9]


def test_marked_copy_refuses_to_overwrite_the_original(tmp_path: Path) -> None:
    src = _xlsx(tmp_path / "a.xlsx", ["A"], [[None]])
    report = scan(src, RULES, record_type="Affiliates", columns={"Affiliate_Name": "A"},
                  groups={"required": ["Affiliate_Name"]})
    with pytest.raises(ValueError, match="different file"):
        marked_copy(report, src)


def test_a_million_row_file_streams_within_a_memory_limit(tmp_path: Path) -> None:
    script = tmp_path / "big.py"
    script.write_text(
        "import csv, sys\n"
        "from pathlib import Path\n"
        f"import sys; sys.path.insert(0, {str(Path(__file__).parent)!r})\n"
        "from memprobe import peak_memory_mb\n"
        "from worker.quality import load_rules, scan, marked_copy, just_these_rows\n"
        "p = Path(sys.argv[1])\n"
        "with open(p, 'w', newline='') as fh:\n"
        "    w = csv.writer(fh); w.writerow(['Journal','GL Account','Investment'])\n"
        "    for i in range(1_000_000):\n"
        "        w.writerow([f'J{i//2}', '6150' if i % 100 == 0 else '1000', '' if i % 7 == 0 else 'D'])\n"
        "base = peak_memory_mb()\n"
        "rep = scan(p, load_rules(), record_type='GL', columns={'Account':'GL Account','Deal':'Investment'},\n"
        "           groups={'required': []}, references={'coa_sources': {'1000'}}, default_breakdown='Account')\n"
        "marked_copy(rep, p.with_name('marked.csv')); just_these_rows(rep, p.with_name('just.csv'))\n"
        "print('ROWS', rep['rows'], 'CAUSES', len(rep['causes']), 'PEAK', peak_memory_mb(), 'BASE', base)\n",
        encoding="utf-8",
    )
    env = {"SYSTEMROOT": os.environ.get("SYSTEMROOT", ""), "PATH": ""}
    done = subprocess.run(
        [sys.executable, "-I", str(script), str(tmp_path / "gl.csv")],
        capture_output=True, text=True, timeout=600, env=env, check=False,
    )
    out = done.stdout.strip().splitlines()[-1] if done.stdout.strip() else done.stderr[-500:]
    assert done.returncode == 0, out
    parts = out.split()
    assert parts[1] == "1000000" and float(parts[5]) < 400  # MB, with the 1,000,000-row file never in memory


def test_tools_find_supply_and_write_files_through_the_agent(tmp_path: Path) -> None:
    import json as _json

    from langgraph.checkpoint.memory import InMemorySaver

    from scripted_model import ScriptedModel, call, say
    from worker.agent import build_agent, decide, pending_gate, start
    from worker.audit import AuditLog
    from worker.context import RunContext
    from worker.staging import Stage

    inbox = tmp_path / "inbox"
    inbox.mkdir()
    rows = [[f"HS-{i:04d}", f"Investor {i}"] for i in range(1, 7)]
    rows[3][1] = None  # row 5
    _xlsx(inbox / "inv.xlsx", ["Row Labels", "Investor"], rows, "Investor")
    out = tmp_path / "run" / "out"
    ctx = RunContext(sponsor="sponsor-a", gp_id="GP-001", inbox_roots=[inbox], out_dir=out,
                     audit=AuditLog(out / "audit.jsonl"), stage=Stage(out))
    cols = {"Investor_ID": "Row Labels", "Investor_Name": "Investor"}
    script = [
        say("check", call("check_data_quality", path="inv.xlsx", record_type="Investors", columns=cols,
                          sheet="Investor", why="what is wrong in this file?")),
        say("files", call("write_quality_files", path="inv.xlsx", kind="marked", why="show where")),
        say("supply", call("supply_value", path="inv.xlsx", field="Investor_Name", value="Alder Co LLC", row=5,
                           why="the human typed the name")),
        say("again", call("check_data_quality", path="inv.xlsx", record_type="Investors", columns=cols,
                          sheet="Investor", why="re-check with the value")),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    result = start(agent, ctx, "check the file", "t1")
    assert [g["name"] for g in pending_gate(result)] == ["supply_value"]  # a typed value needs the human gate
    assert (out / "inv_marked.xlsx").is_file()
    first = next(e for e in ctx.audit.read() if e["action"] == "check_data_quality")
    assert first["outputs"]["rows"] == 6 and first["outputs"]["err"] == 1
    decide(agent, ctx, "t1", {"type": "approve"}, "I typed that name", {"reviewer": "reviewer-a"})
    entries = ctx.audit.read()
    supplied = [e for e in entries if e["action"] == "value.supplied"]
    assert supplied and supplied[0]["outputs"] == {"value": "Alder Co LLC", "inferred": False}
    checks = [e for e in entries if e["action"] == "check_data_quality"]
    assert len(checks) == 2
    causes = [e for e in entries if e["action"] == "quality.cause"]
    assert len(causes) == 1 and _json.dumps(causes[0]["outputs"]).count("5") >= 1  # only the first run had one
