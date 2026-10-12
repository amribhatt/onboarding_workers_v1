"""F4: a GL file through the agent path (scripted model only) writes the console's file; a dirty one is refused."""

from __future__ import annotations

import datetime as dt
import os
import subprocess
import sys
import textwrap
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from langgraph.checkpoint.memory import InMemorySaver

import expected_rules as X
from id_decision import decide_ids
from scripted_model import ScriptedModel, approve, call, say
from worker import console_demo as D
from worker import console_engine as CE
from worker import sample as S
from worker.accounts import key_of
from worker.agent import GateRefused, build_agent, decide, pending_gate, start
from worker.audit import AuditLog
from worker.console import Console
from worker.context import RunContext
from worker.registry import Registry
from worker.staging import Stage

LP_COLUMNS = {"Account": "GL Account", "Deal": "Investment", "Journal": "Journal", "Debit": "Debit",
              "Credit": "Credit", "Investor_ID": "Investor", "Offset": "Investor"}
COLUMNS = {"Account": "GL Account", "Deal": "Investment", "Journal": "Journal", "Debit": "Debit",
           "Credit": "Credit", "Investor_ID": "Investor"}

# (journal, line, account, transaction type, debit, credit, investor): every journal balances to the cent
CLEAN_FUND = [
    ("J-3001", 1, "1000", "Capital call", 500000, 0, ""), ("J-3001", 2, "3000", "Capital call", 0, 500000, ""),
    ("J-3002", 1, "6100", "Expense", 12500, 0, ""), ("J-3002", 2, "1000", "Expense", 0, 12500, ""),
    ("J-3003", 1, "3100", "Distribution", 250000, 0, ""), ("J-3003", 2, "1000", "Distribution", 0, 250000, ""),
]


def write_clean(path: Path, lines: list[tuple[Any, ...]]) -> Path:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Data Source"
    ws.append(S.GLSRC)
    for j, ln, acct, tt, dr, cr, inv in lines:
        ws.append(["2025-12-31", j, ln, acct, "name", tt, dr or None, cr or None, f"{tt} {j}", inv or None, None,
                   None, None, "Fund A", "B1"])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def iso_dates(row: Any) -> list[Any]:
    """A row as the upload's reader shows it: a true Excel date cell (the writer's, since the user chose real dates)
    is its ISO date text, so a file reads the same as it did when the dates were written as text."""
    return [v.date().isoformat() if isinstance(v, dt.datetime) and v.time() == dt.time() else v for v in row]


def read_file(path: Path) -> list[tuple[Any, ...]]:
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = [tuple(iso_dates(r)) for r in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()
    width = len(rows[0])
    return [(*r, *([None] * (width - len(r)))) for r in rows]


def approve_statics(c: Console) -> None:
    """The five static types of Sponsor A, approved: GL lines point at those records."""
    for t in X.STATIC:
        c.act(f"start|Sponsor A|{t}")
        c.act("ask_set|cov|all")
        c.act("ask_sample")
        c.act("ask_go")
        j0 = c.job()
        assert j0 is not None
        if c.ex_pending(j0):
            c.act("ex|no")
        c.act("go")
        decide_ids(c, t, job_id=j0["id"])
        c.act("acceptall")
        decide_ids(c)
        c.act("approve", who="reviewer-a")


def console_file(tmp: Path, typ: str, build: Any) -> tuple[Console, dict[str, Any], Path]:
    """The console's output for a GL file: `build(console)` writes the source that replaces the sample."""
    c = Console(tmp / "console_inbox")
    c.act("prov|paid")
    approve_statics(c)
    if typ == "GL LP":  # LP lines are checked against the fund's own GL, which is approved first
        c.act("start|Sponsor A|GL Fund")
        c.act("ask_set|fund|Fund A")
        c.act("ask_sample")
        c.act("ask_go")
        c.act("go")
        c.act("acceptall")
        decide_ids(c)
        c.act("approve", who="reviewer-a")
    c.act(f"start|Sponsor A|{typ}")
    c.act("ask_set|fund|Fund A")
    c.act("ask_sample")
    src = build(c)
    Path(c.S["ask"]["path"]).write_bytes(src.read_bytes())
    c.act("ask_go")
    j = c.job()
    assert j is not None
    c.act("go")
    return c, j, src


def agent_ctx(tmp: Path, state: Path | None = None) -> RunContext:
    state = state or tmp / "state"
    out = tmp / "run" / "out"
    inbox = tmp / "agent_inbox"
    inbox.mkdir(parents=True, exist_ok=True)
    return RunContext(sponsor="sponsor-a", gp_id="sponsor-a", inbox_roots=[inbox], out_dir=out,
                      audit=AuditLog(out / "audit.jsonl"), stage=Stage(out), accounts_root=state / "accounts",
                      registry=Registry(state, "sponsor-a"))


def picks_for(lines: list[tuple[Any, ...]], c: Console) -> dict[str, str]:
    """The account the console knows for each distinct account + type key (what the person confirms)."""
    return {key_of("", acct, tt): c.known_account("Sponsor A", acct, tt) for _j, _l, acct, tt, *_r in lines}


def test_F4_a_clean_gl_file_through_the_agent_equals_the_consoles_file(tmp_path: Path) -> None:
    c, j, src = console_file(tmp_path, "GL Fund", lambda _c: write_clean(tmp_path / "src" / "gl.xlsx", CLEAN_FUND))
    c.act("acceptall")
    assert CE.tally(c, j)["err"] == 0
    decide_ids(c)
    c.act("approve", who="reviewer-a")
    want = read_file(Path(j["out"]["path"]))
    assert len(want) == 1 + len(CLEAN_FUND)

    ctx = agent_ctx(tmp_path)
    (ctx.inbox_roots[0] / "gl.xlsx").write_bytes(src.read_bytes())
    picks = picks_for(CLEAN_FUND, c)
    script = [
        say("lift", call("lift_gl_lines", path="gl.xlsx", fund_id=c.entity_id("Fund A"), fund_name="Fund A",
                         why="stage the GL lines")),
        say("map", call("map_accounts", stage_id="s1", why="map the GL accounts")),
        say("confirm", call("confirm_account_mappings", stage_id="s2", picks=picks, why="the person confirmed")),
        say("check", call("check_data_quality", path="gl.xlsx", record_type="GL", columns=COLUMNS,
                          why="find every problem")),
        say("write", call("write_upload", stage_id="s3", template="GL", why="deliver the file")),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    result = start(agent, ctx, "build me the GL sheet", "t1")
    assert [g["name"] for g in pending_gate(result)] == ["confirm_account_mappings"]
    result = decide(agent, ctx, "t1", {"type": "approve"}, "these are the accounts", {"reviewer": "reviewer-a"})
    assert [g["name"] for g in pending_gate(result)] == ["write_upload"]
    assert not list(ctx.out_dir.glob("GL_Upload*"))  # nothing is written before the human decides
    approve(agent, ctx, "t1", "checked against the source")
    got = read_file(ctx.out_dir / "GL_Upload.xlsx")
    assert got == want  # same header, same rows, same offsets, same totals
    total = sum(r[got[0].index("Total_Amount")] for r in got[1:])
    assert total == 0
    assert ctx.audit.verify()[0]
    actions = [e["action"] for e in ctx.audit.read()]
    for step in ("lift_gl_lines", "map_accounts", "account.confirmed", "check_data_quality", "gate.approve",
                 "write_upload"):
        assert step in actions
    # the amounts were made by code from Debit and Credit: the staged table carries them signed
    assert [r["Total_Amount"] for r in ctx.stage.rows("s1")] == [500000, -500000, 12500, -12500, 250000, -250000]


def test_F4_gl_lp_offsets_and_investor_lookup_equal_the_console(tmp_path: Path) -> None:
    lp_lines: list[tuple[Any, ...]] = []

    def build(c: Console) -> Path:
        two = [e["source_key"] or e["key"] for e in c.registry("Sponsor A").entries("Investors")[:2]]
        lp_lines.extend([("J-4001", 1, "3000", "Capital call", 0, 300000, two[0]),
                         ("J-4001", 2, "3000", "Capital call", 0, 150000, two[1]),
                         ("J-4002", 1, "3100", "Distribution", 150000, 0, two[1])])
        return write_clean(tmp_path / "src" / "lp.xlsx", lp_lines)

    c, j, src = console_file(tmp_path, "GL LP", build)
    c.act("acceptall")
    assert CE.tally(c, j)["err"] == 0, CE.tally(c, j)
    decide_ids(c)
    c.act("approve", who="reviewer-a")
    want = read_file(Path(j["out"]["path"]))
    assert len(want) == 1 + 2 * len(lp_lines)

    ctx = agent_ctx(tmp_path, state=c.root)  # the console and the agent share one registry
    (ctx.inbox_roots[0] / "lp.xlsx").write_bytes(src.read_bytes())
    script = [
        say("lift", call("lift_gl_lines", path="lp.xlsx", fund_id=c.entity_id("Fund A"), fund_name="Fund A",
                         offsets=True, why="stage the LP lines with their offsets")),
        say("map", call("map_accounts", stage_id="s1", why="map")),
        say("confirm", call("confirm_account_mappings", stage_id="s2", picks=picks_for(lp_lines, c), why="confirmed")),
        say("check", call("check_data_quality", path="lp.xlsx", record_type="GL", columns=LP_COLUMNS, why="check")),
        say("write", call("write_upload", stage_id="s3", template="GL", why="deliver")),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "build me the LP GL sheet", "t1")
    decide(agent, ctx, "t1", {"type": "approve"}, "these are the accounts", {"reviewer": "reviewer-a"})
    approve(agent, ctx, "t1", "checked")
    got = read_file(ctx.out_dir / "GL_Upload.xlsx")
    assert got == want
    header = list(got[0])
    assert sum(1 for r in got[1:] if str(r[header.index("Event_ID")]).endswith("-OFF")) == len(lp_lines)
    assert all(r[header.index("Investor_ID")] for r in got[1:])  # looked up in the shared store, never typed


def test_F4_a_dirty_gl_file_is_refused_at_the_gate_with_locators_and_a_marked_copy(tmp_path: Path) -> None:
    src = S.write_gl(tmp_path / "src" / "gl.xlsx", "GL Fund")
    ctx = agent_ctx(tmp_path)
    (ctx.inbox_roots[0] / "gl.xlsx").write_bytes(src.read_bytes())
    golden = {a: g.split(" ", 1)[0] for a, g, flag in D.ACCTS if flag}
    picks = {key_of("", ln.acct, ln.tt): golden[ln.acct] for ln in S.GLD["GL Fund"]["lines"] if ln.acct in golden}
    script = [
        say("lift", call("lift_gl_lines", path="gl.xlsx", fund_id="5317401", fund_name="Fund A", why="stage")),
        say("map", call("map_accounts", stage_id="s1", why="map")),
        say("confirm", call("confirm_account_mappings", stage_id="s2", picks=picks, why="the person confirmed")),
        say("check", call("check_data_quality", path="gl.xlsx", record_type="GL", columns=COLUMNS, why="check")),
        say("write", call("write_upload", stage_id="s3", template="GL", why="deliver")),
        say("marked", call("write_quality_files", path="gl.xlsx", kind="marked", why="show the problems")),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "build me the GL sheet", "t1")
    result = decide(agent, ctx, "t1", {"type": "approve"}, "these are the accounts", {"reviewer": "reviewer-a"})
    assert [g["name"] for g in pending_gate(result)] == ["write_upload"]
    with pytest.raises(GateRefused) as refused:
        approve(agent, ctx, "t1", "looks fine")
    assert "ERR must be resolved" in str(refused.value)
    found = {f["rule"]: f for f in refused.value.review["findings"] if f["code"] == "data_quality"}
    acct, deal = found["account_not_in_chart"], found["deal_required"]
    assert acct["count"] == 2 and "6150" in acct["locator"]["filter"] and acct["ranges"] == "9, 19"
    assert deal["count"] == 4 and "(Blanks)" in deal["locator"]["filter"]
    assert "journal_unbalanced" in found  # the two off-by-cents journals are reported too
    assert acct["locator"]["where"].startswith("gl.xlsx")
    assert deal["offer"] == {"tool": "write_quality_files", "kind": "marked", "cause_id": deal["cause"]}
    assert not list(ctx.out_dir.glob("GL_Upload*"))
    # the person rejects; the agent offers the marked copy and the original is untouched
    before = (ctx.inbox_roots[0] / "gl.xlsx").read_bytes()
    decide(agent, ctx, "t1", {"type": "reject", "message": "fix the file first"})
    marked = ctx.out_dir / "gl_marked.xlsx"
    assert marked.is_file() and (ctx.inbox_roots[0] / "gl.xlsx").read_bytes() == before
    notes = [str(r[-1]) for r in read_file(marked)[1:] if r[-1]]
    assert any("no match in the chart" in n for n in notes) and any("Deal is blank" in n for n in notes)
    assert not list(ctx.out_dir.glob("GL_Upload*"))


def test_F4_a_gl_file_that_was_never_checked_needs_a_reason_at_the_gate(tmp_path: Path) -> None:
    src = write_clean(tmp_path / "src" / "gl.xlsx", CLEAN_FUND)
    ctx = agent_ctx(tmp_path)
    (ctx.inbox_roots[0] / "gl.xlsx").write_bytes(src.read_bytes())
    script = [
        say("lift", call("lift_gl_lines", path="gl.xlsx", fund_id="5317401", fund_name="Fund A", why="stage")),
        say("write", call("write_upload", stage_id="s1", template="GL", why="deliver")),
        say("done"),
    ]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "build me the GL sheet", "t1")
    with pytest.raises(GateRefused, match="not run through check_data_quality|account"):
        decide(agent, ctx, "t1", {"type": "approve"}, "go", {"reviewer": "reviewer-a"})


def test_F4_the_staged_gl_survives_a_restart_and_replays_from_its_spec(tmp_path: Path) -> None:
    from worker.memory import apply_spec

    src = write_clean(tmp_path / "src" / "gl.xlsx", [("J-4001", 1, "3000", "Capital call", 0, 300000, "HS-1"),
                                                      ("J-4002", 1, "3100", "Distribution", 150000, 0, "")])
    ctx = agent_ctx(tmp_path)
    (ctx.inbox_roots[0] / "gl.xlsx").write_bytes(src.read_bytes())
    script = [say("lift", call("lift_gl_lines", path="gl.xlsx", fund_id="5317401", fund_name="Fund A", offsets=True,
                               why="stage")), say("done")]
    agent = build_agent(ctx, ScriptedModel(script=script), InMemorySaver())
    start(agent, ctx, "stage it", "t1")
    again = Stage(ctx.out_dir)
    assert again.restore() == ["s1"] and again.rows("s1") == ctx.stage.rows("s1")
    assert again.meta("s1")["generated_rows"] == 1 and again.meta("s1")["input_rows"] == 2
    spec = {"kind": "gl_lines", "path": str(ctx.inbox_roots[0] / "gl.xlsx"), "sheet": "Data Source", "columns": {},
            "header_row": 1, "gp_id": "sponsor-a", "fund_id": "5317401", "fund_name": "Fund A", "investors": {},
            "offsets": True}
    replayed = apply_spec(Stage(tmp_path / "other"), "GL", spec)
    assert replayed == "s1"


@pytest.mark.slow
def test_F4_a_million_line_dirty_gl_file_is_checked_by_the_agent_with_locators_and_a_marked_copy(
    tmp_path: Path,
) -> None:
    """The tools the agent uses on a million lines: the streaming check, the locators, the marked copy. The staged
    table of a million lines is not built (the check refuses first), so memory stays bounded."""
    script = tmp_path / "big_agent.py"
    script.write_text(textwrap.dedent('''
        import sys
        from pathlib import Path
        sys.path.insert(0, sys.argv[2])
        import large_gl
        from langgraph.checkpoint.memory import InMemorySaver
        from scripted_model import ScriptedModel, call, say
        from worker.agent import build_agent, start
        from worker.audit import AuditLog
        from memprobe import peak_memory_mb
        from worker.context import RunContext
        from worker.staging import Stage

        root = Path(sys.argv[1])
        inbox = root / "inbox"
        inbox.mkdir()
        large_gl.build(inbox / "gl.csv")
        out = root / "out"
        ctx = RunContext(sponsor="sponsor-a", gp_id="sponsor-a", inbox_roots=[inbox], out_dir=out,
                         audit=AuditLog(out / "audit.jsonl"), stage=Stage(out), accounts_root=root / "accounts")
        from worker.accounts import AccountStore, key_of
        for acct in ("1000", "1500", "1600", "8050", "9010"):  # approved in an earlier job: only 6150 is unmapped
            AccountStore(root / "accounts", "sponsor-a").add(key_of("", acct, "Expense"), "11010",
                                                             source="earlier job", reason="approved earlier")
        cols = {"Account": "GL Account", "Deal": "Investment", "Journal": "Journal"}
        steps = [say("check", call("check_data_quality", path="gl.csv", record_type="GL", columns=cols,
                                   why="find every problem")),
                 say("marked", call("write_quality_files", path="gl.csv", kind="marked", why="show them")),
                 say("done")]
        agent = build_agent(ctx, ScriptedModel(script=steps), InMemorySaver())
        start(agent, ctx, "check the GL", "t1")
        rep = ctx.quality_reports[str((inbox / "gl.csv").resolve())]
        got = {c["rule"]: c["count"] for c in rep["causes"]}
        marked = out / "gl_marked.csv"
        print("RESULT", rep["rows"], got.get("account_not_in_chart"), got.get("deal_required"),
              marked.is_file(), round(peak_memory_mb()))
    '''), encoding="utf-8")
    env = {"SYSTEMROOT": os.environ.get("SYSTEMROOT", ""), "PATH": "", "PYTHONPATH": os.pathsep.join(
        [str(Path("src").resolve()), str(Path("tests").resolve())])}
    done = subprocess.run([sys.executable, str(script), str(tmp_path), str(Path("tests").resolve())],
                          capture_output=True, text=True, timeout=1500, env=env, check=False)
    line = next((x for x in done.stdout.splitlines() if x.startswith("RESULT")), done.stderr[-1500:])
    assert done.returncode == 0, line
    _, rows, unmapped, blank, marked, peak = line.split()
    assert int(rows) == 999_999 and int(unmapped) == 10_214 and int(blank) == 40_000
    assert marked == "True" and int(peak) < 600
