"""The eleven GL reconciliation checks of the GL Recon Checks guide: one engine and eleven declarative specs.

Code decides. The upload side of every check is computed here from the job's own lines (plain group-and-sum); the
client side comes from a report file the analyst supplies for that check, or from the other run of the same entity
(Entity Level vs LP Level). Without a report the client column is blank and the check says so: no client figure is
ever produced here. A guide step that needs a field the file does not carry is reported as a gap, never filled in.

The engine, in the guide's order:
  1. select rows by a predicate (account group / FS line, period ITD / YTD / before this year, investor present or
     blank, Trans Type words),
  2. group them by the check's keys (account, FS line, investment, investor, PR_DATE, Trans Type),
  3. sum a measure (Dr/(Cr), shares, a count),
  4. add the derived lines (FMV = cost + unrealised, net income, ending capital) as small named sums,
  5. compare each key with the client's figure for the same key,
  6. classify by the check's tolerance (controls.yaml: tiered, flat, whole number, percent) and roll up the worst.
The rules (specs, account series, Trans Type words, sections, applicability, tolerances) are the reconcile-gl
skill's data, loaded by `load_rules`; this module holds only the engine. Every sum, difference and tolerance
comparison is exact decimal arithmetic; figures leave as 2-decimal numbers. The same specs write the "Rows it
selects" and "How it runs" text, so the explanation cannot drift from the code.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import yaml

from worker import columnar, rulefiles
from worker.decide import Table, TableError
from worker.money import ZERO, cents, to_decimal, to_float
from worker.rulefiles import SchemaError
from worker.skills import Skill, bundled_skills_dir, find_skill, skills_dir

ORDER = {"ERROR": 4, "WARN": 3, "REVIEW": 2, "PASS": 1}
NOT_NAMED = "(no investment name)"


@dataclass(slots=True)
class Row:
    """One upload line as the checks see it. `inv` is blank on an offset (the guide's blank-investor row)."""

    acct: str
    fsg: str
    fsl: str
    amt: float  # Dr/(Cr): debit positive
    date: date | None
    inv: str
    tt: str
    invst: str = ""
    pr: date | None = None
    shares: float | None = None
    agn: bool | None = None
    offset: bool = False


@dataclass(frozen=True)
class Ctx:
    fy_start: date | None
    fy_end: date | None
    investors: Mapping[str, str]  # investor reference -> display name
    has_pr: bool
    has_shares: bool
    has_agn: bool
    undated: int


# ---- the rules: loaded from the reconcile-gl skill ----------------------------------------------------------
SKILL = "reconcile-gl"
FUND_TYPES = ("DE", "DLP", "EL", "LP")
TIER_STYLES = ("tiered", "flat")

# What a controller may set in controls.yaml. Everything outside these limits is refused at load.
CONTROL_LIMITS: dict[str, Any] = {
    "tolerances": {  # key: (largest pass_up_to, largest warn_up_to, warn must equal pass)
        "money-tiered": (Decimal(5), Decimal(25), False),
        "money-flat": (Decimal(5), Decimal(5), True),
        "int": (ZERO, ZERO, True),
        "pct": (Decimal("0.05"), Decimal("0.05"), True),
    },
    "levels": {  # name: levels it may take (never below the floor)
        "within_warn_band": ("WARN", "ERROR"),
        "over_warn_limit": ("ERROR",),
        "needs_judgement": ("REVIEW", "WARN", "ERROR"),
    },
    "required_floor": {  # a check may need more fields, never fewer
        "is": {"gl_date"}, "cap": {"gl_date"}, "cf": {"gl_date"}, "pcap": {"investor"},
        "irr": {"pr_date", "investor"}, "unf": {"investor"}, "wf": {"pr_date", "investor"},
    },
    "fields": {"gl_date", "pr_date", "investor"},
    "balance_target": ZERO,
}


class RulesError(ValueError):
    """The skill's rules are malformed or outside the limits the code allows."""


def _split(text: Any) -> tuple[str, ...]:
    return tuple(x.strip() for x in str(text).split(",") if x.strip() and x.strip() != "-")


def validate_controls(raw: Mapping[str, Any], check_ids: Iterable[str]) -> dict[str, Any]:
    """The controls inside the code's limits, as exact decimals. Raises RulesError naming the value that is out."""
    lim = CONTROL_LIMITS
    tol: dict[str, tuple[Decimal, Decimal]] = {}
    for key, (max_pass, max_warn, same) in lim["tolerances"].items():
        row = (raw.get("tolerances") or {}).get(key)
        if not isinstance(row, Mapping) or "pass_up_to" not in row or "warn_up_to" not in row:
            raise RulesError(f"controls: tolerance {key!r} needs pass_up_to and warn_up_to")
        p, w = to_decimal(row["pass_up_to"]), to_decimal(row["warn_up_to"])
        if p < 0 or p > max_pass or w < p or w > max_warn or (same and w != p):
            raise RulesError(f"controls: tolerance {key!r} ({p} / {w}) is outside the allowed limits "
                             f"(pass at most {max_pass}, warn at most {max_warn}{', equal to pass' if same else ''})")
        tol[key] = (p, w)
    levels: dict[str, str] = {}
    for name, allowed in lim["levels"].items():
        v = str((raw.get("levels") or {}).get(name, ""))
        if v not in allowed:
            raise RulesError(f"controls: level {name!r} must be one of {list(allowed)} (the floor), got {v!r}")
        levels[name] = v
    style = dict(raw.get("check_style") or {})
    for cid in check_ids:
        if style.get(cid) not in TIER_STYLES:
            raise RulesError(f"controls: check_style for {cid!r} must be one of {list(TIER_STYLES)}")
    bal = raw.get("balance") or {}
    target = to_decimal(bal["target"]) if "target" in bal else None
    if target is None or target != lim["balance_target"]:
        raise RulesError("controls: the balance target must be 0")
    required: dict[str, tuple[str, ...]] = {}
    for cid in check_ids:
        fields = [str(f) for f in (raw.get("required_fields") or {}).get(cid, [])]
        if set(fields) - lim["fields"]:
            raise RulesError(f"controls: required fields of {cid!r} must be among {sorted(lim['fields'])}")
        missing = lim["required_floor"].get(cid, set()) - set(fields)
        if missing:
            raise RulesError(f"controls: {cid!r} cannot drop the required field(s) {sorted(missing)}")
        required[cid] = tuple(f for f in ("gl_date", "pr_date", "investor") if f in fields)
    return {"tolerances": tol, "levels": levels, "check_style": style, "balance_target": target,
            "required": required}


@dataclass
class Rules:
    """What the engine applies: the skill's tables, controls and check specifications, with their identity."""

    skill: Skill
    tables: dict[str, Table]
    controls: dict[str, Any]
    specs: list[dict[str, Any]] = field(default_factory=list)
    sponsor: str = ""
    digest: str = ""
    aliases: dict[str, tuple[str, ...]] = field(default_factory=dict)  # report column -> header words, in priority
    _cache: dict[Any, Any] = field(default_factory=dict, repr=False)

    @property
    def stamp(self) -> str:
        """The human label, `reconcile-gl@1.0`."""
        return self.skill.stamp

    @property
    def ident(self) -> str:
        """Label plus what actually ran: `reconcile-gl@1.0#<content hash>`."""
        return f"{self.stamp}#{self.digest}"

    def decide(self, table: str, **inputs: Any) -> dict[str, Any]:
        key = (table, tuple(sorted(inputs.items())))
        if key not in self._cache:
            self._cache[key] = self.tables[table].decide(inputs).output
        return self._cache[key]

    def series(self, check: str, name: str, field_: str) -> tuple[str, ...]:
        out = self.decide("account-series", check=check, series=name)
        if field_ not in out:
            raise RulesError(f"account-series: check {check!r} asks for series {name!r} by {field_}, and the table "
                             "has no such row")
        return _split(out[field_])

    def words(self, name: str) -> tuple[str, ...]:
        return tuple(w.lower() for w in _split(self.decide("trans-type-words", words=name).get("words_list", "")))

    def flips(self, fsg: str) -> bool:
        return self.decide("presentation-sign", fsg=fsg).get("sign") == "flip"

    def section(self, line: str) -> str:
        return str(self.decide("cash-flow-sections", line=line).get("section", ""))

    def thresholds(self, style: str, kind: str) -> tuple[Decimal, Decimal]:
        key = "int" if kind == "int" else "pct" if kind == "pct" else f"money-{style}"
        return self.controls["tolerances"][key]

    def level(self, name: str) -> str:
        return str(self.controls["levels"][name])

    def combine(self, status: str, kind: str, has_client: bool, vs_zero: bool, diff: float) -> bool:
        """The cross-check table: may this line join a combined break?"""
        return bool(self.decide("cross-check", status=status, kind=kind, has_client=has_client, vs_zero=vs_zero,
                                abs_diff=float(abs(diff))).get("combine"))

    def view(self) -> dict[str, Any]:
        c = self.controls
        return {"tolerances": {k: {"pass_up_to": float(p), "warn_up_to": float(w)}
                               for k, (p, w) in c["tolerances"].items()},
                "levels": dict(c["levels"]), "check_style": dict(c["check_style"]),
                "balance_target": float(c["balance_target"]),
                "required_fields": {k: list(v) for k, v in c["required"].items() if v}}


def _load_skill(root: Path | None) -> Skill:
    for base in (root, skills_dir(), bundled_skills_dir()):
        if base is None:
            continue
        found = find_skill(SKILL, base)
        if found is not None:
            return found
    raise RulesError(f"the {SKILL} skill is not installed")


def _controls_first(skill: Skill) -> None:
    """The engine's own limits give the plainest message for a control that is out of them, so they speak first;
    anything else that is wrong with the files is the strict loader's to say."""
    try:
        controls = yaml.safe_load((skill.path / "controls.yaml").read_text(encoding="utf-8"))
        ids = [d["id"] for d in yaml.safe_load((skill.path / "specs.yaml").read_text(encoding="utf-8"))["checks"]]
    except (OSError, yaml.YAMLError, KeyError, TypeError):
        return
    if isinstance(controls, Mapping):
        validate_controls(controls, ids)


def _strict(skill: Skill) -> tuple[dict[str, Any], dict[str, Any]]:
    """The skill's specs and controls as the strict loader read them (copies: the reading is shared)."""
    try:
        pin = skill.pin
    except SchemaError as exc:
        raise RulesError(f"{SKILL}: " + "; ".join(exc.lines)) from exc
    kinds = {f.kind: f for f in pin.files.values() if not f.overlay}
    if "recon-specs" not in kinds or "controls" not in kinds:
        raise RulesError(f"{SKILL} needs specs.yaml and controls.yaml")
    return copy.deepcopy(dict(kinds["recon-specs"].data)), copy.deepcopy(dict(kinds["controls"].data))


def _text(s: str, p: Decimal, w: Decimal) -> str:
    return s.replace("{pass}", f"{p:f}".rstrip("0").rstrip(".") or "0").replace(
        "{warn}", f"{w:f}".rstrip("0").rstrip(".") or "0")


def _where(rules: Rules, check: str, w: Mapping[str, Any]) -> dict[str, Any]:
    """A selector as the specification writes it (series and word-list names) as the concrete predicate."""
    out: dict[str, Any] = {}
    for k, v in w.items():
        if k in ("fsg", "fsl"):
            names = str(v).split("+")
            out[k] = tuple(dict.fromkeys(x for n in names for x in rules.series(check, n, k)))
        elif k == "fsl_not":
            out[k] = tuple(dict.fromkeys(x for n in v for x in rules.series(check, n, "fsl")))
        elif k in ("fsl_prefix", "fsl_noprefix"):
            out[k] = rules.series(check, str(v), "fsl_prefix")[0]
        elif k == "tt_any":
            out[k] = rules.words(str(v))
        elif k == "tt_none":
            out[k] = tuple(dict.fromkeys(x for n in v for x in rules.words(n)))
        else:
            out[k] = v
    return out


def _build_specs(rules: Rules, raw: Mapping[str, Any]) -> list[dict[str, Any]]:
    specs: list[dict[str, Any]] = []
    for d in raw["checks"]:
        cid = d["id"]
        style = rules.controls["check_style"][cid]
        p, w = rules.thresholds(style, "money")
        app = rules.decide("fund-type-applicability", check=cid)
        req = rules.controls["required"][cid]
        s: dict[str, Any] = {k: v for k, v in d.items() if k not in ("measures", "derived", "key", "client")}
        for k in ("rule", "sel", "layout_text"):
            if k in s:
                s[k] = _text(s[k], p, w)
        s["steps"] = [_text(x, p, w) for x in d["steps"]]
        s["key"] = tuple(d["key"])
        s["scope"] = [ft for ft in FUND_TYPES if app.get(ft) is True]
        s["tol"] = style
        s["required"] = list(req)
        s["pr"] = "pr_date" in req
        s["lp"] = "investor" in req
        if "client" in d:
            c = dict(d["client"])
            c["key"] = tuple(c["key"])
            c["measure"] = tuple(c["measure"])
            s["client"] = c
        ms: dict[str, Any] = {}
        for mid, m in d["measures"].items():
            e = dict(m)
            e["key"] = None if m.get("key") is None else tuple(m["key"])
            e["where"] = [_where(rules, cid, w_) for w_ in m["where"]]
            if e.pop("balance", False):
                e["vs"] = 0.0
            ms[mid] = e
        s["measures"] = ms
        dv: dict[str, Any] = {}
        for did, x in (d.get("derived") or {}).items():
            e = dict(x)
            if e.pop("balance", False):
                e["vs"] = 0.0
            if "section" in e:
                want = e.pop("section")
                e["of"] = [(mid, 1) for mid, m in ms.items() if m.get("line") and rules.section(m["line"]) == want]
            else:
                e["of"] = [(a, b) for a, b in x["of"]]
            if e.get("key") is not None:
                e["key"] = tuple(e["key"])
            dv[did] = e
        if dv:
            s["derived"] = dv
        s["_rules"] = rules
        specs.append(s)
    return specs


def load_rules(sponsor_dir: Path | None = None, overlay: Mapping[str, Sequence[Mapping[str, Any]]] | None = None,
               skills_root: Path | None = None, sponsor: str = "",
               table_override: Mapping[str, Table] | None = None) -> Rules:
    """The reconcile-gl skill's rules. `sponsor_dir` is the sponsor's own folder, `templates/sponsors/<sponsor>`: its
    `skills/reconcile-gl/tables/` rows replace the organisation's by id; `overlay` is promoted rows per table id from
    the improvement gate. Every rule file of the skill goes through the strict loader; a refused file raises
    RulesError listing every problem."""
    skill = _load_skill(skills_root)
    _controls_first(skill)
    raw, raw_controls = _strict(skill)
    tdir = skill.path / "tables"
    tables: dict[str, Table] = {}
    for p in sorted(tdir.glob("*.yaml")) if tdir.is_dir() else []:
        try:
            tables[p.stem] = skill.table(p.stem, sponsor_dir, (overlay or {}).get(p.stem))
        except (TableError, ValueError) as exc:
            raise RulesError(f"{SKILL} table {p.stem}: {exc}") from exc
    tables.update(table_override or {})
    for need in ("account-series", "presentation-sign", "trans-type-words", "cash-flow-sections",
                 "fund-type-applicability", "cross-check"):
        if need not in tables:
            raise RulesError(f"{SKILL} has no table {need!r}")
    ids = [d["id"] for d in raw["checks"]]
    controls = validate_controls(raw_controls, ids)
    rules = Rules(skill, tables, controls, [], sponsor, skill.content_hash(sponsor_dir, overlay))
    rules.specs = _build_specs(rules, raw)
    rules.aliases = _report_aliases(skill, sponsor_dir)
    return rules


def _report_aliases(skill: Skill, sponsor_root: Path | None) -> dict[str, tuple[str, ...]]:
    """The report's header words: `synonyms.yaml` of the skill with the sponsor's own words in front. Add-only: a
    sponsor overlay can put words first and can never take one away."""
    try:
        base = next((f for f in skill.pin.files.values() if f.kind == "synonyms" and not f.overlay), None)
    except SchemaError as exc:
        raise RulesError(f"{SKILL}: " + "; ".join(exc.lines)) from exc
    if base is None:
        raise RulesError(f"{SKILL} needs synonyms.yaml")
    mine: Mapping[str, Any] = {}
    if sponsor_root is not None:
        path = sponsor_root / "skills" / SKILL / "synonyms.yaml"
        if path.is_file():
            try:
                mine = rulefiles.load_text(f"skills/{SKILL}/synonyms.yaml", path.read_text(encoding="utf-8"),
                                           skill=SKILL, base=base).data["concepts"]
            except SchemaError as exc:
                raise RulesError(f"{path}: " + "; ".join(exc.lines)) from exc
    theirs: Mapping[str, Any] = base.data["concepts"]
    out: dict[str, tuple[str, ...]] = {}
    for key in [*theirs, *(k for k in mine if k not in theirs)]:
        seen: set[str] = set()
        out[key] = tuple(w for w in [*mine.get(key, []), *theirs.get(key, [])]
                         if not (w.casefold() in seen or seen.add(w.casefold())))
    return out


def read_overlay(state_dir: str | Path) -> dict[str, list[dict[str, Any]]]:
    """Rows promoted through the improvement gate, per table id (`<state>/table_versions/reconcile-gl/<table>`)."""
    base = Path(state_dir) / "table_versions" / SKILL
    out: dict[str, list[dict[str, Any]]] = {}
    for d in sorted(base.iterdir()) if base.is_dir() else []:
        pointer = d / "CURRENT"
        try:
            n = int(pointer.read_text(encoding="utf-8").strip()) if pointer.is_file() else 0
            rows = json.loads((d / f"v{n}.json").read_text(encoding="utf-8"))["content"].get("rows", []) if n else []
        except (OSError, ValueError, KeyError):
            continue
        if rows:
            out[d.name] = rows
    return out


def rules_for(sponsor_dir: Path | None = None, state_dir: str | Path | None = None, sponsor: str = "") -> Rules:
    """The rules a job runs under: the skill, the rows promoted for the organisation, then the sponsor's own."""
    overlay = read_overlay(state_dir) if state_dir is not None else None
    if sponsor_dir is None and sponsor:  # the sponsor's own folder
        from worker.templates import sponsors_dir

        sponsor_dir = sponsors_dir() / sponsor
    return load_rules(sponsor_dir, overlay, sponsor=sponsor)


_DEFAULT: Rules | None = None


def default_rules() -> Rules:
    global _DEFAULT
    if _DEFAULT is None:
        _DEFAULT = load_rules()
    return _DEFAULT


_BOOT = default_rules()
SPECS: list[dict[str, Any]] = _BOOT.specs
SPEC = {s["id"]: s for s in SPECS}
CHECK_IDS = [s["id"] for s in SPECS]


def rules_of(spec: Mapping[str, Any]) -> Rules:
    return spec.get("_rules") or default_rules()  # type: ignore[no-any-return]


# ---- dates, rows ------------------------------------------------------------------------------------------
def parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            return None
    for f in ("%m/%d/%Y", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def _yes(v: str) -> bool | None:
    s = str(v or "").strip().lower()
    return True if s in ("y", "yes", "true", "1") else False if s in ("n", "no", "false", "0") else None


def make_ctx(rows: list[Row], investors: Mapping[str, str]) -> Ctx:
    dts = [r.date for r in rows if r.date]
    fy = max(dts).year if dts else None
    return Ctx(date(fy, 1, 1) if fy else None, date(fy, 12, 31) if fy else None, investors,
               any(r.pr for r in rows), any(r.shares is not None for r in rows), any(r.agn is not None for r in rows),
               sum(1 for r in rows if not r.date))


def _in_period(r: Row, period: str, ctx: Ctx) -> bool:
    if period == "ITD":
        return True
    if ctx.fy_start is None or r.date is None:
        return False
    if period == "YTD":
        return ctx.fy_start <= r.date <= ctx.fy_end  # type: ignore[operator]
    if period == "PRIOR":
        return r.date < ctx.fy_start
    return False


def _hit(r: Row, w: Mapping[str, Any], ctx: Ctx) -> bool:
    if "fsg" in w and r.fsg not in w["fsg"]:
        return False
    if "fsl" in w and r.fsl not in w["fsl"]:
        return False
    if "fsl_not" in w and r.fsl in w["fsl_not"]:
        return False
    if "fsl_prefix" in w and not r.fsl.startswith(w["fsl_prefix"]):
        return False
    if "fsl_noprefix" in w and r.fsl.startswith(w["fsl_noprefix"]):
        return False
    if "inv" in w and (bool(r.inv) and not r.offset) != (w["inv"] == "present"):
        return False
    tt = r.tt.lower()
    if "tt_any" in w and not any(k in tt for k in w["tt_any"]):
        return False
    if "tt_none" in w and any(k in tt for k in w["tt_none"]):
        return False
    if w.get("agn") and ctx.has_agn and r.agn is not True:
        return False
    if w.get("pr_ne_date") and not (r.pr and r.date and r.pr != r.date):
        return False
    return _in_period(r, w.get("period", "ITD"), ctx)


def _keypart(r: Row, name: str) -> str:
    if name == "acct":
        return r.acct
    if name == "fsl":
        return r.fsl
    if name == "invst":
        return r.invst or NOT_NAMED
    if name == "inv":
        return r.inv
    if name == "pr":
        return r.pr.isoformat() if r.pr else ""
    if name == "tt":
        return r.tt
    raise KeyError(name)


def _sign(r: Row, how: str, rules: Rules) -> int:
    if how == "flip":
        return -1
    if how == "natural":
        return -1 if rules.flips(r.fsg) else 1
    return 1


Items = dict[tuple[tuple[str, ...], str], Decimal]


def measure_items(spec: Mapping[str, Any], rows: Iterable[Row], ctx: Ctx, src: str = "self") -> Items:
    """Steps 1 to 3: select, group, sum. Only measures whose `src` is this source are summed."""
    out: Items = {}
    rules = rules_of(spec)
    ms = {mid: m for mid, m in spec["measures"].items() if m.get("src", "self") == src and not m.get("from_client")}
    for r in rows:
        for mid, m in ms.items():
            if not any(_hit(r, w, ctx) for w in m["where"]):
                continue
            if m["value"] == "shares":
                if r.shares is None:
                    continue
                v = to_decimal(r.shares)
            elif m["value"] == "count":
                v = Decimal(1)
            else:
                v = to_decimal(r.amt) * _sign(r, m["sign"], rules)
            k = tuple(_keypart(r, p) for p in (m["key"] if m["key"] is not None else spec["key"]))
            out[(k, mid)] = out.get((k, mid), ZERO) + v
    return out


def derive(spec: Mapping[str, Any], items: Items) -> Items:
    """Step 4: the derived lines, small named sums over the grouped results."""
    out = dict(items)
    for did, d in spec.get("derived", {}).items():
        parts = [(p, s) for p, s in d["of"]]
        keys = {k for (k, mid) in out if mid in {p for p, _ in parts}}
        if d.get("key") is not None:
            keys = {d["key"]}
        for k in keys:
            if d.get("need") and (k, d["need"]) not in out:
                continue
            if all((k, p) not in out for p, _ in parts):
                continue
            out[(k, did)] = sum((s * out.get((k, p), ZERO) for p, s in parts), ZERO)
    return out


# ---- comparing -------------------------------------------------------------------------------------------
def classify(a: Any, b: Any, kind: str, tol: str, rules: Rules | None = None) -> tuple[float, str]:
    """Difference (2 decimals) and flag, from the controls' thresholds for this tolerance style and kind."""
    r = rules or default_rules()
    d = cents(cents(to_decimal(a)) - cents(to_decimal(b)))
    ad = abs(d)
    low, high = r.thresholds(tol, kind)
    if ad <= low:
        return to_float(d), "PASS"
    return to_float(d), r.level("within_warn_band") if ad <= high else r.level("over_warn_limit")


def worst(rows: list[dict[str, Any]]) -> str:
    w = ""
    for r in rows:
        s = r.get("status", "")
        if s in ORDER and (not w or ORDER[s] > ORDER[w]):
            w = s
    return w


def _kind(spec: Mapping[str, Any], mid: str) -> str:
    m = spec["measures"].get(mid)
    return m["kind"] if m else "money"


def _label(spec: Mapping[str, Any], mid: str) -> str:
    m = spec["measures"].get(mid)
    return (m or spec.get("derived", {}).get(mid, {})).get("label", "")


def _keyname(spec: Mapping[str, Any], mid: str) -> tuple[str, ...]:
    m = spec["measures"].get(mid)
    if m is not None:
        return m["key"] if m["key"] is not None else spec["key"]
    d = spec.get("derived", {}).get(mid, {})
    if d.get("key") is not None:
        return d["key"]
    of = d.get("of") or []
    return _keyname(spec, of[0][0]) if of else spec["key"]


def key_text(spec: Mapping[str, Any], names: tuple[str, ...], key: tuple[str, ...], ctx: Ctx,
             coa: Mapping[str, Mapping[str, str]]) -> str:
    out = []
    for n, v in zip(names, key, strict=False):
        if n == "acct":
            nm = coa.get(v, {}).get("Account_Name", "")
            out.append(f"{v} {nm}".strip())
        elif n == "inv":
            out.append(ctx.investors.get(v, v) or "(blank investor)")
        elif n == "tt" and not v:
            continue
        else:
            out.append(v or "(blank)")
    return " · ".join(out)


# ---- the client's report ----------------------------------------------------------------------------------
class ReportError(ValueError):
    """A report file the app cannot read as the documented layout."""


def _norm_head(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _rows_of(path: Path) -> list[list[Any]]:
    return [r for _n, r in columnar.rows(path)]


def _amount(v: Any) -> Decimal | None:
    """A report cell as an exact amount: "$", "%" and thousands commas dropped, (1,234) negative; None when blank
    or not a number."""
    if v is None or isinstance(v, str) and v.strip() == "":
        return None
    if isinstance(v, str):
        v = v.strip().replace("$", "").replace("%", "")
    try:
        return to_decimal(v)
    except ValueError:
        return None


def read_report(path: str | Path, aliases: Mapping[str, Sequence[str]] | None = None) -> list[dict[str, Any]]:
    """The client's report as normalized rows {line, investor, date, tt, measure, period, amount}. Reads the
    documented long layout, and for the waterfall a table with dates as columns and investors as rows. The header
    words are `aliases` (a sponsor's, from `Rules.aliases`) or the skill's `synonyms.yaml`."""
    aliases = aliases if aliases is not None else default_rules().aliases
    raw = [r for r in _rows_of(Path(path)) if any(str(c or "").strip() for c in r)]
    if not raw:
        raise ReportError("the report file has no rows")
    head = [_norm_head(h) for h in raw[0]]
    col: dict[str, int] = {}
    for key, names in aliases.items():
        for n in map(_norm_head, names):
            if n in head:
                col.setdefault(key, head.index(n))
                break
    out: list[dict[str, Any]] = []
    if "amount" not in col:
        dates = {i: parse_date(raw[0][i]) for i in range(len(head)) if parse_date(raw[0][i])}
        who = col.get("investor", 0)
        if not dates:
            raise ReportError("no Amount column and no date columns: use the layout from the Reconcile page")
        for r in raw[1:]:
            for i, d in dates.items():
                a = _amount(r[i] if i < len(r) else None)
                if a is not None:
                    out.append({"line": "", "investor": str(r[who] or "").strip(), "date": d.isoformat(), "tt": "",
                                "measure": "", "period": "", "amount": a})
        return out
    for r in raw[1:]:
        def g(k: str, r: list[Any] = r) -> Any:
            i = col.get(k)
            return r[i] if i is not None and i < len(r) else None

        d = parse_date(g("date"))
        a = _amount(g("amount"))
        out.append({"line": str(g("line") or "").strip(), "investor": str(g("investor") or "").strip(),
                    "date": d.isoformat() if d else str(g("date") or "").strip(), "tt": str(g("tt") or "").strip(),
                    "measure": str(g("measure") or "").strip(), "period": str(g("period") or "").strip().upper(),
                    "amount": a})
    if not out:
        raise ReportError("the report file has a header but no rows")
    return out


def client_items(spec: Mapping[str, Any], report: list[dict[str, Any]], ctx: Ctx, resolve_acct: Any,
                 aliases: Mapping[str, str]) -> tuple[Items, list[dict[str, Any]]]:
    """Place each report row on a (key, measure) of the check. Rows that cannot be placed come back as `issues`
    (reviewed by the analyst; never matched by guess)."""
    cl = spec["client"]
    how, table = cl["measure"]
    out: Items = {}
    issues: list[dict[str, Any]] = []
    labelled: dict[str, str] = {}
    for mid in list(spec["measures"]) + list(spec.get("derived", {})):
        lab = _label(spec, mid)
        if lab:
            labelled.setdefault(lab.casefold(), mid)
    for nr in report:
        line = nr["line"]
        low = line.casefold()
        if cl.get("irr") and low == cl["irr"].casefold():
            continue  # read by the IRR step, not placed as a cash flow
        if nr["amount"] is None:
            issues.append({"label": line or nr["investor"], "why": "amount missing or not a number", "b": None})
            continue
        mid: str | None = None
        if cl.get("terminal") and low == cl["terminal"].casefold():
            mid = "end"
        elif how == "fixed":
            tot = labelled.get(low)
            mid = tot if tot and _keyname(spec, tot) == () and tot != table else table
        elif how == "col":
            mid = table.get(nr["measure"].casefold())
        elif table is not None:
            mid = table.get(low)
        else:
            mid = labelled.get(low + (" (ytd)" if nr["period"] == "YTD" else ""))
        if mid is None:
            issues.append({"label": line or nr["investor"], "why": "this line is not one the check lists; match it "
                           "by meaning", "b": to_float(nr["amount"])})
            continue
        key: list[str] = []
        bad = ""
        for n in _keyname(spec, mid):
            if n == "acct":
                g = resolve_acct(line, nr["tt"])
                if not g:
                    bad = f"client account {line!r} has no Ledgerline account in the account mapping"
                key.append(g or line)
            elif n == "inv":
                ref = aliases.get((nr["investor"] or "").casefold())
                if ref is None:
                    bad = f"investor {nr['investor']!r} is not one of this entity's investors"
                key.append(ref or nr["investor"])
            elif n == "pr":
                key.append(nr["date"])
            elif n == "tt":
                key.append(nr["tt"])
            else:
                key.append(line)
        if bad:
            issues.append({"label": line or nr["investor"], "why": bad, "b": to_float(nr["amount"])})
            continue
        k = (tuple(key), mid)
        out[k] = out.get(k, ZERO) + to_decimal(nr["amount"])
    return out, issues


# ---- running a check ---------------------------------------------------------------------------------------
def _gaps(spec: Mapping[str, Any], ctx: Ctx, rows: list[Row]) -> list[str]:
    g: list[str] = []
    req = spec.get("required", [])
    if "gl_date" in req and ctx.fy_start is None:
        g.append("GL Date is blank or unreadable on every line, so the current fiscal year cannot be found and the "
                 "year-to-date figures cannot be computed.")
    if "pr_date" in req and not ctx.has_pr:
        g.append("PR_DATE is not in this GL file (the upload carries Event_Date and Effective_Date only). The guide "
                 "needs PR_DATE as the cash-flow date, so no figure is shown rather than using the GL date.")
    if "investor" in req and not any(r.inv and not r.offset for r in rows):
        g.append("This is not an LP-level upload: no line has an investor.")
    return g


def _pair_rows(spec: Mapping[str, Any], items: Items, ctx: Ctx, coa: Mapping[str, Mapping[str, str]]
               ) -> list[dict[str, Any]]:
    accts = sorted({k for (k, mid) in items if mid in ("fund", "lp", "off")})
    out = []
    for k in accts:
        f, lp = cents(items.get((k, "fund"), ZERO)), cents(items.get((k, "lp"), ZERO))
        d, st = classify(f, lp, "money", spec["tol"], rules_of(spec))
        out.append({"label": key_text(spec, ("acct",), k, ctx, coa), "a": to_float(f), "b": to_float(lp), "diff": d,
                    "kind": "money",
                    "status": st, "note": "", "mid": "fund"})
    for k in accts:
        t = cents(items.get((k, "offtest"), ZERO))
        d, st = classify(t, ZERO, "money", spec["tol"], rules_of(spec))
        out.append({"label": key_text(spec, ("acct",), k, ctx, coa) + " · LP + offset", "a": to_float(t), "b": 0.0, "diff": d,
                    "kind": "money", "status": st, "note": "should net to zero", "mid": "offtest"})
    return out


def run_check(spec: Mapping[str, Any], rows: list[Row], ctx: Ctx, coa: Mapping[str, Mapping[str, str]],
              report: list[dict[str, Any]] | None, resolve_acct: Any, aliases: Mapping[str, str],
              fund_rows: list[Row] | None = None, lp_rows: list[Row] | None = None,
              incomplete: bool = False) -> dict[str, Any]:
    """One check: {rows, state, notes}. `state` is gap | needs-report | run. `rows` carry the upload figure `a`,
    the client figure `b` (None when there is no report), the difference and the flag."""
    notes: list[str] = []
    rules = rules_of(spec)
    review = rules.level("needs_judgement")
    gaps = _gaps(spec, ctx, rows)
    if gaps:
        return {"rows": [], "state": "gap", "notes": gaps}
    ms = spec["measures"]
    if spec.get("pair"):
        items: Items = {}
        for src, rs in (("fund", fund_rows), ("lp", lp_rows)):
            items.update(measure_items(spec, rs or [], ctx, src))
        return {"rows": _pair_rows(spec, derive(spec, items), ctx, coa), "state": "run", "notes": notes}
    items = measure_items(spec, rows, ctx)
    if any(m.get("src") == "fund" for m in ms.values()):
        if fund_rows is not None:
            items.update(measure_items(spec, fund_rows, ctx, "fund"))
        else:
            notes.append("The fund-level total is compared only when the fund run of this entity is in the app; it "
                         "has not been run here yet.")
    have = report is not None
    cl_items: Items = {}
    issues: list[dict[str, Any]] = []
    if have:
        cl_items, issues = client_items(spec, report or [], ctx, resolve_acct, aliases)
        if spec.get("collapse_tt") and report and not any(r["tt"] for r in report):
            items = _collapse(items)
        for k, v in cl_items.items():
            if ms.get(k[1], {}).get("from_client"):
                items[k] = v
    items = derive(spec, items)
    if spec["id"] == "soi" and not ctx.has_shares:
        notes.append("NO_OF_SHARES is not in this GL file, so shares cannot be shown or compared.")
    no_agn = spec["id"] == "unf" and not ctx.has_agn
    if no_agn:
        notes.append("AGNST_COMMIT is not in this GL file, so contributions called cannot be limited to the rows "
                     "counted against commitment: every contribution is counted and each figure is a REVIEW.")
    keys_by_mid: dict[str, set[tuple[str, ...]]] = {}
    for (k, mid) in list(items) + list(cl_items):
        keys_by_mid.setdefault(mid, set()).add(k)
    entries: list[tuple[Any, str, tuple[str, ...]]] = []
    for idx, mid in enumerate(spec["order"]):
        for k in keys_by_mid.get(mid, ()):
            entries.append(((k, idx) if spec.get("interleave") else (idx, k), mid, k))
    entries.sort(key=lambda e: e[0])
    out: list[dict[str, Any]] = []
    for _sk, mid, k in entries:
        m = ms.get(mid)
        d = spec.get("derived", {}).get(mid, {})
        if (m and (m.get("hidden") or m.get("from_client"))) or (d.get("ytd") and ctx.fy_start is None) or \
                (m and m.get("ytd") and ctx.fy_start is None):
            continue
        kind = m["kind"] if m else d.get("kind", "money")
        label = " · ".join(x for x in (key_text(spec, _keyname(spec, mid), k, ctx, coa), _label(spec, mid)) if x)
        a, b = items.get((k, mid)), (cl_items.get((k, mid)) if have else None)
        base = {"label": label, "kind": kind, "mid": mid}
        if m and m.get("review"):
            if a:
                why = ("these lines carry a PR_DATE that differs from the GL date: the analyst confirms the date"
                       if mid == "pr_diff" else "tag each as paid in cash or accrued")
                out.append({**base, "a": to_float(a), "b": None, "diff": None, "status": review, "note": why})
            continue
        row_a = cents(a) if a is not None else ZERO
        vs = (m or d).get("vs")
        if vs is not None:
            dd, st = classify(row_a, vs, kind, spec["tol"], rules)
            note = "tested against zero"
            if incomplete and st != "PASS":
                st, note = review, "tested against zero; not conclusive while lines are still left out"
            out.append({**base, "a": to_float(row_a), "b": vs, "diff": dd, "status": st, "note": note})
            continue
        if m and m.get("need") == "shares" and not ctx.has_shares:
            if b is not None:
                out.append({**base, "a": None, "b": to_float(b), "diff": None, "status": review,
                            "note": "NO_OF_SHARES is not in this file; shares cannot be compared"})
            continue
        if not have:
            out.append({**base, "a": to_float(row_a), "b": None, "diff": None, "status": "", "note": ""})
            continue
        if b is None and a is not None and spec["unmatched"] == "review":
            out.append({**base, "a": to_float(row_a), "b": None, "diff": None, "status": review,
                        "note": "no client line of this name; if it is a rename, match it by meaning"})
            continue
        if a is None and spec["unmatched"] == "review":
            out.append({**base, "a": None, "b": to_float(b or ZERO), "diff": None, "status": review,
                        "note": "no upload line of this name; if it is a rename, match it by meaning"})
            continue
        bb = cents(b) if b is not None else ZERO
        note = "not on the client's report" if b is None else "not in the upload" if a is None else ""
        if spec["unmatched"] == "error" and (a is None or b is None):
            note += " (a date or an amount differs)"
        dd, st = classify(row_a, bb, kind, spec["tol"], rules)
        out.append({**base, "a": to_float(row_a), "b": to_float(bb), "diff": dd, "status": st, "note": note})
    for it in issues:
        out.append({"label": it["label"], "a": None, "b": it["b"], "diff": None, "kind": "money", "status": review,
                    "note": it["why"], "mid": ""})
    if no_agn:
        for r in out:
            if r["mid"] in ("called", "unf") and r["status"] in ("", "PASS"):
                r["status"], r["note"] = review, "AGNST_COMMIT missing: every contribution counted"
    if spec["id"] == "irr" and have:
        out.extend(_irr_rows(items, report or [], aliases, ctx, rules))
    if spec["id"] == "unf" and have and not any(k[1] == "commit" for k in cl_items):
        notes.append("The client's report gave no Commitment per investor, so the unfunded commitment cannot be "
                     "worked out here; contributions called and recallable distributions are shown.")
    return {"rows": out, "state": "run" if have else "needs-report", "notes": notes}


def _collapse(items: Items) -> Items:
    """The client's table has no Trans Type: compare per investor and date across Trans Types."""
    out: Items = {}
    for (k, mid), v in items.items():
        kk = (k[0], k[1], "") if len(k) == 3 else k
        out[(kk, mid)] = out.get((kk, mid), ZERO) + v
    return out


# ---- investor IRR (step 5 of the guide) -------------------------------------------------------------------
def xirr(flows: list[tuple[date, float]]) -> float | None:
    """Annualised internal rate of return of dated flows (actual/365), by bisection; None when it has no root."""
    if len(flows) < 2 or not any(a > 0 for _, a in flows) or not any(a < 0 for _, a in flows):
        return None
    t0 = min(d for d, _ in flows)

    def npv(rate: float) -> float:
        return sum(a / (1 + rate) ** ((d - t0).days / 365.0) for d, a in flows)

    lo, hi = -0.99, 10.0
    if npv(lo) * npv(hi) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return round((lo + hi) / 2 * 100, 2)


def _irr_rows(items: Items, report: list[dict[str, Any]], aliases: Mapping[str, str], ctx: Ctx,
              rules: Rules) -> list[dict[str, Any]]:
    out = []
    term = {aliases.get(r["investor"].casefold(), r["investor"]): r for r in report
            if r["line"].casefold() == "ending capital" and r["amount"] is not None}
    shown = {aliases.get(r["investor"].casefold(), r["investor"]): to_float(r["amount"]) for r in report
             if r["line"].casefold() == "irr" and r["amount"] is not None}
    for inv, r in term.items():
        d = parse_date(r["date"])
        if d is None:
            continue
        flows = [(parse_date(k[1]), float(v)) for (k, mid), v in items.items()
                 if mid == "cf" and k[0] == inv and parse_date(k[1])]
        flows.append((d, float(items.get(((inv,), "end"), to_decimal(r["amount"])))))
        rate = xirr([(x, v) for x, v in flows if x])
        name = ctx.investors.get(inv, inv)
        if rate is None:
            out.append({"label": f"{name} · IRR", "a": None, "b": shown.get(inv), "diff": None, "kind": "pct",
                        "status": rules.level("needs_judgement"),
                        "note": "no rate can be worked out from these flows"})
        elif inv in shown:
            dd, st = classify(rate, shown[inv], "pct", "flat", rules)
            out.append({"label": f"{name} · IRR (%)", "a": rate, "b": shown[inv], "diff": dd, "kind": "pct",
                        "status": st, "note": ""})
        else:
            out.append({"label": f"{name} · IRR (%)", "a": rate, "b": None, "diff": None, "kind": "pct",
                        "status": "", "note": "the client's table shows no IRR to compare"})
    return out


# ---- replay: the known GL fixtures of the skill's cases ----------------------------------------------------
def _fixture_rows(doc: Mapping[str, Any], name: str, coa: Mapping[str, Mapping[str, str]]) -> list[Row]:
    fx = doc["fixtures"][name]
    rows: list[Row] = []
    for ln in fx["lines"]:
        acct = doc["accounts"][ln["acct"]]
        c = coa[acct]
        rows.append(Row(acct, c["FS_Grouping"], c["FS_Line_Item"], to_float(to_decimal(ln["dr"]) - to_decimal(ln["cr"])),
                        parse_date(ln.get("date")), ln.get("inv") or "", ln.get("tt") or "", ln.get("invst") or "",
                        parse_date(ln.get("pr")), ln.get("shares"), _yes(ln.get("agn") or "")))
    if fx.get("offsets"):
        rows += [Row(r.acct, r.fsg, r.fsl, -r.amt, r.date, "", r.tt, offset=True) for r in rows]
    return rows


def replay_case(rules: Rules, doc: Mapping[str, Any], case: Mapping[str, Any],
                coa: Mapping[str, Mapping[str, str]]) -> dict[str, Any]:
    """Run one recorded case with these rules: the state, the worst flag and every figure by label."""
    spec = {s["id"]: s for s in rules.specs}[case["check"]]
    rows = _fixture_rows(doc, case["ledger"], coa)
    other = _fixture_rows(doc, case["fund_ledger"], coa) if case.get("fund_ledger") else None
    ctx = make_ctx(rows, doc.get("investors", {}))

    def resolve(line: str, _tt: str) -> str:
        t = line.strip().split(" ")[0]
        return t if t in coa else ""

    res = run_check(spec, rows, ctx, coa, case.get("report"), resolve, doc.get("aliases", {}), other, rows,
                    incomplete=bool(case.get("incomplete")))
    return {"state": res["state"], "worst": worst(res["rows"]),
            "figures": {r["label"]: r["a"] for r in res["rows"]}}


def case_signature(result: Mapping[str, Any], labels: Iterable[str]) -> str:
    """What a case is compared on: the state, the worst flag and the figures it names."""
    figures = {k: result["figures"].get(k) for k in labels}
    return json.dumps({"state": result["state"], "worst": result["worst"], "figures": figures}, sort_keys=True)


# ---- storage of the client's reports -----------------------------------------------------------------------
def report_dir(job_dir: Path) -> Path:
    return job_dir / "recon_reports"


def load_report(job_dir: Path, check: str) -> dict[str, Any] | None:
    p = report_dir(job_dir) / f"{check}.json"
    if not p.is_file():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def save_report(job_dir: Path, check: str, name: str, data: bytes, who: str, audit: Any,
                rules: Rules | None = None) -> dict[str, Any]:
    """Read the file, then write the audit record, then store it: if the record cannot be written, nothing is kept."""
    if check not in SPEC or SPEC[check].get("pair"):
        raise ReportError("this check takes no client report")
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(name).name).strip(". ") or "report.xlsx"
    if Path(safe).suffix.lower() not in (".xlsx", ".csv", ".txt"):
        raise ReportError("send an .xlsx or .csv file in the layout of this check")
    d = report_dir(job_dir)
    d.mkdir(parents=True, exist_ok=True)
    raw = d / f"{check}{Path(safe).suffix.lower()}"
    raw_old = raw.read_bytes() if raw.is_file() else None
    raw.write_bytes(data)
    try:
        rows = read_report(raw, (rules or default_rules()).aliases)
    except Exception as exc:
        if raw_old is None:
            raw.unlink(missing_ok=True)
        else:
            raw.write_bytes(raw_old)
        raise ReportError(f"could not read the report: {exc}") from exc
    sha = hashlib.sha256(data).hexdigest()
    rules = rules or default_rules()
    meta = {"check": check, "file": safe, "sha256": sha, "by": who,
            "rows": [{**r, "amount": None if r["amount"] is None else to_float(r["amount"])} for r in rows],
            "skill": rules.stamp,
            "skill_hash": rules.digest,
            "at": datetime.now().strftime("%Y-%m-%d %H:%M")}
    try:
        audit.record("recon.report_uploaded",
                     f"{who} supplied the client's report for the {SPEC[check]['n']} check",
                     {"check": check, "file": safe, "sha256": sha, "skill": rules.stamp,
                      "skill_hash": rules.digest}, {"rows": len(rows)})
    except Exception:
        raw.unlink(missing_ok=True)
        if raw_old is not None:
            raw.write_bytes(raw_old)
        raise
    (d / f"{check}.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return meta


# ---- the blank layout, per check ---------------------------------------------------------------------------
def layout_labels(spec: Mapping[str, Any]) -> list[str]:
    if not spec.get("fixed_labels"):
        return []
    out = []
    for mid in spec["order"]:
        if mid in spec["measures"] and spec["measures"][mid].get("hidden"):
            continue
        if mid in spec.get("derived", {}) and spec["derived"][mid].get("vs") is not None:
            continue
        if mid in spec["measures"] and spec["measures"][mid].get("review"):
            continue
        lab = _label(spec, mid)
        if lab:
            out.append(lab)
    return out


def layout_bytes(check: str, rules: Rules | None = None) -> bytes:
    import openpyxl

    spec = {s["id"]: s for s in (rules or default_rules()).specs}[check]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Layout"
    ws.append(list(spec["layout"]))
    if "Line" in spec["layout"]:
        li = spec["layout"].index("Line")
        for lab in layout_labels(spec):
            row = [""] * len(spec["layout"])
            row[li] = lab
            ws.append(row)
    nt = wb.create_sheet("Notes")
    nt.append([f"{spec['n']}: the client's report"])
    nt.append([spec["layout_text"]])
    nt.append(["Rule: " + spec["rule"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def layout_info(spec: Mapping[str, Any]) -> dict[str, Any]:
    return {"columns": spec.get("layout", []), "text": spec.get("layout_text", ""),
            "labels": layout_labels(spec)}


def combined_breaks(checks: Sequence[Mapping[str, Any]], fund_type: str, rules: Rules | None = None
                    ) -> list[dict[str, Any]]:
    """The same difference in more than one check is one break. Which lines may join is the cross-check table;
    that equal absolute differences are one break, and that it takes two checks, is the engine's rule."""
    r = rules or default_rules()
    seen: dict[float, list[str]] = {}
    for c in checks:
        if fund_type and fund_type not in c["scope"] or c["state"] != "run":
            continue
        for row in c["rows"]:
            if row.get("diff") is None:
                continue
            vs_zero = str(row.get("note", "")).startswith("tested against zero")
            if r.combine(row["status"], row["kind"], row.get("b") is not None, vs_zero, row["diff"]):
                names = seen.setdefault(abs(row["diff"]), [])
                if c["n"] not in names:
                    names.append(c["n"])
    return [{"amount": amt, "checks": names} for amt, names in sorted(seen.items(), reverse=True) if len(names) > 1]


CONTROLS_FILE = "controls.yaml"


def _controls_text_v1(text: str) -> str:
    """A controls text recorded before the files had a header (`version: 1`), as a v1 file: the rolled-back file
    must load strictly."""
    data = yaml.safe_load(text)
    if not isinstance(data, dict) or "schema_version" in data:
        return text
    header = "schema_version: 1\nkind: controls\n"
    if re.search(r"^version: 1[ \t]*$", text, re.MULTILINE):
        return re.sub(r"^version: 1[ \t]*\n", header, text, count=1, flags=re.MULTILINE)
    lines = text.splitlines(keepends=True)
    n = 0
    while n < len(lines) and lines[n].startswith("#"):
        n += 1
    return "".join(lines[:n]) + header + "".join(lines[n:])


def apply_controls(skill_dir: Path, new: Mapping[str, Any], proposer: str, controller: str, reason: str,
                   audit: Any, *, controllers: Iterable[str] | None = None,
                   approver_roles: Sequence[str] = ("controller",), proposal: Mapping[str, Any] | None = None,
                   rollback_of: Mapping[str, Any] | None = None, text: str | None = None) -> dict[str, Any]:
    """Change the controls: the one path that writes `controls.yaml` (behind `improver.approve_control`).

    Only a named controller from `controllers` who is not the proposer may (a rollback of the latest change may be
    made by any controller). Values outside CONTROL_LIMITS are refused. Order: validate, compute the pin the skill
    will have, write the `control.changed` record (bound, direction, pins, proposal hash, controllers list hash),
    write the file, re-hash; a file that does not read back as approved is put back and `control.write_failed` is
    recorded. If the record cannot be written, nothing changes. The improver has no path here."""
    from worker.audit import AuditError
    from worker.improver import controller_refusal, controllers_sha256

    names = sorted({c.strip().lower() for c in controllers or () if c and c.strip()})
    refusal = controller_refusal(controller, proposer, names, rollback=rollback_of is not None)
    if refusal:
        raise PermissionError(refusal)
    if not reason.strip():
        raise ValueError("a reason is required to change the controls")
    if audit is None:
        raise AuditError("a control change needs an audit log; without one the change does not happen")
    path = skill_dir / CONTROLS_FILE
    old_bytes = path.read_bytes()
    old_text = old_bytes.decode("utf-8")
    old = yaml.safe_load(old_text) or {}
    if text is not None:
        text = _controls_text_v1(text)
    # the file keeps its header whoever built `new`: a controls file without one is refused at the next load
    head = {k: old[k] for k in ("schema_version", "kind") if k in old}
    new = {**head, **{k: v for k, v in dict(new).items() if k != "version"}}
    ids = [d["id"] for d in yaml.safe_load((skill_dir / "specs.yaml").read_text(encoding="utf-8"))["checks"]]
    validate_controls(new, ids)
    if text is None:
        lines = old_text.splitlines(keepends=True)
        n = 0
        while n < len(lines) and lines[n].startswith("#"):
            n += 1
        text = "".join(lines[:n]) + yaml.safe_dump(dict(new), sort_keys=False)
    elif yaml.safe_load(text) != dict(new):
        raise ValueError("the controls text does not say the same as the new controls")
    new_bytes = text.encode("utf-8")

    def flat(d: Mapping[str, Any], prefix: str = "") -> dict[str, Any]:
        out: dict[str, Any] = {}
        for k, v in (d or {}).items():
            if isinstance(v, Mapping):
                out.update(flat(v, f"{prefix}{k}."))
            else:
                out[f"{prefix}{k}"] = sorted(str(x) for x in v) if isinstance(v, list) else v
        return out

    before, after = flat(old), flat(new)
    changed = sorted(k for k in set(before) | set(after) if before.get(k) != after.get(k))
    if not changed:
        raise ValueError("nothing in the controls changes")

    def judge(key: str, a: Any, b: Any) -> tuple[dict[str, Any], str]:
        part = key.split(".")
        lim = CONTROL_LIMITS
        if a is None or b is None:
            return {}, "looser"
        if part[0] == "tolerances" and len(part) == 3 and part[1] in lim["tolerances"]:
            max_pass, max_warn, same = lim["tolerances"][part[1]]
            bound = {"max": str(max_pass if part[2] == "pass_up_to" else max_warn), "warn_equals_pass": same}
            return bound, "looser" if to_decimal(b) > to_decimal(a) else "stricter"
        if part[0] == "levels":
            return ({"allowed": list(lim["levels"].get(part[1], ()))},
                    "looser" if ORDER.get(str(b), 0) < ORDER.get(str(a), 0) else "stricter")
        if part[0] == "check_style":
            return {"allowed": list(TIER_STYLES)}, "looser" if b == "tiered" else "stricter"
        if part[0] == "required_fields":
            return ({"floor": sorted(lim["required_floor"].get(part[1], set())), "fields": sorted(lim["fields"])},
                    "looser" if set(a) - set(b) else "stricter")
        if part[0] == "balance":
            return {"equals": str(lim["balance_target"])}, "looser"
        return {}, "looser"  # an unknown key is treated as loosening

    judged = {k: judge(k, before.get(k), after.get(k)) for k in changed}
    direction = "looser" if any(d == "looser" for _b, d in judged.values()) else "stricter"

    def pin(override: bytes | None) -> str:  # the same content hash a job records (Skill.content_hash)
        h = hashlib.sha256()
        for p in sorted(p for p in skill_dir.rglob("*") if p.is_file() and "__pycache__" not in p.parts):
            data = override if override is not None and p == path else p.read_bytes()
            h.update(p.relative_to(skill_dir).as_posix().encode())
            h.update(b"\0" + data.replace(b"\r\n", b"\n") + b"\0")
        return h.hexdigest()[:12]

    def sha(b: bytes) -> str:
        return hashlib.sha256(b.replace(b"\r\n", b"\n")).hexdigest()

    pin_before = pin(None)
    if pin_before != Skill(SKILL, "", "", skill_dir).content_hash():
        raise RulesError("the skill's pin cannot be computed the way jobs record it; nothing changed")
    pin_after = pin(new_bytes)
    prop = dict(proposal or {})
    inputs: dict[str, Any] = {
        "proposal_id": str(prop.get("id", "")), "proposal_sha256": str(prop.get("sha256", "")),
        "proposed_by": proposer.strip(), "proposed_at": str(prop.get("proposed_at", "")),
        "approved_by": controller.strip(), "approver_roles": list(approver_roles), "skill": SKILL,
        "control_id": changed[0] if len(changed) == 1 else changed,
        "old": {k: before.get(k) for k in changed}, "new": {k: after.get(k) for k in changed},
        "bound": {k: b for k, (b, _d) in judged.items()}, "direction": direction,
        "directions": {k: d for k, (_b, d) in judged.items()},
        "evidence": prop.get("evidence") or {},
        "replay": {"info_only": True, "result": prop.get("replay")},
        "pin_before": pin_before, "pin_after": pin_after,
        "controls_sha256_before": sha(old_bytes), "controls_sha256_after": sha(new_bytes),
        "controllers": names, "controllers_list_sha256": controllers_sha256(names),
        "controls_text_before": old_text,
    }
    action = "control.changed"
    if rollback_of is not None:
        action = "control.rolled_back"
        inputs["reverses_seq"] = rollback_of.get("seq")
    entry = audit.record(action, reason, inputs, {"pin_after": pin_after,
                                                  "controls_sha256_after": sha(new_bytes)})
    try:
        path.write_bytes(new_bytes)
        now = pin(None)
    except OSError as exc:
        now = f"not written ({exc})"
    if now != pin_after:
        path.write_bytes(old_bytes)
        audit.record("control.write_failed", f"controls.yaml did not read back as approved ({now}); the old file "
                     "was put back", {"reverses_seq": entry["seq"], "pin_after": pin_after, "found": now},
                     {"restored_pin": pin(None)})
        raise RulesError("the controls file did not read back as approved; the old file was put back")
    return {"controls": dict(new), "pin_before": pin_before, "pin_after": pin_after, "seq": entry["seq"],
            "direction": direction}
