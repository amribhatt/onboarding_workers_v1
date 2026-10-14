"""The data-quality engine: what a problem in the data is, and where each one is in the source file.

Rules are data (`skills/check-data-quality/checks.yaml`), grouped in families: completeness, validity, uniqueness,
consistency, balance. The file is streamed once; rows that fail the same rule on the same field become ONE cause.
Every cause carries a locator computed here, never by a model: the file, sheet and row ranges, and the Excel filter
that finds the rows again, with an optional breakdown (a GL account, an entity) of its own filters and ranges.
The two writers (marked copy, just these rows) also stream, so a million-row file is never held in memory.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from collections import Counter
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from worker import columnar, rulefiles
from worker.money import balance_status, to_decimal
from worker.skills import skills_dir

MARK_COLUMN = "Worker check"
SAMPLE_ROWS = 20
MAX_VALUES = 200
MAX_RANGES_SHOWN = 12
FAMILIES = ("completeness", "validity", "uniqueness", "consistency", "balance")
_HIGHLIGHT = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d-%b-%Y", "%Y/%m/%d", "%d.%m.%Y")


@dataclass(frozen=True)
class Rule:
    id: str
    family: str
    op: str
    fields: tuple[str, ...] | str
    severity: str
    message: str
    fix: str = ""
    record_types: tuple[str, ...] = ()
    when: Mapping[str, Any] | None = None
    breakdown: str = ""
    options: Mapping[str, Any] = field(default_factory=dict)


_OPTION_KEYS = ("max", "set", "group", "debit", "credit", "offset")  # what an op may carry beyond the common fields


def load_rules(path: str | Path | None = None) -> list[Rule]:
    """The rules of `checks.yaml`, read by the strict loader: a file with no header, an unknown field, an unknown
    family, severity or op, or a rule that is looser than its control is refused (`rulefiles.SchemaError`). A key
    the schema does not know is never passed on as an option."""
    p = Path(path) if path else skills_dir() / "check-data-quality" / "checks.yaml"
    data = rulefiles.load_file(p).data
    rules = []
    for r in data["rules"]:
        if "op" not in r:  # a guard names the engine code that enforces it; this engine evaluates ops only
            continue
        if r["family"] not in FAMILIES:
            raise ValueError(f"rule {r['id']}: family must be one of {FAMILIES}")
        fields = r.get("fields", ())
        options = {k: r[k] for k in _OPTION_KEYS if k in r}
        if "tolerance" in r:  # a journal that is out by no more than this is balanced
            options["tol_pass"] = options["tol_warn"] = r["tolerance"]
        rules.append(
            Rule(
                id=r["id"], family=r["family"], op=r["op"],
                fields=fields if isinstance(fields, str) else tuple(fields),
                severity=r["severity"], message=r["message"], fix=r.get("fix", ""),
                record_types=tuple(r.get("record_types", ())), when=r.get("when"),
                breakdown=r.get("breakdown", ""), options=options,
            )
        )
    return rules


# ---- reading a source file as a stream -----------------------------------------------------------------
def column_letter(index: int) -> str:
    """0-based column index -> A, B, ... AA."""
    return get_column_letter(index + 1)


def _norm(text: Any) -> str:
    return " ".join(str(text).split()).lower()


class Source:
    """A CSV or XLSX sheet read as (row number, values) without loading it: header row first, then data."""

    def __init__(self, path: str | Path, sheet: str | None = None, header_row: int = 1) -> None:
        self.path = Path(path)
        self.sheet = sheet
        self.header_row = header_row
        self.is_csv = self.path.suffix.lower() == ".csv"
        self.headers: list[str] = []
        self.sheet_name = sheet or ("" if self.is_csv else None)
        self.headers = [str(v).strip() if v is not None else "" for v in self._header()]

    def stream(self) -> Iterator[tuple[int, list[Any]]]:
        if not self.is_csv:
            self.sheet_name = self.sheet or columnar.first_sheet(self.path)
        yield from columnar.rows(self.path, self.sheet_name if not self.is_csv else None)

    def _header(self) -> list[Any]:
        for n, row in self.stream():
            if n == self.header_row:
                return row
            if n > self.header_row:
                break
        return []

    def rows(self) -> Iterator[tuple[int, list[Any]]]:
        """Data rows after the header; wholly blank rows are skipped."""
        for n, row in self.stream():
            if n <= self.header_row:
                continue
            if any(v is not None and str(v).strip() != "" for v in row):
                yield n, row

    def index_of(self, header: str) -> int | None:
        want = _norm(header)
        for i, h in enumerate(self.headers):
            if _norm(h) == want:
                return i
        return None


# ---- row ranges -----------------------------------------------------------------------------------------
class Ranges:
    """Row numbers as runs: 1204-7203, 256000-258999; stragglers are merged. Flat slots: a GL holds 100,000s."""
    __slots__ = ("count", "late", "runs")
    def __init__(self) -> None:
        self.runs: list[int] = []  # a, b, a, b, ...
        self.late: list[int] | None = None
        self.count = 0

    def add(self, n: int) -> None:
        self.count += 1
        if not self.runs or n > self.runs[-1] + 1:
            self.runs += (n, n)
        elif n == self.runs[-1] + 1:
            self.runs[-1] = n
        else:  # out of order (or repeated): merged, de-duplicated, when the ranges are read
            self.late = self.late or []
            self.late.append(n)

    def final(self) -> list[tuple[int, int]]:
        out: list[list[int]] = []
        for a, b in sorted([*zip(self.runs[::2], self.runs[1::2]), *((r, r) for r in set(self.late or ()))]):
            if out and a <= out[-1][1] + 1:
                out[-1][1] = max(out[-1][1], b)
            else:
                out.append([a, b])
        return [(a, b) for a, b in out]


class _Journal(Ranges):  # a journal's rows and its net amount
    __slots__ = ("net",)
    def __init__(self) -> None:
        super().__init__()
        self.net = Decimal(0)


def ranges_text(runs: Sequence[tuple[int, int]], shown: int = MAX_RANGES_SHOWN) -> str:
    parts = [f"{a}" if a == b else f"{a}–{b}" for a, b in runs[:shown]]
    more = len(runs) - shown
    return ", ".join(parts) + (f" and {more} more ranges" if more > 0 else "")


def values_text(values: Sequence[str], shown: int = 3) -> str:
    vs = list(values)
    head = ", ".join(vs[:shown])
    return head + (f" and {len(vs) - shown} more" if len(vs) > shown else "")


# ---- cell checks -----------------------------------------------------------------------------------------
def _blank(v: Any) -> bool:
    return v is None or str(v).strip() == ""


def _is_date(v: Any) -> bool:
    if isinstance(v, dt.datetime | dt.date):
        return True
    s = str(v).strip()
    for fmt in _DATE_FORMATS:
        try:
            dt.datetime.strptime(s, fmt)
            return True
        except ValueError:
            continue
    try:
        dt.datetime.fromisoformat(s)
        return True
    except ValueError:
        return False


def _is_number(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, int | float | Decimal):
        return True
    s = str(v).strip().replace(",", "")
    if s.startswith("(") and s.endswith(")"):
        s = s[1:-1]
    try:
        Decimal(s)
        return True
    except InvalidOperation:
        return False


def _in_ranges(value: Any, ranges: Sequence[str]) -> bool:
    m = re.match(r"\s*(\d+)", str(value))
    if not m:
        return False
    head = m.group(1)
    for r in ranges:
        lo, hi = (int(x) for x in str(r).split("-"))
        width = len(str(lo))
        if lo <= int(head[:width] if len(head) >= width else head) <= hi:
            return True
    return False


def _natural(values: Sequence[str]) -> list[str]:
    return sorted(values, key=lambda s: (not s.isdigit(), int(s) if s.isdigit() else 0, s))


# ---- scanning ---------------------------------------------------------------------------------------------
@dataclass
class _Acc:
    rule: Rule
    field: str
    ranges: Ranges = field(default_factory=Ranges)
    samples: list[dict[str, Any]] = field(default_factory=list)
    values: Counter[str] = field(default_factory=Counter)
    when_values: Counter[str] = field(default_factory=Counter)
    breakdown: dict[str, Ranges] = field(default_factory=dict)
    seen: Counter[str] = field(default_factory=Counter)  # non-blank values of the field (blank-date suggestion)
    no_column: bool = False


def expand_fields(rule: Rule, groups: Mapping[str, Sequence[str]]) -> list[str]:
    if isinstance(rule.fields, str):
        key = rule.fields.lstrip("@")
        return list(groups.get(key, ()))
    return list(rule.fields)


def groups_from_spec(spec: Any, required: Sequence[str] | None = None) -> dict[str, list[str]]:
    """Field groups (`@required`, `@ids`, `@names`, `@dates`, `@amounts`, `@key`) from a TemplateSpec."""
    fields = list(spec.fields)
    return {
        "required": list(required if required is not None else spec.required),
        "ids": [f for f in fields if f.endswith("_ID") and f != "GP_ID"],
        "names": [f for f in fields if f.endswith("_Name")],
        "dates": [f for f in fields if f.endswith("_Date")],
        "amounts": [f for f in fields if f.endswith("_Amount")],
        "key": list(spec.key),
    }


def scan(
    path: str | Path,
    rules: Sequence[Rule],
    *,
    columns: Mapping[str, str],
    record_type: str,
    sheet: str | None = None,
    header_row: int = 1,
    groups: Mapping[str, Sequence[str]] | None = None,
    references: Mapping[str, set[str]] | None = None,
    levels: Mapping[str, str] | None = None,
    supplied: Mapping[tuple[int, str], Any] | None = None,
    supplied_all: Mapping[tuple[str, str], Any] | None = None,
    default_breakdown: str = "",
    generated: Sequence[str] = (),
) -> dict[str, Any]:
    """Stream `path` once and return the report: causes (one per rule and field), counts and locators.

    `columns` maps each logical field to its header in the file. `supplied` ({(row, field): value}) and
    `supplied_all` ({(rule id, field): value}) are values a person typed; they replace blanks only, and the check
    runs again with them. `generated` lists ID fields the worker fills when blank (never a completeness error).
    """
    src = Source(path, sheet, header_row)
    groups = groups or {}
    references = references or {}
    supplied = supplied or {}
    supplied_all = supplied_all or {}
    index = {f: src.index_of(h) for f, h in columns.items()}
    accs: list[_Acc] = []
    plan: list[tuple[_Acc, Rule]] = []
    balance: list[tuple[Rule, dict[str, _Journal]]] = []
    for rule in rules:
        if rule.record_types and record_type not in rule.record_types:
            continue
        if rule.op == "balance":
            balance.append((rule, {}))
            continue
        for fld in expand_fields(rule, groups):
            if rule.op == "blank" and fld in generated:
                continue
            acc = _Acc(rule, fld, no_column=index.get(fld) is None)
            accs.append(acc)
            plan.append((acc, rule))
    seen_keys: dict[tuple[str, str], dict[str, int]] = {}
    total = 0
    cols = {f: i for f, i in index.items() if i is not None}
    sample_fields = list(cols)[:8]
    for n, row in src.rows():
        total += 1

        def get(fld: str, _row: list[Any] = row, _n: int = n) -> Any:
            i = cols.get(fld)
            v = _row[i] if i is not None and i < len(_row) else None
            if _blank(v):
                if (_n, fld) in supplied:
                    return supplied[(_n, fld)]
                for (_rid, f2), val in supplied_all.items():
                    if f2 == fld:
                        return val
            return v

        for acc, rule in plan:
            if acc.no_column:
                continue
            value = get(acc.field)
            if rule.op == "blank" and not _blank(value) and len(acc.seen) < 3:
                acc.seen[str(value).strip()] += 1
            bad = False
            when_field = (rule.when or {}).get("field")
            if rule.op == "blank":
                if rule.when is not None:
                    gate = get(when_field) if when_field else None
                    if not _in_ranges(gate, rule.when.get("ranges", [])):
                        continue
                bad = _blank(value)
            elif _blank(value):
                continue
            elif rule.op == "max_length":
                bad = len(str(value).strip()) > int(rule.options["max"])
            elif rule.op == "date":
                bad = not _is_date(value)
            elif rule.op == "number":
                bad = not _is_number(value)
            elif rule.op == "unique":
                key = _norm(value)
                firsts = seen_keys.setdefault((rule.id, acc.field), {})
                if key in firsts:
                    first = firsts[key]
                    if first:
                        acc.ranges.add(first)
                        acc.values[str(value).strip()] += 1
                        firsts[key] = 0  # the first row is reported once
                    bad = True
                else:
                    firsts[key] = n
            elif rule.op == "in_set":
                bad = _norm(value) in references.get(rule.options.get("set", "existing"), set())
            elif rule.op == "reference":
                refs = references.get(rule.options.get("set", ""), None)
                bad = refs is not None and _norm(value) not in refs
            if not bad:
                continue
            acc.ranges.add(n)
            if len(acc.samples) < SAMPLE_ROWS:
                acc.samples.append({"row": n, "values": {f: _show(get(f)) for f in sample_fields}})
            shown = "(Blanks)" if rule.op == "blank" else str(value).strip()
            if len(acc.values) < MAX_VALUES or shown in acc.values:
                acc.values[shown] += 1
            if when_field:
                wv = str(get(when_field)).strip()
                if len(acc.when_values) < MAX_VALUES or wv in acc.when_values:
                    acc.when_values[wv] += 1
            bd_field = rule.breakdown or default_breakdown
            if bd_field and bd_field in cols:
                bv = str(get(bd_field)).strip() or "(Blanks)"
                if bv in acc.breakdown or len(acc.breakdown) < MAX_VALUES:
                    acc.breakdown.setdefault(bv, Ranges()).add(n)
        for rule, journals in balance:
            g, d, c = get(rule.options["group"]), get(rule.options["debit"]), get(rule.options["credit"])
            if _blank(g):
                continue
            key = str(g).strip()
            entry = journals.get(key) or journals.setdefault(key, _Journal())
            offset = rule.options.get("offset")  # an LP line with an investor is offset on the fund's side
            if offset and offset in cols and not _blank(get(offset)):
                d = c = 0
            try:
                entry.net += to_decimal(d) - to_decimal(c)
            except ValueError:  # a cell that is not a number adds nothing; the `number` rule reports it
                pass
            entry.add(n)
    causes = [c for c in (_cause(a, src, columns, cols, total, levels, default_breakdown) for a in accs) if c]
    for rule, journals in balance:
        tol_pass, tol_warn = rule.options.get("tol_pass", 0), rule.options.get("tol_warn", 0)
        bad_journals = {k: v for k, v in journals.items() if balance_status(v.net, 0, tol_pass, tol_warn)[0] == "err"}
        if not bad_journals:
            continue
        acc = _Acc(rule, rule.options["group"])
        for k in sorted(bad_journals):
            for a, b in bad_journals[k].final():
                for n in range(a, b + 1):
                    acc.ranges.add(n)
            acc.values[k] = bad_journals[k].count
        causes.append(_cause(acc, src, columns, cols, total, levels, default_breakdown, balance=bad_journals))
    causes = [c for c in causes if c]
    causes.sort(key=lambda c: (c["severity"] != "ERR", -c["count"], c["id"]))
    return {
        "file": src.path.name,
        "path": str(src.path),
        "sheet": src.sheet_name or "",
        "header_row": header_row,
        "rows": total,
        "columns": {f: {"header": columns[f], "letter": column_letter(i)} for f, i in cols.items()},
        "causes": causes,
        "err": sum(1 for c in causes if c["severity"] == "ERR"),
        "warn": sum(1 for c in causes if c["severity"] == "WARN"),
        "families": dict(Counter(c["family"] for c in causes)),
    }


def _show(v: Any) -> Any:
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def _cond_text(label: str, letter: str, header: str, values: Sequence[str]) -> str:
    return f"column {letter} “{header}” = {values_text(values)}"


def _cause(
    acc: _Acc,
    src: Source,
    columns: Mapping[str, str],
    cols: Mapping[str, int],
    total: int,
    levels: Mapping[str, str] | None,
    default_breakdown: str,
    balance: Mapping[str, Any] | None = None,
) -> dict[str, Any] | None:
    rule, fld = acc.rule, acc.field
    # a level from the template or the sponsor may raise a rule; it never goes below the rule's own severity, which
    # is itself held at or above the floor of controls.yaml (looser_than_control at load)
    chosen = (levels or {}).get(rule.id)
    severity = rule.severity if chosen not in ("ERR", "WARN") or rule.severity == "ERR" else chosen
    base = {
        "id": f"{rule.family}.{rule.id}:{fld}", "rule": rule.id, "family": rule.family, "severity": severity,
        "field": fld,
    }
    if acc.no_column:
        if rule.op != "blank":
            return None
        header = columns.get(fld, fld)
        return {
            **base, "message": f"This file has no “{header}” column", "fix": rule.fix.format(field=fld),
            "count": total, "row_ranges": [], "ranges_text": "all rows", "sample_rows": [], "no_column": True,
            "locator": {"file": src.path.name, "sheet": src.sheet_name or "", "header_row": src.header_row,
                        "filter": "", "where": f"{src.path.name} · no column for {fld}"},
            "breakdown": None, "suggestion": "", "copy_values": [],
        }
    runs = acc.ranges.final()
    count = sum(b - a + 1 for a, b in runs)
    if count == 0:
        return None
    letter = column_letter(cols[fld]) if fld in cols else ""
    header = columns.get(fld, fld)
    conditions: list[tuple[str, str, str, list[str]]] = []  # (field, letter, header, values)
    values = _natural(list(acc.values)) if rule.op != "blank" else ["(Blanks)"]
    if balance is not None:
        values = _natural(list(balance))
    conditions.append((fld, letter, header, values))
    when_field = (rule.when or {}).get("field")
    if when_field and when_field in cols:
        conditions.append((when_field, column_letter(cols[when_field]), columns.get(when_field, when_field),
                           _natural(list(acc.when_values))))
    filt = ", and ".join(_cond_text(f, lt, h, vs) for f, lt, h, vs in conditions)
    locator_filter = f"In Excel: Data → Filter, {filt} → {count:,} lines"
    row_text = ranges_text(runs)
    if len(runs) == 1 and runs[0][0] == runs[0][1]:
        n = runs[0][0]
        last = column_letter(len(src.headers) - 1) if src.headers else "A"
        where = f"{src.path.name} · sheet {src.sheet_name or '-'}, row {n} · {src.sheet_name or 'Sheet'}!A{n}:{last}{n}"
    else:
        where = f"{src.path.name} · sheet {src.sheet_name or '-'}, rows {row_text}"
    breakdown = None
    if len(acc.breakdown) > 1:
        bd_field = rule.breakdown or default_breakdown
        bl, bh = (column_letter(cols[bd_field]), columns.get(bd_field, bd_field)) if bd_field in cols else ("", bd_field)
        breakdown = []
        for value in _natural(list(acc.breakdown)):
            r = acc.breakdown[value]
            final = r.final()
            base_conds = [f"{h} = {'(Blanks)' if rule.op == 'blank' else values_text(vs, 99)}"
                          for f, _lt, h, vs in conditions if f != bd_field and f == fld]
            breakdown.append({
                "value": value, "count": r.count, "row_ranges": [list(x) for x in final],
                "ranges_text": ranges_text(final),
                "filter": f"column {bl} “{bh}” = {value}"
                          + (f", and column {letter} “{header}” = (Blanks)" if rule.op == "blank" and fld != bd_field else ""),
                "copy_filter": "; ".join([*base_conds, f"{bh} = {value}"]),
            })
    suggestion = ""
    if rule.op == "blank" and (fld.endswith("_Date") or "date" in header.lower()) and len(acc.seen) == 1:
        only = next(iter(acc.seen))
        suggestion = f"Use {only}, the {header.lower()} of the other rows in this batch"
    return {
        **base,
        "message": rule.message.format(field=fld, max=rule.options.get("max", "")),
        "fix": rule.fix.format(field=fld, n=count),
        "count": count,
        "row_ranges": [list(r) for r in runs],
        "ranges_text": row_text,
        "sample_rows": acc.samples,
        "locator": {
            "file": src.path.name, "sheet": src.sheet_name or "", "header_row": src.header_row,
            "conditions": [{"field": f, "column": lt, "header": h, "values": vs} for f, lt, h, vs in conditions],
            "filter": locator_filter, "where": where,
        },
        "breakdown": breakdown,
        "suggestion": suggestion,
        "copy_values": values if rule.op != "blank" else (conditions[1][3] if len(conditions) > 1 else []),
        "more_samples_in_marked_copy": count > len(acc.samples),
    }


# ---- values a person supplies ---------------------------------------------------------------------------
@dataclass
class Supplied:
    """Values typed by a person for blanks. Applied to blanks only; always recorded as the person's own."""

    rows: dict[tuple[int, str], Any] = field(default_factory=dict)
    all: dict[tuple[str, str], Any] = field(default_factory=dict)
    log: list[dict[str, Any]] = field(default_factory=list)

    def set_row(self, row: int, fld: str, value: Any, who: str, reason: str = "") -> None:
        if _blank(value):
            raise ValueError("a supplied value cannot be blank")
        self.rows[(row, fld)] = value
        self.log.append({"row": row, "field": fld, "value": value, "by": who, "reason": reason or "supplied by hand",
                         "inferred": False})

    def set_all(self, rule_id: str, fld: str, value: Any, who: str, reason: str = "") -> None:
        if _blank(value):
            raise ValueError("a supplied value cannot be blank")
        self.all[(rule_id, fld)] = value
        self.log.append({"rule": rule_id, "field": fld, "value": value, "by": who, "scope": "all blank rows of the cause",
                         "reason": reason or "supplied by hand", "inferred": False})


# ---- streaming writers ---------------------------------------------------------------------------------
def _marks(report: Mapping[str, Any], cause_ids: Sequence[str] | None, only: tuple[str, str] | None
           ) -> list[tuple[str, list[tuple[int, int]]]]:
    chosen = []
    for c in report["causes"]:
        if cause_ids is not None and c["id"] not in cause_ids:
            continue
        if c.get("no_column"):
            continue
        runs = [tuple(r) for r in c["row_ranges"]]
        if only is not None:
            bd = [b for b in (c.get("breakdown") or []) if b["value"] == only[1]]
            if not bd:
                continue
            runs = [tuple(r) for r in bd[0]["row_ranges"]]
        chosen.append((c["message"], runs))
    return chosen


class _Marker:
    """Answers 'which causes mark row n' for ascending n, walking each cause's ranges once."""

    def __init__(self, marks: list[tuple[str, list[tuple[int, int]]]]) -> None:
        self.marks = marks
        self.pos = [0] * len(marks)

    def at(self, n: int) -> list[str]:
        out = []
        for k, (name, runs) in enumerate(self.marks):
            p = self.pos[k]
            while p < len(runs) and runs[p][1] < n:
                p += 1
            self.pos[k] = p
            if p < len(runs) and runs[p][0] <= n:
                out.append(name)
        return out


def marked_copy(
    report: Mapping[str, Any],
    out: str | Path,
    *,
    cause_ids: Sequence[str] | None = None,
    only: tuple[str, str] | None = None,
) -> Path:
    """The user's file untouched, plus a copy with a `Worker check` column, affected rows highlighted, filter on.

    `only` = (breakdown field, value) limits it to that value (e.g. one GL account). CSV sources get a CSV copy
    with the column (a CSV cannot hold highlighting); XLSX sources get a highlighted XLSX.
    """
    src_path, out = Path(report["path"]), Path(out)
    if out.resolve() == src_path.resolve():
        raise ValueError("the marked copy must be a different file; the original is never changed")
    src = Source(src_path, report["sheet"] or None, report["header_row"])
    marker = _Marker(_marks(report, cause_ids, only))
    out.parent.mkdir(parents=True, exist_ok=True)
    width = max(len(src.headers), 1)
    if src.is_csv:
        with open(out, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            for n, row in src.stream():
                if n == src.header_row:
                    w.writerow([*row, MARK_COLUMN])
                elif n > src.header_row:
                    w.writerow([*row, "; ".join(marker.at(n))])
                else:
                    w.writerow(row)
        return out
    book = openpyxl.Workbook(write_only=True)
    ws = book.create_sheet(report["sheet"] or "Sheet1")
    count = 0
    for n, row in src.stream():
        count = n
        padded = [*row, *([None] * (width - len(row)))][:width]
        if n == src.header_row:
            ws.append([*padded, MARK_COLUMN])
        elif n > src.header_row:
            names = marker.at(n)
            if names:
                cells = []
                for v in padded:
                    c = WriteOnlyCell(ws, value=v)
                    c.fill = _HIGHLIGHT
                    cells.append(c)
                last = WriteOnlyCell(ws, value="; ".join(names))
                last.fill = _HIGHLIGHT
                ws.append([*cells, last])
            else:
                ws.append([*padded, None])
        else:
            ws.append([*padded, None])
    ws.auto_filter.ref = f"A{src.header_row}:{column_letter(width)}{max(count, src.header_row)}"
    book.save(out)
    return out


def just_these_rows(
    report: Mapping[str, Any],
    out: str | Path,
    *,
    cause_ids: Sequence[str] | None = None,
    only: tuple[str, str] | None = None,
) -> Path:
    """Only the affected rows, each with its row number in the original file (`Original row` first)."""
    src_path, out = Path(report["path"]), Path(out)
    src = Source(src_path, report["sheet"] or None, report["header_row"])
    marker = _Marker(_marks(report, cause_ids, only))
    out.parent.mkdir(parents=True, exist_ok=True)

    def lines() -> Iterator[list[Any]]:
        for n, row in src.stream():
            if n == src.header_row:
                yield ["Original row", *row, MARK_COLUMN]
            elif n > src.header_row:
                names = marker.at(n)
                if names:
                    yield [n, *row, "; ".join(names)]

    if src.is_csv:
        with open(out, "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerows(lines())
        return out
    book = openpyxl.Workbook(write_only=True)
    ws = book.create_sheet(report["sheet"] or "Sheet1")
    for line in lines():
        ws.append(line)
    book.save(out)
    return out

